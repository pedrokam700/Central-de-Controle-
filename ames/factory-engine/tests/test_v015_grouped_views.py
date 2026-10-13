import unittest, tempfile
from pathlib import Path
from openpyxl import load_workbook
from reports.relatorio_v015 import gerar_excel_integrado

class V015GroupedViewsTests(unittest.TestCase):
    def _traces(self):
        r3074={"items":[{
            "pcba_sn":"CUR",
            "base_query":{"rows":[
                {"Barcode":"CUR","Aut Cmf7":"C1","Bind Time":"2026-10-01 10:00:00","Unbind Time":"2026-10-02 10:00:00","Iss Status":"N"},
                {"Barcode":"CUR","Aut Cmf7":"C2","Bind Time":"2026-10-03 10:00:00","Unbind Time":"","Iss Status":"Y"},
            ]},
            "defect_traces":[{
                "status":"COMPLETE",
                "defect":{"line":"L3","product_model":"CPH","defect_time":"2026-10-03 11:00:00","defect_code":"CAM","defect_desc":"Camera failed","repair_comment":""},
                "trace_3074":{"summary":{"componentes_no_momento_da_falha":1},"results":[{
                    "item_sn":"MAT","item_type":"Rear Camera","status":"SEGUNDO_USO_NO_MOMENTO_DA_FALHA","previous_pcba_count":1,
                    "previous_pcbas":["OLD"],"pcbas_until_defect":["OLD","CUR"],"current_pcba_confirmed":True,
                    "total_pcba_count_known_now":2,"reuse_count_known_now":1,"inactive_now_pcbas":["OLD"],"active_now_pcbas":["CUR"],
                    "all_pcba_history":[{"pcba_sn":"OLD","iss_status_now":"N"},{"pcba_sn":"CUR","iss_status_now":"Y"}],
                }]}
            }]
        }]}
        r2114={"pcbas":[
            {"pcba_sn":"CUR","status":"COMPLETE","result":{"shift":"1st Shift","rows":[
                {"Defect Hist Seq":"0","Defect Code":"CAM","Defect Description":"Camera failed","Repair Status":"Y","Defect Type":"MainBoard"},
                {"Defect Hist Seq":"1","Defect Code":"CAM","Defect Description":"Camera failed","Repair Status":"N","Defect Type":"MainBoard"},
            ]}},
            {"pcba_sn":"OLD","status":"COMPLETE","result":{"shift":"1st Shift","rows":[
                {"Defect Hist Seq":"0","Defect Code":"CAM","Defect Description":"Camera failed","Repair Status":"Y","Defect Type":"Phone_Disassembly"},
                {"Defect Hist Seq":"1","Defect Code":"RF","Defect Description":"Antenna failed","Repair Status":"Y","Defect Type":"MainBoard"},
            ]}},
        ]}
        return r3074,r2114

    def test_detail_sheets_are_grouped_and_collapsed(self):
        r3074,r2114=self._traces()
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'x.xlsx'; gerar_excel_integrado(r3074,r2114,p)
            wb=load_workbook(p)
            self.assertEqual(wb.sheetnames[:4],["INDICADORES","ANALISE","DETALHE_PCBA","DETALHE_MATERIAL"])
            wp=wb["DETALHE_PCBA"]; wm=wb["DETALHE_MATERIAL"]
            self.assertTrue(any(d.hidden and d.outlineLevel>=1 for d in wp.row_dimensions.values()))
            self.assertTrue(any(d.hidden and d.outlineLevel>=1 for d in wm.row_dimensions.values()))
            self.assertTrue(any(d.hidden and d.outlineLevel>=2 for d in wm.row_dimensions.values()))

if __name__ == '__main__': unittest.main()
