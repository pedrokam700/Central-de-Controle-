import tempfile, unittest
from pathlib import Path
from openpyxl import load_workbook
from reports.relatorio_integrado_2114 import gerar_excel_integrado, _build_views

class ExcelV012Tests(unittest.TestCase):
    def sample3074(self):
        defect={'line':'L3','product_model':'CPH','defect_time':'2026-10-06 10:00:00','defect_code':'GN10-YS01','defect_desc':'Speaker: noise','repair_comment':'x','defect_reason_type':'CS','defect_reason_desc':'reason'}
        return {'items':[{'pcba_sn':'CUR','defect_traces':[{'status':'COMPLETE','defect':defect,'trace_3074':{'status':'OK','results':[{
            'item_sn':'MAT','item_type':'Speaker','status':'SEGUNDO_USO_NO_MOMENTO_DA_FALHA','previous_pcbas':['OLD'],'previous_pcba_count':1,
            'current_pcba_confirmed':True,'total_pcba_count_known_now':2,'reuse_count_known_now':1,'active_now_pcbas':['CUR'],'inactive_now_pcbas':['OLD'],
            'all_pcba_history':[{'pcba_sn':'OLD','iss_status_now':'N'},{'pcba_sn':'CUR','iss_status_now':'Y'}]
        }]}}]}]}
    def sample2114(self):
        return {'pcbas':[
            {'pcba_sn':'CUR','status':'COMPLETE','error':None,'result':{'rows':[
                {'Defect Hist Seq':'0','Defect Code':'GN10-YS01','Defect Description':'Speaker: noise','Defect Type':'Phone_Disassembly','Repair Status':'Y'},
                {'Defect Hist Seq':'1','Defect Code':'GN10-YS01','Defect Description':'Speaker: noise','Defect Type':'MainBoard','Repair Status':'N'}], 'no_data':False}},
            {'pcba_sn':'OLD','status':'COMPLETE','error':None,'result':{'rows':[
                {'Defect Hist Seq':'0','Defect Code':'GN10-YS01','Defect Description':'Speaker: noise','Defect Type':'MainBoard','Repair Status':'Y'}], 'no_data':False}},
        ],'contexts':[],'current_contexts':[]}

    def test_three_visible_views_no_panes(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'x.xlsx'; stats=gerar_excel_integrado(self.sample3074(),self.sample2114(),p)
            wb=load_workbook(p)
            self.assertEqual(wb.sheetnames, ['ANALISE','HIST_PCBA','HIST_REUSO'])
            for ws in wb.worksheets:
                self.assertIsNone(ws.freeze_panes)
                self.assertIsNone(ws.sheet_view.pane)
            self.assertEqual(stats['materiais_reuso'],1)
            self.assertEqual(wb['ANALISE']['K2'].value,1)  # repetição adicional da mesma falha
            self.assertIn('Speaker', wb['ANALISE']['Q2'].value)
            self.assertEqual(wb['HIST_REUSO']['X2'].value,'MESMA FALHA')

    def test_reuse_view_cannot_silently_disappear(self):
        v=_build_views(self.sample3074(), self.sample2114())
        self.assertEqual(v['stats']['materiais_reuso'],1)
        self.assertEqual(len(v['hist_reuso_rows']),1)
        self.assertEqual(v['analise_rows'][0]['Materiais de reuso'],1)

if __name__ == '__main__': unittest.main()
