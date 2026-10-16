from collections import Counter, defaultdict
from copy import copy
from pathlib import Path
import json

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.utils import get_column_letter

from reports.relatorio_v015 import (
    _build_views,
    _material_name,
    _unique_histories,
    _set_internal_link,
)

HEADER_FILL = "1F4E78"
HEADER_FONT = "FFFFFF"
SECTION_FILL = "D9EAF7"
KPI_FILL = "EAF2F8"
ALT_FILL = "F5F9FC"
FILL_SAME = "E2F0D9"
FILL_SIMILAR = "FFF2CC"
FILL_DIFF = "DDEBF7"
FILL_NONE = "E7E6E6"
LIGHT_BORDER = Side(style="thin", color="D9E1F2")
FONT_NAME = "Arial"


def _s(v):
    return "" if v is None else str(v)


def _ival(v):
    try:
        return int(v or 0)
    except Exception:
        return 0


def _is_camera(name):
    t = _s(name).lower()
    return "camera" in t or "摄像头" in t


def _is_screen(name):
    t = _s(name).lower()
    return "screen" in t or "lcd" in t or "显示屏" in t


def _unique_latest_pcba(rows):
    out = {}
    for r in rows:
        pcba = _s(r.get("PCBA")).strip()
        if not pcba:
            continue
        if pcba not in out or _s(r.get("Defect Time")) > _s(out[pcba].get("Defect Time")):
            out[pcba] = r
    return list(out.values())


def _material_maps(views):
    all_materials = {}
    for r in views.get("materiais", []):
        sn = _s(r.get("Material SN")).strip()
        if sn:
            all_materials.setdefault(sn, r)
    reused = {}
    for r in views.get("reusos", []):
        sn = _s(r.get("Material SN")).strip()
        if sn:
            reused.setdefault(sn, r)
    summary_by_sn = defaultdict(lambda: {"history":0,"same":0,"similar":0,"different":0,"no_history":0})
    representative_summary = {}
    for r in views.get("material_summary_rows", []):
        sn = _s(r.get("Material SN")).strip()
        if not sn:
            continue
        representative_summary.setdefault(sn, r)
        d = summary_by_sn[sn]
        d["history"] = max(d["history"], _ival(r.get("Falhas 2114 nas desvinculadas")))
        d["same"] = max(d["same"], _ival(r.get("PCBAs com mesma falha")))
        d["similar"] = max(d["similar"], _ival(r.get("PCBAs com mesma família")))
        d["different"] = max(d["different"], _ival(r.get("PCBAs com falha diferente")))
        d["no_history"] = max(d["no_history"], _ival(r.get("PCBAs sem histórico de falha")))
    return all_materials, reused, summary_by_sn, representative_summary


def _build_indicators_v016(views):
    rows = views["analise_rows"]
    pcba_rows = _unique_latest_pcba(rows)

    pcba_first = sum(1 for r in pcba_rows if _ival(r.get("_pcba_use_index")) == 1)
    pcba_second = sum(1 for r in pcba_rows if _ival(r.get("_pcba_use_index")) == 2)
    pcba_multi = sum(1 for r in pcba_rows if _ival(r.get("_pcba_use_index")) >= 3)
    pcba_unknown = sum(1 for r in pcba_rows if _ival(r.get("_pcba_use_index")) == 0)
    pcba_fail2 = sum(1 for r in pcba_rows if _ival(r.get("Falhas registradas 2114")) >= 2)
    pcba_fail3 = sum(1 for r in pcba_rows if _ival(r.get("Falhas registradas 2114")) >= 3)
    pcba_same = sum(1 for r in pcba_rows if _ival(r.get("Mesma falha anterior")) > 0)
    pcba_sim = sum(1 for r in pcba_rows if _ival(r.get("Mesma família anterior")) > 0)
    pcba_reused_one_failure = sum(1 for r in pcba_rows if _ival(r.get("_pcba_use_index")) >= 2 and _ival(r.get("Falhas registradas 2114")) == 1)
    pcba_reused_with_prior_failure = sum(1 for r in pcba_rows if _ival(r.get("_pcba_use_index")) >= 2 and _ival(r.get("Falhas anteriores 2114")) > 0)

    all_materials, reused, summary_by_sn, representative_summary = _material_maps(views)
    by_type = defaultdict(lambda: {
        "all":set(),"reuse":set(),"second":set(),"multi":set(),
        "with_old_failure":set(),"without_old_failure":set(),"same":set(),"similar":set()
    })
    for sn, r in all_materials.items():
        by_type[_material_name(r.get("Tipo material"))]["all"].add(sn)
    for sn, r in reused.items():
        name = _material_name(r.get("Tipo material"))
        d = by_type[name]
        d["reuse"].add(sn)
        prev = _ival(r.get("Reusos anteriores a falha"))
        if prev == 1: d["second"].add(sn)
        elif prev >= 2: d["multi"].add(sn)
        sig = summary_by_sn[sn]
        if sig["history"] > 0: d["with_old_failure"].add(sn)
        else: d["without_old_failure"].add(sn)
        if sig["same"] > 0: d["same"].add(sn)
        elif sig["similar"] > 0: d["similar"].add(sn)

    component_types = []
    for name, d in sorted(by_type.items(), key=lambda kv: (-len(kv[1]["reuse"]), kv[0])):
        total, reuse = len(d["all"]), len(d["reuse"])
        component_types.append({
            "Tipo material": name,
            "Materiais únicos no lote": total,
            "Reutilizados únicos": reuse,
            "2º uso": len(d["second"]),
            "3º+ uso": len(d["multi"]),
            "Com falha nas PCBAs desvinculadas": len(d["with_old_failure"]),
            "Sem falha nas PCBAs desvinculadas": len(d["without_old_failure"]),
            "Com mesma falha": len(d["same"]),
            "Com mesma família": len(d["similar"]),
            "% reuso": round(reuse / total * 100.0, 1) if total else 0,
        })

    mat_second = sum(1 for r in reused.values() if _ival(r.get("Reusos anteriores a falha")) == 1)
    mat_multi = sum(1 for r in reused.values() if _ival(r.get("Reusos anteriores a falha")) >= 2)
    mat_with_old = sum(1 for sn in reused if summary_by_sn[sn]["history"] > 0)
    mat_without_old = len(reused) - mat_with_old
    mat_same = sum(1 for sn in reused if summary_by_sn[sn]["same"] > 0)
    mat_sim = sum(1 for sn in reused if summary_by_sn[sn]["same"] == 0 and summary_by_sn[sn]["similar"] > 0)
    mat_diff_only = sum(1 for sn in reused if summary_by_sn[sn]["history"] > 0 and summary_by_sn[sn]["same"] == 0 and summary_by_sn[sn]["similar"] == 0)
    screen_count = sum(1 for sn, r in reused.items() if _is_screen(_material_name(r.get("Tipo material"))))
    camera_count = sum(1 for sn, r in reused.items() if _is_camera(_material_name(r.get("Tipo material"))) or _is_camera(r.get("Tipo material")))

    link_relation = {}
    for r in views.get("hist_material_rows", []):
        sn, pcba = r.get("Material SN"), r.get("PCBA desvinculada")
        if sn and pcba:
            link_relation[(sn, pcba)] = r.get("Classificação desta PCBA")
    links_same = sum(1 for v in link_relation.values() if v == "MESMA FALHA")
    links_sim = sum(1 for v in link_relation.values() if v == "MESMA FAMÍLIA")

    hist_unique = _unique_histories(views.get("hist_pcba_rows", []))
    types = Counter(_s(r.get("Defect Type 2114")).strip() or "(vazio)" for r in hist_unique)
    families = Counter(r.get("Família 2114") or "OUTRO" for r in hist_unique if r.get("Defect Description 2114") != "SEM HISTÓRICO 2114")
    current_process = Counter((r.get("Situação reparo atual") or "SEM HISTÓRICO 2114") for r in pcba_rows)

    attention_pcba = set()
    for r in views.get("hist_pcba_rows", []):
        if r.get("Posição no histórico") == "ANTERIOR" and r.get("Relação com falha atual") in ("MESMA FALHA","MESMA FAMÍLIA") and r.get("Situação reparo") == "REPARO DE PLACA CONCLUÍDO":
            attention_pcba.add(r.get("PCBA"))

    line_data = defaultdict(lambda:{
        "falhas":0,"pcbas":set(),"pcba2":set(),"pcba3":set(),"pcba_one_failure":set(),
        "mat2":set(),"mat3":set(),"mat_old":set(),"mat_no_old":set(),"same":set(),"family":set()
    })
    for r in rows:
        line = r.get("Linha") or "SEM LINHA"; d=line_data[line]
        d["falhas"] += 1; d["pcbas"].add(r.get("PCBA")); uses=_ival(r.get("_pcba_use_index"))
        if uses==2: d["pcba2"].add(r.get("PCBA"))
        if uses>=3: d["pcba3"].add(r.get("PCBA"))
        if uses>=2 and _ival(r.get("Falhas registradas 2114"))==1: d["pcba_one_failure"].add(r.get("PCBA"))
        if _ival(r.get("Mesma falha anterior"))>0: d["same"].add(r.get("PCBA"))
        if _ival(r.get("Mesma família anterior"))>0: d["family"].add(r.get("PCBA"))
    for sn,r in reused.items():
        line=r.get("Linha") or "SEM LINHA"; d=line_data[line]; prev=_ival(r.get("Reusos anteriores a falha"))
        if prev==1: d["mat2"].add(sn)
        if prev>=2: d["mat3"].add(sn)
        if summary_by_sn[sn]["history"]>0: d["mat_old"].add(sn)
        else: d["mat_no_old"].add(sn)
    lines=[]
    for line,d in sorted(line_data.items()):
        lines.append({
            "Linha":line,"Falhas":d["falhas"],"PCBAs únicas":len(d["pcbas"]),"PCBA 2º uso":len(d["pcba2"]),"PCBA 3º+ uso":len(d["pcba3"]),
            "PCBA 2º+ uso / só 1 falha":len(d["pcba_one_failure"]),"Materiais 2º uso":len(d["mat2"]),"Materiais 3º+ uso":len(d["mat3"]),
            "Materiais com falha anterior":len(d["mat_old"]),"Materiais sem falha anterior":len(d["mat_no_old"]),"PCBAs mesma falha":len(d["same"]),"PCBAs mesma família":len(d["family"]),
        })

    return {
        "overview_kpis":[("Falhas atuais 3028",len(rows)),("PCBAs únicas com falha",len(pcba_rows)),("Linhas no lote",len(line_data)),("Materiais únicos presentes",len(all_materials))],
        "pcba_kpis":[
            ("PCBA em 1º uso",pcba_first),("PCBA em 2º uso",pcba_second),("PCBA em 3º+ uso",pcba_multi),("PCBAs 2º+ uso com só 1 falha 2114",pcba_reused_one_failure),("PCBA uso não determinado",pcba_unknown),
            ("PCBAs 2º+ uso com falha anterior",pcba_reused_with_prior_failure),("PCBAs com 2+ falhas 2114",pcba_fail2),("PCBAs com 3+ falhas 2114",pcba_fail3),("PCBAs que repetiram a mesma falha",pcba_same),
            ("PCBAs com recorrência da mesma família",pcba_sim),("Recorrência após MainBoard + Y",len(attention_pcba)),
        ],
        "material_kpis":[
            ("Materiais únicos 2º+ uso",len(reused)),("Materiais em 2º uso",mat_second),("Materiais em 3º+ uso",mat_multi),("Reuso COM falha em PCBA desvinculada",mat_with_old),
            ("Reuso SEM falha em PCBA desvinculada",mat_without_old),("Reuso com a MESMA falha",mat_same),("Reuso com falha da MESMA família",mat_sim),("Reuso com somente outras falhas",mat_diff_only),
            ("Screens/LCDs reutilizados",screen_count),("Câmeras reutilizadas",camera_count),("Tipos de componente com reuso",sum(1 for x in component_types if x["Reutilizados únicos"]>0)),
        ],
        "correlation_kpis":[("Vínculos material↔PCBA com mesma falha",links_same),("Vínculos material↔PCBA com mesma família",links_sim)],
        "lines":lines,"component_types":component_types,
        "current_repair_states":sorted(current_process.items(), key=lambda x:(-x[1],x[0])),
        "families":sorted(families.items(), key=lambda x:(-x[1],x[0])),
        "defect_types":sorted(types.items(), key=lambda x:(-x[1],x[0])),
        "pcba_rows":pcba_rows,"all_materials":all_materials,"reused_materials":reused,"summary_by_sn":summary_by_sn,
        "representative_material":representative_summary,"link_relation":link_relation,"attention_pcba":attention_pcba,
    }


def _font(cell, bold=False, color="000000", size=9):
    cell.font = Font(name=FONT_NAME, size=size, bold=bold, color=color)


def _style_header(ws, row, c1, c2):
    for col in range(c1,c2+1):
        c=ws.cell(row,col); c.fill=PatternFill("solid",fgColor=HEADER_FILL); _font(c,True,HEADER_FONT,9)
        c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); c.border=Border(bottom=LIGHT_BORDER)
    ws.row_dimensions[row].height=30


def _style_body(ws, min_row=2, max_row=None, max_col=None):
    max_row=max_row or ws.max_row; max_col=max_col or ws.max_column
    for r in range(min_row,max_row+1):
        if r % 2 == 0:
            fill=PatternFill("solid",fgColor=ALT_FILL)
        else:
            fill=None
        for col in range(1,max_col+1):
            c=ws.cell(r,col); _font(c,False,"000000",9); c.alignment=Alignment(vertical="top",wrap_text=True); c.border=Border(bottom=LIGHT_BORDER)
            if fill: c.fill=fill


def _add_table(ws, ref, name):
    tab=Table(displayName=name,ref=ref)
    tab.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showFirstColumn=False,showLastColumn=False,showRowStripes=True,showColumnStripes=False)
    ws.add_table(tab)
    ws.auto_filter.ref=ref
    return tab


def _write_flat_sheet(wb,name,rows,headers,table_name,widths=None,style_body=True):
    ws=wb.create_sheet(name); ws.sheet_view.showGridLines=False; ws.sheet_view.zoomScale=85; ws.freeze_panes=None; ws.sheet_view.pane=None
    ws.append(headers)
    for r in rows: ws.append([r.get(h,"") for h in headers])
    _style_header(ws,1,1,len(headers))
    if style_body:
        _style_body(ws,2,ws.max_row,len(headers))
    if rows:
        _add_table(ws,f"A1:{get_column_letter(len(headers))}{len(rows)+1}",table_name)
    for i,h in enumerate(headers,1):
        w=(widths or {}).get(h)
        if w is None:
            if h.startswith("Reuso "): w=28
            elif "Description" in h or "Desc" in h or "Comment" in h or "Motivo" in h or "Situação" in h or "Ciclos" in h: w=30
            elif "PCBA" in h or "Material SN" in h: w=27
            else: w=16
        ws.column_dimensions[get_column_letter(i)].width=w
    return ws


def _color_relation_cells(ws, headers, names):
    hmap={h:i+1 for i,h in enumerate(headers)}
    for name in names:
        col=hmap.get(name)
        if not col: continue
        for r in range(2,ws.max_row+1):
            c=ws.cell(r,col); v=_s(c.value).upper()
            if "MESMA FALHA" in v or "RECORRÊNCIA DA MESMA FALHA" in v: color=FILL_SAME
            elif "MESMA FAMÍLIA" in v: color=FILL_SIMILAR
            elif "DIFERENTE" in v or "OUTRAS" in v: color=FILL_DIFF
            elif "SEM HIST" in v or "SEM REPET" in v: color=FILL_NONE
            else: continue
            c.fill=PatternFill("solid",fgColor=color); _font(c,True,"000000",9)


def _flat_hist_pcba(views):
    rows=[]; seen_pcba=set(); seen_hist=set()
    by_key_analysis={r.get("_failure_key"):r for r in views.get("analise_rows",[])}
    for r in views.get("hist_pcba_rows",[]):
        x={k:v for k,v in r.items() if not str(k).startswith("_")}
        a=by_key_analysis.get(r.get("_failure_key"),{})
        x["Mesma falha anterior"]=a.get("Mesma falha anterior",0); x["Mesma família anterior"]=a.get("Mesma família anterior",0); x["Falhas diferentes anteriores"]=a.get("Falhas diferentes anteriores",0)
        pkey=r.get("_failure_key")
        x["Peso PCBA atual"]=0 if pkey in seen_pcba else 1; seen_pcba.add(pkey)
        hs=(r.get("PCBA"),r.get("Hist Seq"),r.get("Defect Code 2114"),r.get("Defect Description 2114"),r.get("Repair Status 2114"),r.get("Defect Type 2114"))
        has_real = r.get("Defect Description 2114") != "SEM HISTÓRICO 2114" and bool(r.get("Defect Code 2114") or r.get("Defect Description 2114"))
        x["Peso falha 2114"]=1 if has_real and hs not in seen_hist else 0
        if has_real: seen_hist.add(hs)
        x["Voltar"]="← ANALISE"
        rows.append(x)
    return rows


def _flat_hist_material(views):
    rows=[]; seen_mat=set(); seen_link=set(); seen_fail=set(); seen_cls=set()
    for r in views.get("hist_material_rows",[]):
        x={k:v for k,v in r.items() if not str(k).startswith("_")}; x["Voltar"]="← ANALISE"
        sn=_s(r.get("Material SN")); pcba=_s(r.get("PCBA desvinculada")); link=(sn,pcba)
        x["Peso material reutilizado"]=0 if sn in seen_mat else 1; seen_mat.add(sn)
        x["Peso PCBA desvinculada"]=0 if link in seen_link else 1; seen_link.add(link)
        fail=(sn,pcba,r.get("Hist Seq"),r.get("Defect Code antigo"),r.get("Defect Description antigo"),r.get("Repair Status antigo"),r.get("Defect Type antigo"))
        real=bool(r.get("Defect Code antigo") or (r.get("Defect Description antigo") and r.get("Defect Description antigo")!="SEM HISTÓRICO 2114"))
        x["Peso falha histórica material"]=1 if real and fail not in seen_fail else 0
        if real: seen_fail.add(fail)
        cls=r.get("Classificação desta PCBA")
        ckey=(link,cls)
        first_cls=ckey not in seen_cls
        if first_cls: seen_cls.add(ckey)
        x["Peso PCBA mesma falha"]=1 if first_cls and cls=="MESMA FALHA" else 0
        x["Peso PCBA mesma família"]=1 if first_cls and cls=="MESMA FAMÍLIA" else 0
        x["Peso PCBA falha diferente"]=1 if first_cls and cls=="FALHA DIFERENTE" else 0
        rows.append(x)
    return rows


def _build_detail_indicator_rows(views, ind):
    out=[]; order=[]
    def add(label, rows):
        if label not in order: order.append(label)
        for r in rows:
            item={
                "Abrir":"Abrir ANALISE ↗","Indicador":label,"Linha":r.get("Linha",r.get("line","")),"Modelo":r.get("Modelo",r.get("product_model","")),
                "PCBA atual":r.get("PCBA",r.get("PCBA atual",r.get("PCBA da falha",""))),"Defect Time atual":r.get("Defect Time",r.get("Defect Time atual","")),
                "Defect Code atual":r.get("Defect Code",r.get("Defect Code atual","")),"Defect Desc atual":r.get("Defect Desc",r.get("Defect Desc atual","")),
                "Uso PCBA":r.get("Uso da PCBA na falha",r.get("Uso PCBA na falha","")),"Falhas 2114":r.get("Falhas registradas 2114",0),
                "Mesma falha anterior":r.get("Mesma falha anterior",0),"Mesma família anterior":r.get("Mesma família anterior",0),"Situação PCBA":r.get("Situação da PCBA",""),
                "Tipo material":r.get("Tipo material",""),"Material SN":r.get("Material SN",""),"Uso material":r.get("Uso material na falha",""),
                "PCBA desvinculada":r.get("PCBA desvinculada",""),"Falhas PCBA desvinculada":r.get("Falhas 2114 desta PCBA",r.get("Falhas PCBA desvinculada",0)),
                "Classificação":r.get("Classificação desta PCBA",r.get("Evidência histórica","")),"Observação":r.get("Observação","")
            }
            item["_analysis_row"]=r.get("_analysis_row",2); out.append(item)

    rows=views.get("analise_rows",[]); pcba=ind["pcba_rows"]
    add("Falhas atuais 3028", rows); add("PCBAs únicas com falha", pcba)
    add("PCBA em 1º uso",[r for r in pcba if _ival(r.get("_pcba_use_index"))==1])
    add("PCBA em 2º uso",[r for r in pcba if _ival(r.get("_pcba_use_index"))==2])
    add("PCBA em 3º+ uso",[r for r in pcba if _ival(r.get("_pcba_use_index"))>=3])
    add("PCBA uso não determinado",[r for r in pcba if _ival(r.get("_pcba_use_index"))==0])
    add("PCBAs 2º+ uso com só 1 falha 2114",[r for r in pcba if _ival(r.get("_pcba_use_index"))>=2 and _ival(r.get("Falhas registradas 2114"))==1])
    add("PCBAs 2º+ uso com falha anterior",[r for r in pcba if _ival(r.get("_pcba_use_index"))>=2 and _ival(r.get("Falhas anteriores 2114"))>0])
    add("PCBAs com 2+ falhas 2114",[r for r in pcba if _ival(r.get("Falhas registradas 2114"))>=2])
    add("PCBAs com 3+ falhas 2114",[r for r in pcba if _ival(r.get("Falhas registradas 2114"))>=3])
    add("PCBAs que repetiram a mesma falha",[r for r in pcba if _ival(r.get("Mesma falha anterior"))>0])
    add("PCBAs com recorrência da mesma família",[r for r in pcba if _ival(r.get("Mesma família anterior"))>0])
    add("Recorrência após MainBoard + Y",[r for r in pcba if r.get("PCBA") in ind["attention_pcba"]])

    rep=ind["representative_material"]
    reused=ind["reused_materials"]; sig=ind["summary_by_sn"]
    def mats(pred):
        data=[]
        for sn,r0 in reused.items():
            if not pred(sn,r0): continue
            r=copy(rep.get(sn) or r0); r["Material SN"]=sn; r["Tipo material"]=_material_name(r0.get("Tipo material")); r["_analysis_row"]=(rep.get(sn) or {}).get("_analysis_row",2); data.append(r)
        return data
    add("Materiais únicos 2º+ uso",mats(lambda sn,r:True))
    add("Materiais em 2º uso",mats(lambda sn,r:_ival(r.get("Reusos anteriores a falha"))==1))
    add("Materiais em 3º+ uso",mats(lambda sn,r:_ival(r.get("Reusos anteriores a falha"))>=2))
    add("Reuso COM falha em PCBA desvinculada",mats(lambda sn,r:sig[sn]["history"]>0))
    add("Reuso SEM falha em PCBA desvinculada",mats(lambda sn,r:sig[sn]["history"]==0))
    add("Reuso com a MESMA falha",mats(lambda sn,r:sig[sn]["same"]>0))
    add("Reuso com falha da MESMA família",mats(lambda sn,r:sig[sn]["same"]==0 and sig[sn]["similar"]>0))
    add("Reuso com somente outras falhas",mats(lambda sn,r:sig[sn]["history"]>0 and sig[sn]["same"]==0 and sig[sn]["similar"]==0))
    add("Screens/LCDs reutilizados",mats(lambda sn,r:_is_screen(_material_name(r.get("Tipo material")))))
    add("Câmeras reutilizadas",mats(lambda sn,r:_is_camera(_material_name(r.get("Tipo material"))) or _is_camera(r.get("Tipo material"))))
    type_rows=[]
    for x in ind["component_types"]:
        if x["Reutilizados únicos"]>0:
            type_rows.append({"Tipo material":x["Tipo material"],"Observação":f"{x['Reutilizados únicos']} Material SN reutilizado(s) de {x['Materiais únicos no lote']} no lote"})
    add("Tipos de componente com reuso",type_rows)

    # Um vínculo por Material SN + PCBA desvinculada.
    link_first={}
    for r in views.get("hist_material_rows",[]):
        key=(r.get("Material SN"),r.get("PCBA desvinculada"))
        link_first.setdefault(key,r)
    add("Vínculos material↔PCBA com mesma falha",[r for k,r in link_first.items() if r.get("Classificação desta PCBA")=="MESMA FALHA"])
    add("Vínculos material↔PCBA com mesma família",[r for k,r in link_first.items() if r.get("Classificação desta PCBA")=="MESMA FAMÍLIA"])
    return out, order


def _write_details_indicators(wb, rows, order):
    headers=["Abrir","Indicador","Linha","Modelo","PCBA atual","Defect Time atual","Defect Code atual","Defect Desc atual","Uso PCBA","Falhas 2114","Mesma falha anterior","Mesma família anterior","Situação PCBA","Tipo material","Material SN","Uso material","PCBA desvinculada","Falhas PCBA desvinculada","Classificação","Observação"]
    # Ordenar pela mesma sequência visual dos cards.
    rank={label:i for i,label in enumerate(order)}
    rows=sorted(rows,key=lambda r:(rank.get(r.get("Indicador"),9999),_s(r.get("Linha")),_s(r.get("PCBA atual")),_s(r.get("Material SN"))))
    ws=_write_flat_sheet(wb,"DETALHES_INDICADORES",rows,headers,"DetalhesIndicadores",widths={"Abrir":18,"Indicador":34,"PCBA atual":27,"Defect Desc atual":34,"Material SN":30,"Observação":35})
    anchors={}; current=None
    for i,r in enumerate(rows,2):
        label=r.get("Indicador")
        if label not in anchors: anchors[label]=i
        ar=_ival(r.get("_analysis_row")) or 2
        _set_internal_link(ws.cell(i,1),"ANALISE",f"A{ar}","Abrir ANALISE ↗")
    _color_relation_cells(ws,headers,["Situação PCBA","Classificação"])
    return ws, anchors


def _write_kpi_cards(ws,start_row,title,pairs,anchors,per_row=4):
    max_col=per_row*2
    ws.merge_cells(start_row=start_row,start_column=1,end_row=start_row,end_column=max_col)
    c=ws.cell(start_row,1); c.value=title; c.fill=PatternFill("solid",fgColor=HEADER_FILL); _font(c,True,HEADER_FONT,12); c.alignment=Alignment(vertical="center")
    row=start_row+1
    for i,(label,value) in enumerate(pairs):
        g=i%per_row
        if i and g==0: row+=3
        c1=1+g*2; c2=c1+1
        ws.merge_cells(start_row=row,start_column=c1,end_row=row,end_column=c2); ws.merge_cells(start_row=row+1,start_column=c1,end_row=row+1,end_column=c2)
        l=ws.cell(row,c1); v=ws.cell(row+1,c1); l.value=label; v.value=value
        l.fill=PatternFill("solid",fgColor=KPI_FILL); _font(l,True,"44546A",9); l.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
        v.fill=PatternFill("solid",fgColor="FFFFFF"); _font(v,True,HEADER_FILL,16); v.alignment=Alignment(horizontal="center",vertical="center")
        for rr in (row,row+1):
            for cc in range(c1,c2+1): ws.cell(rr,cc).border=Border(left=LIGHT_BORDER,right=LIGHT_BORDER,top=LIGHT_BORDER,bottom=LIGHT_BORDER)
        if _ival(value)>0 and label in anchors: _set_internal_link(v,"DETALHES_INDICADORES",f"A{anchors[label]}",str(value))
    groups=(len(pairs)+per_row-1)//per_row
    return start_row+1+groups*3


def _write_indicators(wb,ind,anchors):
    ws=wb.create_sheet("INDICADORES"); ws.sheet_view.showGridLines=False; ws.sheet_view.zoomScale=90
    ws.merge_cells("A1:J1"); ws["A1"]="A-MES — Indicadores de rastreabilidade, reuso e recorrência"; ws["A1"].fill=PatternFill("solid",fgColor=HEADER_FILL); _font(ws["A1"],True,HEADER_FONT,15); ws["A1"].alignment=Alignment(vertical="center"); ws.row_dimensions[1].height=28
    ws.merge_cells("A2:J2"); ws["A2"]="Clique nos valores azuis dos cards para abrir o recorte exato. Os filtros das tabelas estão ativos."; _font(ws["A2"],False,"666666",9); ws["A2"].alignment=Alignment(wrap_text=True)
    r=4
    r=_write_kpi_cards(ws,r,"VISÃO GERAL DO LOTE",ind["overview_kpis"],anchors,4)+1
    r=_write_kpi_cards(ws,r,"PCBA — USO x FALHAS 2114",ind["pcba_kpis"],anchors,5)+1
    r=_write_kpi_cards(ws,r,"MATERIAIS / COMPONENTES — REUSO E EVIDÊNCIA",ind["material_kpis"],anchors,4)+1
    r=_write_kpi_cards(ws,r,"VÍNCULOS MATERIAL ↔ PCBA DESVINCULADA",ind["correlation_kpis"],anchors,4)+2

    ws.cell(r,1).value="POR LINHA"; ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=12); ws.cell(r,1).fill=PatternFill("solid",fgColor=SECTION_FILL); _font(ws.cell(r,1),True,"000000",11); r+=1
    line_headers=["Linha","Falhas","PCBAs únicas","PCBA 2º uso","PCBA 3º+ uso","PCBA 2º+ uso / só 1 falha","Materiais 2º uso","Materiais 3º+ uso","Materiais com falha anterior","Materiais sem falha anterior","PCBAs mesma falha","PCBAs mesma família"]
    line_header=r
    for c,h in enumerate(line_headers,1): ws.cell(r,c).value=h
    _style_header(ws,r,1,len(line_headers)); r+=1; line_start=r
    for d in ind["lines"]:
        for c,h in enumerate(line_headers,1): ws.cell(r,c).value=d.get(h,"")
        r+=1
    line_end=r-1; _style_body(ws,line_start,line_end,len(line_headers)); r+=2

    ws.cell(r,1).value="TODOS OS COMPONENTES DO LOTE — MATERIAL SN ÚNICO"; ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=10); ws.cell(r,1).fill=PatternFill("solid",fgColor=SECTION_FILL); _font(ws.cell(r,1),True,"000000",11); r+=1
    comp_headers=["Tipo material","Materiais únicos no lote","Reutilizados únicos","2º uso","3º+ uso","Com falha nas PCBAs desvinculadas","Sem falha nas PCBAs desvinculadas","Com mesma falha","Com mesma família","% reuso"]
    comp_header=r
    for c,h in enumerate(comp_headers,1): ws.cell(r,c).value=h
    _style_header(ws,r,1,len(comp_headers)); r+=1; comp_start=r
    for d in ind["component_types"]:
        for c,h in enumerate(comp_headers,1): ws.cell(r,c).value=d.get(h,"")
        ws.cell(r,10).number_format='0.0"%"'; r+=1
    comp_end=r-1; _style_body(ws,comp_start,comp_end,len(comp_headers)); r+=2

    ws.cell(r,1).value="SITUAÇÃO ATUAL DO REPARO"; ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=4); ws.cell(r,1).fill=PatternFill("solid",fgColor=SECTION_FILL); _font(ws.cell(r,1),True,"000000",11)
    ws.cell(r,6).value="FAMÍLIAS DE FALHA 2114"; ws.merge_cells(start_row=r,start_column=6,end_row=r,end_column=9); ws.cell(r,6).fill=PatternFill("solid",fgColor=SECTION_FILL); _font(ws.cell(r,6),True,"000000",11); r+=1
    ws.cell(r,1).value="Situação"; ws.cell(r,2).value="Qtd"; _style_header(ws,r,1,2); ws.cell(r,6).value="Família"; ws.cell(r,7).value="Qtd"; _style_header(ws,r,6,7)
    n=max(len(ind["current_repair_states"]),len(ind["families"]),1)
    for i in range(n):
        rr=r+1+i
        if i<len(ind["current_repair_states"]): ws.cell(rr,1).value=ind["current_repair_states"][i][0]; ws.cell(rr,2).value=ind["current_repair_states"][i][1]
        if i<len(ind["families"]): ws.cell(rr,6).value=ind["families"][i][0]; ws.cell(rr,7).value=ind["families"][i][1]
    r=r+n+3
    ws.cell(r,1).value="DEFECT TYPE — TODOS OS VALORES PRESERVADOS"; ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=4); ws.cell(r,1).fill=PatternFill("solid",fgColor=SECTION_FILL); _font(ws.cell(r,1),True,"000000",11); r+=1
    ws.cell(r,1).value="Defect Type"; ws.cell(r,2).value="Qtd"; _style_header(ws,r,1,2)
    for i,(name,count) in enumerate(ind["defect_types"],1): ws.cell(r+i,1).value=name; ws.cell(r+i,2).value=count
    r=r+max(len(ind["defect_types"]),1)+3

    if comp_end>=comp_start:
        chart=BarChart(); chart.type="bar"; chart.style=10; chart.title="Reuso por tipo de componente — Material SN único"; chart.x_axis.title="Quantidade"; chart.y_axis.title="Componente"
        data=Reference(ws,min_col=3,min_row=comp_header,max_row=comp_end); cats=Reference(ws,min_col=1,min_row=comp_start,max_row=comp_end); chart.add_data(data,titles_from_data=True); chart.set_categories(cats); chart.width=19; chart.height=max(8,min(16,4+len(ind["component_types"])*0.45)); ws.add_chart(chart,f"A{r}"); r+=24
    if line_end>=line_start:
        chart2=BarChart(); chart2.type="col"; chart2.style=10; chart2.title="Falhas e PCBAs únicas por linha"; chart2.y_axis.title="Quantidade"
        data=Reference(ws,min_col=2,max_col=3,min_row=line_header,max_row=line_end); cats=Reference(ws,min_col=1,min_row=line_start,max_row=line_end); chart2.add_data(data,titles_from_data=True); chart2.set_categories(cats); chart2.width=19; chart2.height=8; ws.add_chart(chart2,f"A{r}")
    for col,w in {"A":35,"B":18,"C":35,"D":18,"E":35,"F":18,"G":35,"H":18,"I":35,"J":14,"K":18,"L":18}.items(): ws.column_dimensions[col].width=w
    for row in ws.iter_rows():
        for c in row:
            if c.font.name != FONT_NAME: c.font=copy(c.font); c.font=Font(name=FONT_NAME,size=c.font.sz or 9,bold=c.font.bold,color=c.font.color)
            c.alignment=copy(c.alignment); c.alignment=Alignment(horizontal=c.alignment.horizontal,vertical=c.alignment.vertical or "center",wrap_text=True)
    return ws


def _raw_3028_map(resultado_3074):
    out={}
    for item in (resultado_3074 or {}).get("items",[]) or []:
        for d in item.get("current_defects",[]) or []:
            key=(_s(d.get("pcba_sn")),_s(d.get("defect_time")),_s(d.get("defect_code")))
            out[key]=d
    return out


def _build_base_rows(views, resultado_3074, resultado_2114):
    base=[]; analysis_by_key={r.get("_failure_key"):r for r in views.get("analise_rows",[])}; raw3028=_raw_3028_map(resultado_3074)
    def common_from_analysis(a):
        if not a: return {}
        return {
            "ID análise":"|".join(map(_s,a.get("_failure_key") or (a.get("PCBA"),a.get("Defect Time"),a.get("Defect Code")))),"Linha":a.get("Linha",""),"Modelo":a.get("Modelo",""),"PCBA atual":a.get("PCBA",""),
            "Defect Time atual":a.get("Defect Time",""),"Defect Code atual":a.get("Defect Code",""),"Defect Desc atual":a.get("Defect Desc",""),"Repair Comment atual":a.get("Repair Comment",""),"Reason Type":a.get("Reason Type",""),"Reason Desc":a.get("Reason Desc",""),
            "Uso PCBA na falha":a.get("Uso da PCBA na falha",""),"Usos conhecidos PCBA":a.get("Usos conhecidos da PCBA",0),"Reusos anteriores PCBA":a.get("Reusos anteriores da PCBA",0),"Falhas registradas 2114":a.get("Falhas registradas 2114",0),
            "Falhas anteriores 2114":a.get("Falhas anteriores 2114",0),"Mesma falha anterior":a.get("Mesma falha anterior",0),"Mesma família anterior":a.get("Mesma família anterior",0),"Falhas diferentes anteriores":a.get("Falhas diferentes anteriores",0),
            "Situação da PCBA":a.get("Situação da PCBA",""),"Repair Status atual":a.get("Repair Status atual",""),"Defect Type atual":a.get("Defect Type atual",""),"Situação reparo atual":a.get("Situação reparo atual",""),"Registro falha atual":a.get("Registro falha atual",""),
        }
    def apply_3028(row,a):
        key=(row.get("PCBA atual",""),row.get("Defect Time atual",""),row.get("Defect Code atual","")); d=raw3028.get(tuple(map(_s,key)),{})
        mp={"defect_oper":"3028 Defect Oper","defect_oper_desc":"3028 Defect Oper Desc","defect_return_category":"3028 Defect Return Category","defect_detail_category":"3028 Defect Detail Category","reason_code":"3028 Reason Code","repair_code":"3028 Repair Code","repair_desc":"3028 Repair Desc","repair_user":"3028 Repair User","defect_material":"3028 Defect Material","defect_material_desc":"3028 Defect Material Desc","defect_category":"3028 Defect Category","defect_source":"3028 Defect Source","color":"3028 Color","imei":"3028 IMEI","time_zone":"3028 Time Zone","work_shift":"3028 Work Shift","order_id":"3028 Order ID","user_id":"3028 User ID","manual_or_auto":"3028 Manual/Auto"}
        for k,h in mp.items(): row[h]=d.get(k,"")
    for a in views.get("analise_rows",[]):
        row={"Tipo registro":"FALHA_ATUAL","Fonte original":"3028","Subtipo registro":"Falha atual",**common_from_analysis(a),"Peso falha atual":1,"Peso PCBA atual":1}; apply_3028(row,a); base.append(row)
    for h in views.get("hist_pcba_rows",[]):
        a=analysis_by_key.get(h.get("_failure_key"),{}); row={"Tipo registro":"HIST_PCBA","Fonte original":"2114","Subtipo registro":"Histórico da PCBA atual",**common_from_analysis(a)}
        row.update({"Hist Seq PCBA":h.get("Hist Seq",""),"Defect Code 2114 PCBA":h.get("Defect Code 2114",""),"Defect Description 2114 PCBA":h.get("Defect Description 2114",""),"Defect Oper 2114 PCBA":h.get("Defect Oper 2114",""),"Registro falha 2114 PCBA":h.get("Registro falha 2114",""),"Repair Status 2114 PCBA":h.get("Repair Status 2114",""),"Defect Type 2114 PCBA":h.get("Defect Type 2114",""),"Situação reparo 2114 PCBA":h.get("Situação reparo",""),"Família 2114 PCBA":h.get("Família 2114",""),"Relação PCBA":h.get("Relação com falha atual",""),"Motivo relação PCBA":h.get("Motivo da relação",""),"Ciclos PCBA 3074":h.get("Ciclos de uso PCBA 3074",""),"Peso falha 2114 PCBA":0 if h.get("Defect Description 2114")=="SEM HISTÓRICO 2114" else 1}); apply_3028(row,a); base.append(row)
    for m in views.get("hist_material_rows",[]):
        a=analysis_by_key.get(m.get("_failure_key"),{}); row={"Tipo registro":"HIST_MATERIAL","Fonte original":"3074+2114","Subtipo registro":"Histórico de PCBA desvinculada de material reutilizado",**common_from_analysis(a)}
        for h in ["Tipo material","Material SN","Uso material na falha","Reusos anteriores material","Usos conhecidos material","Reusos conhecidos material","PCBAs desvinculadas únicas","PCBAs com mesma falha","PCBAs com mesma família","PCBAs com falha diferente","Falhas 2114 nas desvinculadas","Evidência histórica","PCBA desvinculada","Estado vínculo atual","Ordem de uso anterior","Falhas 2114 desta PCBA","Classificação desta PCBA","Hist Seq","Defect Code antigo","Defect Description antigo","Defect Oper antigo","Registro falha antigo","Repair Status antigo","Defect Type antigo","Situação reparo antiga","Família falha antiga","Relação desta ocorrência","Motivo da relação"]: row[h if h!="Hist Seq" else "Hist Seq material"]=m.get(h,"")
        apply_3028(row,a); base.append(row)
    # 3074 bruto: uma linha por registro retornado na grade base.
    for item in (resultado_3074 or {}).get("items",[]) or []:
        pcba=item.get("pcba_sn",""); a=next((x for x in views.get("analise_rows",[]) if x.get("PCBA")==pcba),{})
        for rr in ((item.get("base_query") or {}).get("rows") or []):
            row={"Tipo registro":"3074_RAW","Fonte original":"3074","Subtipo registro":"Linha bruta da grade 3074",**common_from_analysis(a)}
            for k,v in rr.items(): row[f"3074 {k}"]=v
            apply_3028(row,a); base.append(row)
        for tr in item.get("defect_traces",[]) or []:
            trace=tr.get("trace_3074") or {}; defect=tr.get("defect") or {}; a2=analysis_by_key.get((_s(defect.get("pcba_sn") or pcba),_s(defect.get("defect_time")),_s(defect.get("defect_code"))),a)
            for res in trace.get("results",[]) or []:
                row={"Tipo registro":"3074_COMPONENTE","Fonte original":"3074","Subtipo registro":"Resultado temporal por componente",**common_from_analysis(a2),"3074D Defect Key":tr.get("defect_key",""),"3074D Status trace":trace.get("status","")}
                mp={"status":"3074D Material Status reuso","bind_time_utc":"3074D Bind Time UTC","unbind_time_utc":"3074D Unbind Time UTC","iss_status_now":"3074D Iss Status Now","previous_pcba_count":"3074D Previous PCBA Count","previous_pcbas":"3074D Previous PCBAs","pcbas_until_defect":"3074D PCBAs until defect","current_pcba_confirmed":"3074D Current PCBA Confirmed","future_association_count_ignored":"3074D Future Associations Ignored","total_pcba_count_known_now":"3074D Total PCBA Count Known Now","reuse_count_known_now":"3074D Reuse Count Known Now","active_now_pcbas":"3074D Active Now PCBAs","inactive_now_pcbas":"3074D Inactive Now PCBAs","all_pcba_history":"3074D All PCBA History","source":"3074D Source","error":"3074D Error"}
                for k,h in mp.items():
                    v=res.get(k,""); row[h]=json.dumps(v,ensure_ascii=False) if isinstance(v,(list,dict)) else v
                apply_3028(row,a2); base.append(row)
    # 2114 bruto: preserva todos os campos retornados inclusive PCBA histórica.
    for item in (resultado_2114 or {}).get("pcbas",[]) or []:
        res=item.get("result") or {}; pcba=item.get("pcba_sn") or res.get("query_sn"); a=next((x for x in views.get("analise_rows",[]) if x.get("PCBA")==pcba),{})
        common={"2114 Query SN":res.get("query_sn",pcba),"2114 Shift":res.get("shift",""),"2114 Model":res.get("model",""),"2114 Color":res.get("color",""),"2114 Material":res.get("material",""),"2114 MFG Order":res.get("mfg_order","")}
        histrows=res.get("rows") or []
        if not histrows: histrows=[{}]
        for hr in histrows:
            row={"Tipo registro":"2114_RAW","Fonte original":"2114","Subtipo registro":"Linha bruta 2114",**common_from_analysis(a),**common}
            for k,v in hr.items(): row[f"2114 {k}"]=v
            apply_3028(row,a); base.append(row)
    return base


def _base_headers():
    return [
        "Tipo registro","ID análise","Linha","Modelo","PCBA atual","Defect Time atual","Defect Code atual","Defect Desc atual","Repair Comment atual","Reason Type","Reason Desc","Uso PCBA na falha","Usos conhecidos PCBA","Reusos anteriores PCBA","Falhas registradas 2114","Falhas anteriores 2114","Mesma falha anterior","Mesma família anterior","Falhas diferentes anteriores","Situação da PCBA","Repair Status atual","Defect Type atual","Situação reparo atual","Registro falha atual",
        "Hist Seq PCBA","Defect Code 2114 PCBA","Defect Description 2114 PCBA","Defect Oper 2114 PCBA","Registro falha 2114 PCBA","Repair Status 2114 PCBA","Defect Type 2114 PCBA","Situação reparo 2114 PCBA","Família 2114 PCBA","Relação PCBA","Motivo relação PCBA","Ciclos PCBA 3074",
        "Tipo material","Material SN","Uso material na falha","Reusos anteriores material","Usos conhecidos material","Reusos conhecidos material","PCBAs desvinculadas únicas","PCBAs com mesma falha","PCBAs com mesma família","PCBAs com falha diferente","Falhas 2114 nas desvinculadas","Evidência histórica material","PCBA desvinculada","Estado vínculo atual","Ordem de uso anterior","Falhas 2114 desta PCBA","Classificação desta PCBA","Hist Seq material","Defect Code antigo","Defect Description antigo","Defect Oper antigo","Registro falha antigo","Repair Status antigo","Defect Type antigo","Situação reparo antiga","Família falha antiga","Relação ocorrência material","Motivo relação material",
        "Peso falha atual","Peso PCBA atual","Peso falha 2114 PCBA","Peso material reutilizado","Peso vínculo material-PCBA","Peso falha histórica material","Peso vínculo mesma falha","Peso vínculo mesma família","Peso vínculo falha diferente","Fonte original","Subtipo registro",
        "3028 Defect Oper","3028 Defect Oper Desc","3028 Defect Return Category","3028 Defect Detail Category","3028 Reason Code","3028 Repair Code","3028 Repair Desc","3028 Repair User","3028 Defect Material","3028 Defect Material Desc","3028 Defect Category","3028 Defect Source","3028 Color","3028 IMEI","3028 Time Zone","3028 Work Shift","3028 Order ID","3028 User ID","3028 Manual/Auto",
        "3074 Barcode","3074 Sn Type","3074 Batch Id","3074 Sn Seq","3074 Batch Count","3074 Mat Gear","3074 Order Id","3074 Line Id","3074 Line Id__2","3074 Operation","3074 Product Model","3074 Product Id","3074 Whether to use","3074 Prod Time","3074 Sn Name","3074 Material Code","3074 Material SN","3074 Bind Time","3074 Unbind Time","3074 Init Flag","3074 Iss Status","3074 Lims Hold Status","3074 Create User Id","3074 Create Time","3074 Aut Cmf1","3074 Aut Cmf2","3074 Aut Cmf3","3074 Aut Cmf4","3074 Aut Cmf5","3074 Aut Cmf6","3074 Aut Cmf7","3074 Aut Cmf8","3074 Aut Cmf9","3074 Aut Cmf10",
        "2114 Query SN","2114 Shift","2114 Model","2114 Color","2114 Material","2114 MFG Order","2114 &#160;","2114 Lot Id","2114 Defect Hist Seq","2114 Defect Location","2114 Defect Material ID","2114 Defect Code","2114 Defect Description","2114 Defect Oper","2114 Repair Status","2114 TestTools Auto Defect Or Defect By Hand","2114 Defect Type","2114 Cancel","2114 Defect Type Class","2114 Defect Type Meaning PT",
        "3074D Defect Key","3074D Status trace","3074D Material Status reuso","3074D Bind Time UTC","3074D Unbind Time UTC","3074D Iss Status Now","3074D Previous PCBA Count","3074D Previous PCBAs","3074D PCBAs until defect","3074D Current PCBA Confirmed","3074D Future Associations Ignored","3074D Total PCBA Count Known Now","3074D Reuse Count Known Now","3074D Active Now PCBAs","3074D Inactive Now PCBAs","3074D All PCBA History","3074D Source","3074D Error"
    ]


def gerar_excel_integrado(resultado_3074, resultado_2114, output_path):
    views=_build_views(resultado_3074,resultado_2114); ind=_build_indicators_v016(views)
    wb=Workbook(); wb.remove(wb.active)

    detail_rows,detail_order=_build_detail_indicator_rows(views,ind)
    ws_details,anchors=_write_details_indicators(wb,detail_rows,detail_order)
    _write_indicators(wb,ind,anchors)

    headers_a=["Linha","Modelo","PCBA","Defect Time","Defect Code","Defect Desc","Repair Comment","Reason Type","Reason Desc","Uso da PCBA na falha","Usos conhecidos da PCBA","Reusos anteriores da PCBA","Falhas registradas 2114","Falhas anteriores 2114","Mesma falha anterior","Mesma família anterior","Falhas diferentes anteriores","Situação da PCBA","Repair Status atual","Defect Type atual","Situação reparo atual","Registro falha atual","Materiais na falha","Materiais 2º+ uso"]+[f"Reuso {i}" for i in range(1,views.get("max_reusos",0)+1)]+["Abrir histórico PCBA","Abrir materiais reutilizados"]
    analise_rows=[]
    for r in views["analise_rows"]:
        x={k:v for k,v in r.items() if not str(k).startswith("_")}; x["Abrir histórico PCBA"]="Histórico PCBA ↗"; x["Abrir materiais reutilizados"]="Materiais ↗"; x["_failure_key"]=r.get("_failure_key"); analise_rows.append(x)
    ws_a=_write_flat_sheet(wb,"ANALISE",analise_rows,headers_a,"T_ANALISE_V016",widths={"Linha":13,"Modelo":15,"PCBA":27,"Defect Time":20,"Defect Desc":34,"Repair Comment":30,"Situação da PCBA":30,"Situação reparo atual":34,"Abrir histórico PCBA":20,"Abrir materiais reutilizados":23})
    _color_relation_cells(ws_a,headers_a,["Situação da PCBA"])

    hp=_flat_hist_pcba(views)
    headers_hp=["Voltar","Linha","Modelo","PCBA","Uso PCBA na falha","Usos conhecidos PCBA","Reusos anteriores PCBA","Falhas registradas 2114","Falhas anteriores 2114","Mesma falha anterior","Mesma família anterior","Falhas diferentes anteriores","Defect Time atual","Defect Code atual","Defect Desc atual","Repair Comment atual","Repair Status atual","Defect Type atual","Situação reparo atual","Hist Seq","Defect Code 2114","Defect Description 2114","Defect Oper 2114","Registro falha 2114","Repair Status 2114","Defect Type 2114","Situação reparo","Família 2114","Relação com falha atual","Motivo da relação","Ciclos de uso PCBA 3074","Peso PCBA atual","Peso falha 2114"]
    ws_hp=_write_flat_sheet(wb,"HIST_PCBA",hp,headers_hp,"T_HIST_PCBA_V016",widths={"Voltar":12,"PCBA":27,"Defect Desc atual":34,"Defect Description 2114":34,"Motivo da relação":38,"Ciclos de uso PCBA 3074":42})
    _color_relation_cells(ws_hp,headers_hp,["Relação com falha atual"])

    hm=_flat_hist_material(views)
    headers_hm=["Voltar","Linha","Modelo","PCBA atual","Defect Time atual","Defect Code atual","Defect Desc atual","Tipo material","Material SN","Uso material na falha","Reusos anteriores material","Usos conhecidos material","Reusos conhecidos material","PCBAs desvinculadas únicas","PCBAs com mesma falha","PCBAs com mesma família","PCBAs com falha diferente","Falhas 2114 nas desvinculadas","Evidência histórica","PCBA desvinculada","Estado vínculo atual","Ordem de uso anterior","Falhas 2114 desta PCBA","Classificação desta PCBA","Hist Seq","Defect Code antigo","Defect Description antigo","Defect Oper antigo","Registro falha antigo","Repair Status antigo","Defect Type antigo","Situação reparo antiga","Família falha antiga","Relação desta ocorrência","Motivo da relação","Peso material reutilizado","Peso PCBA desvinculada","Peso falha histórica material","Peso PCBA mesma falha","Peso PCBA mesma família","Peso PCBA falha diferente"]
    ws_hm=_write_flat_sheet(wb,"HIST_MATERIAL",hm,headers_hm,"T_HIST_MATERIAL_V016",widths={"Voltar":12,"PCBA atual":27,"Defect Desc atual":34,"Tipo material":22,"Material SN":30,"PCBA desvinculada":27,"Defect Description antigo":34,"Motivo da relação":38})
    _color_relation_cells(ws_hm,headers_hm,["Evidência histórica","Classificação desta PCBA","Relação desta ocorrência"])

    # Links e âncoras das tabelas planas.
    pcba_anchor={}; mat_anchor={}; failure_mat_anchor={}
    for i,r in enumerate(views["hist_pcba_rows"],2): pcba_anchor.setdefault(r.get("_failure_key"),i)
    for i,r in enumerate(views["hist_material_rows"],2):
        mat_anchor.setdefault(r.get("_material_key"),i); failure_mat_anchor.setdefault(r.get("_failure_key"),i)
    a_hmap={h:i+1 for i,h in enumerate(headers_a)}
    for i,r in enumerate(views["analise_rows"],2):
        key=r.get("_failure_key")
        if key in pcba_anchor: _set_internal_link(ws_a.cell(i,a_hmap["Abrir histórico PCBA"]),"HIST_PCBA",f"A{pcba_anchor[key]}","Histórico PCBA ↗")
        if key in failure_mat_anchor: _set_internal_link(ws_a.cell(i,a_hmap["Abrir materiais reutilizados"]),"HIST_MATERIAL",f"A{failure_mat_anchor[key]}","Materiais ↗")
        for n in range(1,views.get("max_reusos",0)+1):
            mk=r.get(f"_reuso_{n}_key")
            if mk in mat_anchor and f"Reuso {n}" in a_hmap: _set_internal_link(ws_a.cell(i,a_hmap[f"Reuso {n}"]),"HIST_MATERIAL",f"A{mat_anchor[mk]}")
    hp_hmap={h:i+1 for i,h in enumerate(headers_hp)}
    for i,r in enumerate(views["hist_pcba_rows"],2): _set_internal_link(ws_hp.cell(i,hp_hmap["Voltar"]),"ANALISE",f"A{r.get('_analysis_row',2)}","← ANALISE")
    hm_hmap={h:i+1 for i,h in enumerate(headers_hm)}
    for i,r in enumerate(views["hist_material_rows"],2): _set_internal_link(ws_hm.cell(i,hm_hmap["Voltar"]),"ANALISE",f"A{r.get('_analysis_row',2)}","← ANALISE")

    base_rows=_build_base_rows(views,resultado_3074,resultado_2114); base_headers=_base_headers()
    ws_base=_write_flat_sheet(wb,"BASE_DADOS",base_rows,base_headers,"BaseDadosCompleta",widths={"Tipo registro":18,"ID análise":38,"PCBA atual":27,"Defect Desc atual":34,"Material SN":30,"PCBA desvinculada":27,"3074D All PCBA History":45},style_body=False)

    # Reordenar abas no padrão validado.
    order=["INDICADORES","ANALISE","HIST_PCBA","HIST_MATERIAL","BASE_DADOS","DETALHES_INDICADORES"]
    wb._sheets=[wb[n] for n in order]
    wb.active=0

    # Regras de consistência obrigatórias.
    unique_reuse=len({_s(r.get("Material SN")) for r in views.get("reusos",[]) if r.get("Material SN")})
    if unique_reuse != sum(_ival(r.get("Peso material reutilizado")) for r in hm): raise RuntimeError("HIST_MATERIAL: peso de Material SN reutilizado não bate com DISTINCT Material SN.")
    camera_count=next((v for k,v in ind["material_kpis"] if k=="Câmeras reutilizadas"),0)
    if camera_count>unique_reuse: raise RuntimeError("Indicador impossível: câmeras reutilizadas > materiais reutilizados únicos.")
    if len(views["analise_rows"]) != len((resultado_3074 or {}).get("items",[]) or []) and (resultado_3074 or {}).get("records_count") == (resultado_3074 or {}).get("unique_pcba_count"):
        # Só alerta implicitamente via stats; não bloquear lotes com múltiplas falhas por PCBA.
        pass

    output_path=Path(output_path); output_path.parent.mkdir(parents=True,exist_ok=True); wb.save(output_path)
    return {
        "falhas":len(views["analise_rows"]),"pcbas_unicas":len({_s(r.get('PCBA')) for r in views['analise_rows'] if r.get('PCBA')}),"materiais_reuso":unique_reuse,
        "linhas_hist_pcba":len(hp),"linhas_hist_material":len(hm),"linhas_base_dados":len(base_rows),"detalhes_indicadores":len(detail_rows),
        "pcba_primeiro_uso":next((v for k,v in ind['pcba_kpis'] if k=='PCBA em 1º uso'),0),"pcba_segundo_uso":next((v for k,v in ind['pcba_kpis'] if k=='PCBA em 2º uso'),0),
        "pcba_multiuso":next((v for k,v in ind['pcba_kpis'] if k=='PCBA em 3º+ uso'),0),"pcba_uso_nao_determinado":next((v for k,v in ind['pcba_kpis'] if k=='PCBA uso não determinado'),0),
        "cameras_reutilizadas":camera_count,
    }
