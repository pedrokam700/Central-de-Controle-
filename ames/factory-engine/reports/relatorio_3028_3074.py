from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


HEADER_FILL = "1F4E78"
HEADER_FONT = "FFFFFF"
REUSE_FILL = "FFF2CC"
MULTI_FILL = "FCE4D6"
LIGHT_BORDER = Side(style="thin", color="D9E1F2")


def _join(values):
    return " | ".join(str(x) for x in (values or []) if x not in (None, ""))


def _history_text(r):
    hist = r.get("all_pcba_history") or []
    if hist:
        return " -> ".join(
            f"{x.get('pcba_sn','')} [{x.get('iss_status_now') or '-'}]"
            for x in hist
            if x.get("pcba_sn")
        )
    # Compatibilidade com JSON V0.8 já concluído.
    pcs = list(r.get("previous_pcbas") or [])
    current = None
    until = r.get("pcbas_until_defect") or []
    if until:
        current = until[-1]
    if current and current not in pcs:
        pcs.append(current)
    return " -> ".join(pcs)


def _known_uses(r):
    value = r.get("total_pcba_count_known_now")
    if value is not None:
        return value
    # JSON V0.8 antigo só permite afirmar o total ATÉ A FALHA.
    prev = r.get("previous_pcba_count")
    if prev is None:
        return None
    return int(prev) + (1 if r.get("current_pcba_confirmed", True) else 0)


def _known_reuses(r):
    value = r.get("reuse_count_known_now")
    if value is not None:
        return value
    uses = _known_uses(r)
    return max(0, uses - 1) if uses is not None else None


def _component_card(row):
    """Texto compacto de UM material reutilizado para a view Falhas_3028.

    Cada componente fica em sua propria coluna. As PCBAs antigas permanecem
    junto do componente correspondente para nao perder a associacao visual.
    """
    tipo = row.get("Tipo material") or "Material"
    sn = row.get("Material SN") or "-"
    uses = row.get("Usos totais conhecidos")
    if uses is None:
        uses = row.get("Usos ate a falha")
    reuses = row.get("Reusos totais conhecidos")
    if reuses is None:
        reuses = row.get("Reusos anteriores a falha")

    antigas = [x.strip() for x in str(row.get("PCBAs antigas") or "").split("|") if x.strip()]
    atuais = [x.strip() for x in str(row.get("PCBAs atuais Y") or "").split("|") if x.strip()]

    lines = [
        str(tipo),
        f"SN: {sn}",
        f"Usos: {uses if uses is not None else '-'} | Reusos: {reuses if reuses is not None else '-'}",
    ]
    if antigas:
        lines.append("PCBAs antigas:")
        lines.extend(antigas)
    else:
        lines.append("PCBAs antigas: -")

    if atuais:
        lines.append("Y agora: " + " | ".join(atuais))
    elif row.get("Iss Status agora"):
        lines.append("Iss agora: " + str(row.get("Iss Status agora")))

    return "\n".join(lines)


def extrair_linhas(resultado):
    falhas = []
    materiais = []
    reuso = []

    for item in resultado.get("items", []) or []:
        pcba = item.get("pcba_sn", "")
        for dstate in item.get("defect_traces", []) or []:
            defect = dstate.get("defect", {}) or {}
            trace = dstate.get("trace_3074", {}) or {}
            results = trace.get("results", []) or []

            material_rows = []
            for r in results:
                previous = list(r.get("previous_pcbas") or [])
                uses_until_defect = (r.get("previous_pcba_count") or 0) + (1 if r.get("current_pcba_confirmed", True) else 0)
                reuses_until_defect = r.get("previous_pcba_count") or 0
                known_uses = _known_uses(r)
                known_reuses = _known_reuses(r)
                inactive_now = r.get("inactive_now_pcbas")
                if inactive_now is None:
                    inactive_now = previous
                current_now = r.get("active_now_pcbas") or []
                history = _history_text(r)

                row = {
                    "Linha": defect.get("line", ""),
                    "Modelo": defect.get("product_model", ""),
                    "PCBA da falha": pcba,
                    "Defect Time": defect.get("defect_time", ""),
                    "Defect Code": defect.get("defect_code", ""),
                    "Defect Desc": defect.get("defect_desc", ""),
                    "Repair Comment": defect.get("repair_comment", ""),
                    "Reason Type": defect.get("defect_reason_type", ""),
                    "Reason Desc": defect.get("defect_reason_desc", ""),
                    "Tipo material": r.get("item_type", ""),
                    "Material SN": r.get("item_sn", ""),
                    "Iss Status agora": r.get("iss_status_now", ""),
                    "Status reuso na falha": r.get("status", ""),
                    "Reusos anteriores a falha": reuses_until_defect,
                    "Usos ate a falha": uses_until_defect,
                    "Reusos totais conhecidos": known_reuses,
                    "Usos totais conhecidos": known_uses,
                    # "PCBAs antigas" = usos ANTERIORES ao uso da PCBA da falha.
                    # Não usar inactive_now_pcbas aqui, porque se a PCBA atual já foi
                    # desmontada ela também fica N e não pode virar "PCBA antiga".
                    "PCBAs antigas": _join(previous),
                    "PCBAs inativas agora": _join(inactive_now),
                    "PCBAs atuais Y": _join(current_now),
                    "Historico PCBA": history,
                    "Bind Time": r.get("bind_time_utc", ""),
                    "Unbind Time": r.get("unbind_time_utc", ""),
                    "Erro item": r.get("error") or "",
                }
                materiais.append(row)
                material_rows.append(row)
                if (known_reuses or 0) > 0 or reuses_until_defect > 0:
                    reuso.append(row)

            reused_here = [x for x in material_rows if (x.get("Reusos totais conhecidos") or 0) > 0 or (x.get("Reusos anteriores a falha") or 0) > 0]
            # Ordem estavel ajuda a comparar linhas no Excel. Cada material de reuso
            # recebe sua propria coluna na view Falhas_3028.
            reused_here.sort(key=lambda x: (str(x.get("Tipo material") or ""), str(x.get("Material SN") or "")))

            falha_row = {
                "Linha": defect.get("line", ""),
                "Modelo": defect.get("product_model", ""),
                "PCBA": pcba,
                "Defect Time": defect.get("defect_time", ""),
                "Defect Code": defect.get("defect_code", ""),
                "Defect Desc": defect.get("defect_desc", ""),
                "Repair Comment": defect.get("repair_comment", ""),
                "Reason Type": defect.get("defect_reason_type", ""),
                "Reason Desc": defect.get("defect_reason_desc", ""),
                "Status temporal": trace.get("status", ""),
                "Materiais na falha": len(results),
                "Materiais de reuso": len(reused_here),
            }
            for idx, material in enumerate(reused_here, 1):
                falha_row[f"Reuso {idx}"] = _component_card(material)
            falhas.append(falha_row)

    return falhas, materiais, reuso


def _write_sheet(wb, name, rows, headers, freeze_panes="A2", width_overrides=None):
    ws = wb.create_sheet(title=name)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])

    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.font = Font(bold=True, color=HEADER_FONT)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.freeze_panes = freeze_panes
    ws.auto_filter.ref = ws.dimensions

    if ws.max_row >= 2:
        ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        table = Table(displayName=f"T_{name.replace(' ','_').replace('-','_')[:20]}", ref=ref)
        style = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        table.tableStyleInfo = style
        ws.add_table(table)

    for col_idx, header in enumerate(headers, 1):
        max_len = len(str(header))
        for row_idx in range(2, min(ws.max_row, 500) + 1):
            value = ws.cell(row_idx, col_idx).value
            if value is not None:
                max_len = max(max_len, len(str(value)))
        if width_overrides and header in width_overrides:
            width = width_overrides[header]
        elif str(header).startswith("Reuso "):
            width = 36
        else:
            width = min(max(max_len + 2, 11), 52)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = Border(bottom=LIGHT_BORDER)

    return ws


def gerar_excel_3028_3074(resultado, output_path):
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    falhas, materiais, reuso = extrair_linhas(resultado)

    wb = Workbook()
    default = wb.active
    wb.remove(default)

    max_reusos_por_falha = max((int(x.get("Materiais de reuso") or 0) for x in falhas), default=0)
    headers_falhas = [
        "Linha", "Modelo", "PCBA", "Defect Time", "Defect Code", "Defect Desc",
        "Repair Comment", "Reason Type", "Reason Desc", "Status temporal",
        "Materiais na falha", "Materiais de reuso",
    ] + [f"Reuso {i}" for i in range(1, max_reusos_por_falha + 1)]
    headers_materiais = [
        "Linha", "Modelo", "PCBA da falha", "Defect Time", "Defect Desc", "Repair Comment",
        "Tipo material", "Material SN", "Iss Status agora", "Status reuso na falha",
        "Reusos anteriores a falha", "Usos ate a falha", "Reusos totais conhecidos",
        "Usos totais conhecidos", "PCBAs antigas", "PCBAs inativas agora", "PCBAs atuais Y", "Historico PCBA",
        "Bind Time", "Unbind Time", "Erro item",
    ]
    headers_reuso = [
        "Material SN", "Tipo material", "Reusos totais conhecidos", "Usos totais conhecidos",
        "PCBAs antigas", "PCBAs inativas agora", "PCBAs atuais Y", "Historico PCBA", "PCBA da falha",
        "Defect Time", "Defect Desc", "Repair Comment", "Linha", "Modelo",
        "Iss Status agora", "Reusos anteriores a falha", "Usos ate a falha",
    ]

    ws_falhas = _write_sheet(
        wb, "Falhas_3028", falhas, headers_falhas, freeze_panes="D2",
        width_overrides={
            "Linha": 14, "Modelo": 16, "PCBA": 28, "Defect Time": 20,
            "Defect Code": 16, "Defect Desc": 34, "Repair Comment": 36,
            "Reason Type": 14, "Reason Desc": 30, "Status temporal": 36,
            "Materiais na falha": 16, "Materiais de reuso": 16,
        },
    )
    # Na view principal, deixa as colunas de reuso com leitura vertical limpa.
    if ws_falhas.max_row >= 2:
        for row_idx in range(2, ws_falhas.max_row + 1):
            if (ws_falhas.cell(row_idx, 12).value or 0) > 0:
                ws_falhas.row_dimensions[row_idx].height = 92

    _write_sheet(wb, "Materiais", materiais, headers_materiais)
    ws_reuso = _write_sheet(wb, "REUSO", reuso, headers_reuso)

    # Destaca as linhas reutilizadas sem criar dashboard: é apenas uma view operacional.
    if ws_reuso.max_row >= 2:
        for row_idx in range(2, ws_reuso.max_row + 1):
            reuses = ws_reuso.cell(row_idx, 3).value or 0
            fill = MULTI_FILL if isinstance(reuses, (int, float)) and reuses >= 2 else REUSE_FILL
            for c in ws_reuso[row_idx]:
                c.fill = PatternFill("solid", fgColor=fill)

    wb.save(output_path)
    return output_path


def listar_reusos_terminal(resultado):
    _, _, reuso = extrair_linhas(resultado)
    seen = set()
    out = []
    for r in reuso:
        key = (r.get("Material SN"), r.get("PCBA da falha"), r.get("Defect Time"))
        if key in seen:
            continue
        seen.add(key)
        out.append(r)
    return out
