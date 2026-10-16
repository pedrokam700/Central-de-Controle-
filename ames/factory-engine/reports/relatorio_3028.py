from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

FIELDS = [
    "pcba_sn","defect_code","defect_desc","repair_comment","line","defect_time",
    "product_model","repair_user","defect_reason_type","defect_reason_desc",
    "repair_code","repair_desc","defect_material","defect_material_desc",
    "defect_category","defect_oper","defect_oper_desc","imei","work_shift",
    "order_id","manual_or_auto"
]

def gerar_excel_3028(resultado, output_path):
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    registros = []
    for r in resultado["records"]:
        registros.append({k: r.get(k) for k in FIELDS})

    resumo = [{
        "Arquivo": resultado["source"],
        "Linhas de falha": resultado["records_count"],
        "PCBAs unicas": resultado["unique_pcba_count"],
        "Line Id filtro": resultado["search_condition"].get("Line Id",""),
        "Date From": resultado["search_condition"].get("Date From",""),
        "Start Time Zone": resultado["search_condition"].get("Start Time Zone",""),
        "Date To": resultado["search_condition"].get("Date To",""),
        "End Time Zone": resultado["search_condition"].get("End Time Zone",""),
        "Prod Model": resultado["search_condition"].get("Prod Model",""),
        "Quantity": resultado["overall_summary"].get("Quantity",""),
        "Total DefectQty": resultado["overall_summary"].get("Total DefectQty",""),
        "Fpy": resultado["overall_summary"].get("Fpy",""),
        "Warnings": " | ".join(resultado.get("warnings",[])),
    }]

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        pd.DataFrame(resumo).to_excel(writer, sheet_name="Resumo_3028", index=False)
        pd.DataFrame(registros).to_excel(writer, sheet_name="Falhas_3028", index=False)

    wb = load_workbook(output_path)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for c in ws[1]:
            c.font = Font(bold=True)
            c.alignment = Alignment(horizontal="center", vertical="center")
        for ci in range(1, ws.max_column+1):
            max_len = 0
            for ri in range(1, min(ws.max_row, 300)+1):
                v = ws.cell(ri,ci).value
                max_len = max(max_len, len("" if v is None else str(v)))
            ws.column_dimensions[get_column_letter(ci)].width = min(max(max_len+2,10),55)
    wb.save(output_path)
    return output_path
