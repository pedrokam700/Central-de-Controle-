from pathlib import Path
import sqlite3
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


ABAS = {
    "current_defects": "Falhas_Atuais",
    "pcba_components": "Itens_Vinculados",
    "item_links": "Historico_Vinculos",
    "defect_history": "Falhas_Anteriores",
    "process_events": "Timeline_3022",
    "correlations": "Correlacao",
    "errors": "Erros_Log"
}


def _ajustar(path):
    wb = load_workbook(path)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for idx, col in enumerate(ws.columns, 1):
            max_len = 0
            for cell in list(col)[:500]:
                val = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, len(val))
            ws.column_dimensions[get_column_letter(idx)].width = min(max(max_len + 2, 10), 45)
    wb.save(path)


def gerar_relatorio(db_path, output_path):
    db_path = Path(db_path)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    con = sqlite3.connect(db_path)
    try:
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            resumo = []
            for tabela, aba in ABAS.items():
                try:
                    df = pd.read_sql_query(f"SELECT * FROM {tabela}", con)
                except Exception:
                    df = pd.DataFrame()
                df.to_excel(writer, sheet_name=aba, index=False)
                resumo.append({"Tabela": tabela, "Registros": len(df)})
            pd.DataFrame(resumo).to_excel(writer, sheet_name="Resumo", index=False)
    finally:
        con.close()

    _ajustar(output_path)
    return output_path
