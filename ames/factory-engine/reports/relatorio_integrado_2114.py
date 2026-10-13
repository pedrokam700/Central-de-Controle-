from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

from reports.relatorio_3028_3074 import extrair_linhas
from core.correlacao_falhas import analisar_historico_pcba_atual, comparar_falhas, familia_falha


HEADER_FILL = "1F4E78"
HEADER_FONT = "FFFFFF"
LIGHT_BORDER = Side(style="thin", color="D9E1F2")
FILL_SAME = "E2F0D9"
FILL_SIMILAR = "FFF2CC"
FILL_DIFF = "DDEBF7"
FILL_NONE = "E7E6E6"
FILL_ALERT = "FCE4D6"


def _join(values):
    return " | ".join(str(x) for x in (values or []) if x not in (None, ""))


def _split_pipe(value):
    return [x.strip() for x in str(value or "").split("|") if x.strip()]


def _failure_key(pcba, defect_time, defect_code):
    # Deliberadamente não usa a descrição. Assim diferenças de pontuação/Unicode
    # não quebram a ligação Falha 3028 -> materiais 3074.
    return (str(pcba or "").strip(), str(defect_time or "").strip(), str(defect_code or "").strip())


def _pcba_result_map(resultado_2114):
    out = {}
    for item in (resultado_2114 or {}).get("pcbas", []) or []:
        pcba = str(item.get("pcba_sn") or "").strip()
        if pcba:
            out[pcba] = item
    return out


def _histories(pcba_map, pcba):
    state = pcba_map.get(str(pcba or "").strip()) or {}
    if state.get("status") not in (None, "COMPLETE"):
        return []
    result = state.get("result") or {}
    return result.get("rows") or []


def _relation_label(relation):
    return {
        "MESMA_FALHA": "MESMA FALHA",
        "FALHA_SIMILAR": "FALHA SIMILAR",
        "FALHA_DIFERENTE": "FALHA DIFERENTE",
        "SEM_HISTORICO": "SEM HISTÓRICO 2114",
        "SEM_HISTORICO_2114": "SEM HISTÓRICO 2114",
    }.get(relation, str(relation or ""))


def _own_status_label(status):
    return {
        "RECORRENCIA_MESMA_FALHA": "RECORRÊNCIA DA MESMA FALHA",
        "MESMA_FALHA_REPETIDA": "RECORRÊNCIA DA MESMA FALHA",
        "POSSIVEL_RECORRENCIA_SIMILAR": "POSSÍVEL RECORRÊNCIA SIMILAR",
        "FALHA_SIMILAR_NA_MESMA_PCBA": "POSSÍVEL RECORRÊNCIA SIMILAR",
        "HISTORICO_COM_OUTRAS_FALHAS": "HISTÓRICO COM OUTRAS FALHAS",
        "OUTRAS_FALHAS_NA_MESMA_PCBA": "HISTÓRICO COM OUTRAS FALHAS",
        "SEM_REPETICAO_IDENTIFICADA": "SEM REPETIÇÃO IDENTIFICADA",
        "APENAS_UM_HISTORICO": "SEM REPETIÇÃO IDENTIFICADA",
        "SEM_HISTORICO_2114": "SEM HISTÓRICO 2114",
    }.get(status, str(status or ""))


def _defect_type_meaning(value):
    raw = str(value or "").strip().lower().replace(" ", "_")
    if raw in ("phone_disassembly", "phonedisassembly"):
        return "Aparelho desmontado e retornado ao fluxo/linha"
    if raw in ("mainboard", "main_board"):
        return "Falha atribuída à placa principal no fluxo de reparo"
    return ""


def _material_name(raw):
    value = str(raw or "").strip()
    mapping = [
        ("后置副摄像头", "Rear secondary camera"),
        ("后置摄像头2", "Rear camera 2"),
        ("前置摄像头", "Front camera"),
        ("显示屏", "Screen"),
        ("指纹模组", "Fingerprint"),
        ("上盖", "TopCover"),
        ("U板", "UBoard"),
        ("A板", "ABoard"),
        ("C板", "CBoard"),
        ("电池盖", "BatteryCover"),
        ("电池", "Battery"),
    ]
    for token, label in mapping:
        if token in value:
            return label
    return value or "Material"


def _own_analysis(base, pcba_map):
    pcba = base.get("PCBA") or base.get("PCBA da falha")
    histories = _histories(pcba_map, pcba)
    analysis = analisar_historico_pcba_atual(
        base.get("Defect Code", ""), base.get("Defect Desc", ""), histories
    )
    analysis["histories"] = histories
    return analysis


def _reuse_analysis(row, pcba_map):
    old_pcbas = _split_pipe(row.get("PCBAs antigas"))
    details = []
    for use_index, pcba in enumerate(old_pcbas, 1):
        histories = _histories(pcba_map, pcba)
        if not histories:
            details.append({
                "pcba": pcba,
                "historical_use_index": use_index,
                "history": None,
                "relation": "SEM_HISTORICO_2114",
                "reason": "PCBA anterior sem histórico retornado na 2114",
                "current_family": familia_falha(row.get("Defect Desc", "")),
                "old_family": "",
            })
            continue
        for hist in histories:
            cmp = comparar_falhas(
                row.get("Defect Code", ""), row.get("Defect Desc", ""),
                hist.get("Defect Code", ""), hist.get("Defect Description", ""),
            )
            details.append({
                "pcba": pcba,
                "historical_use_index": use_index,
                "history": hist,
                **cmp,
            })

    same = sum(1 for x in details if x["relation"] == "MESMA_FALHA")
    similar = sum(1 for x in details if x["relation"] == "FALHA_SIMILAR")
    different = sum(1 for x in details if x["relation"] == "FALHA_DIFERENTE")
    if same:
        strongest = "MESMA_FALHA"
    elif similar:
        strongest = "FALHA_SIMILAR"
    elif different:
        strongest = "FALHA_DIFERENTE"
    else:
        strongest = "SEM_HISTORICO_2114"
    return {
        "old_pcbas": old_pcbas,
        "details": details,
        "same_count": same,
        "similar_count": similar,
        "different_count": different,
        "strongest": strongest,
    }


def _reuse_card(row, analysis):
    previous_reuses = int(row.get("Reusos anteriores a falha") or 0)
    use_number = previous_reuses + 1
    uses_known = row.get("Usos totais conhecidos")
    reuses_known = row.get("Reusos totais conhecidos")
    lines = [
        _material_name(row.get("Tipo material")),
        f"SN: {row.get('Material SN') or '-'}",
        f"Uso na falha: {use_number}º | Reusos antes: {previous_reuses}",
        f"Usos conhecidos agora: {uses_known if uses_known is not None else '-'} | Reusos agora: {reuses_known if reuses_known is not None else '-'}",
        "PCBAs anteriores: " + (_join(analysis["old_pcbas"]) or "-"),
        f"Correlação histórica: {_relation_label(analysis['strongest'])}",
        f"Iguais: {analysis['same_count']} | Similares: {analysis['similar_count']} | Diferentes: {analysis['different_count']}",
    ]
    return "\n".join(lines)


def _build_views(resultado_3074, resultado_2114):
    falhas, materiais, reusos = extrair_linhas(resultado_3074)
    pcba_map = _pcba_result_map(resultado_2114)

    reuso_by_key = {}
    for r in reusos:
        key = _failure_key(r.get("PCBA da falha"), r.get("Defect Time"), r.get("Defect Code"))
        reuso_by_key.setdefault(key, []).append(r)

    analise_rows = []
    hist_pcba_rows = []
    hist_reuso_rows = []
    max_reusos = 0
    cards_total = 0

    for f in falhas:
        key = _failure_key(f.get("PCBA"), f.get("Defect Time"), f.get("Defect Code"))
        own = _own_analysis(f, pcba_map)
        related_reuses = sorted(
            reuso_by_key.get(key, []),
            key=lambda x: (str(x.get("Tipo material") or ""), str(x.get("Material SN") or "")),
        )
        cards_total += len(related_reuses)
        max_reusos = max(max_reusos, len(related_reuses))

        main = {
            "Linha": f.get("Linha", ""),
            "Modelo": f.get("Modelo", ""),
            "PCBA": f.get("PCBA", ""),
            "Defect Time": f.get("Defect Time", ""),
            "Defect Code": f.get("Defect Code", ""),
            "Defect Desc": f.get("Defect Desc", ""),
            "Repair Comment": f.get("Repair Comment", ""),
            "Reason Type": f.get("Reason Type", ""),
            "Reason Desc": f.get("Reason Desc", ""),
            "Históricos 2114 da PCBA": own["total"],
            "Repetições da mesma falha": own["same_repetitions"],
            "Falhas similares na PCBA": own["similar_count"],
            "Falhas diferentes na PCBA": own["different_count"],
            "Situação da PCBA": _own_status_label(own["status"]),
            "Materiais na falha": f.get("Materiais na falha", 0),
            "Materiais de reuso": len(related_reuses),
        }

        for idx, r in enumerate(related_reuses, 1):
            ra = _reuse_analysis(r, pcba_map)
            main[f"Reuso {idx}"] = _reuse_card(r, ra)

            prev = int(r.get("Reusos anteriores a falha") or 0)
            use_number = prev + 1
            for detail in ra["details"]:
                hist = detail.get("history") or {}
                hist_reuso_rows.append({
                    "Linha": r.get("Linha", ""),
                    "Modelo": r.get("Modelo", ""),
                    "PCBA atual": r.get("PCBA da falha", ""),
                    "Defect Time atual": r.get("Defect Time", ""),
                    "Defect Code atual": r.get("Defect Code", ""),
                    "Defect Desc atual": r.get("Defect Desc", ""),
                    "Tipo material": _material_name(r.get("Tipo material", "")),
                    "Material SN": r.get("Material SN", ""),
                    "Uso na falha": f"{use_number}º uso",
                    "Reusos anteriores": prev,
                    "Usos conhecidos agora": r.get("Usos totais conhecidos"),
                    "Reusos conhecidos agora": r.get("Reusos totais conhecidos"),
                    "PCBA anterior": detail.get("pcba", ""),
                    "Ordem de uso anterior": detail.get("historical_use_index", ""),
                    "Hist Seq": hist.get("Defect Hist Seq", ""),
                    "Defect Location antigo": hist.get("Defect Location", ""),
                    "Defect Material ID antigo": hist.get("Defect Material ID", ""),
                    "Defect Code antigo": hist.get("Defect Code", ""),
                    "Defect Description antigo": hist.get("Defect Description", "") or ("SEM HISTÓRICO 2114" if not hist else ""),
                    "Defect Oper antigo": hist.get("Defect Oper", ""),
                    "Repair Status antigo": hist.get("Repair Status", ""),
                    "Defect Type antigo": hist.get("Defect Type", ""),
                    "Interpretação Defect Type": _defect_type_meaning(hist.get("Defect Type", "")),
                    "Relação com falha atual": _relation_label(detail.get("relation", "SEM_HISTORICO_2114")),
                    "Motivo da relação": detail.get("reason", ""),
                })

        analise_rows.append(main)

        histories = own["histories"]
        if not histories:
            hist_pcba_rows.append({
                "Linha": f.get("Linha", ""), "Modelo": f.get("Modelo", ""), "PCBA": f.get("PCBA", ""),
                "Defect Time atual": f.get("Defect Time", ""), "Defect Code atual": f.get("Defect Code", ""),
                "Defect Desc atual": f.get("Defect Desc", ""), "Repair Comment atual": f.get("Repair Comment", ""),
                "Hist Seq": "", "Defect Location 2114": "", "Defect Material ID 2114": "",
                "Defect Code 2114": "", "Defect Description 2114": "SEM HISTÓRICO 2114",
                "Defect Oper 2114": "", "Repair Status 2114": "", "Defect Type 2114": "",
                "Interpretação Defect Type": "", "Relação com falha atual": "SEM HISTÓRICO 2114",
                "Motivo da relação": "PCBA sem histórico retornado na 2114",
            })
        else:
            for hist in histories:
                cmp = comparar_falhas(
                    f.get("Defect Code", ""), f.get("Defect Desc", ""),
                    hist.get("Defect Code", ""), hist.get("Defect Description", ""),
                )
                hist_pcba_rows.append({
                    "Linha": f.get("Linha", ""), "Modelo": f.get("Modelo", ""), "PCBA": f.get("PCBA", ""),
                    "Defect Time atual": f.get("Defect Time", ""), "Defect Code atual": f.get("Defect Code", ""),
                    "Defect Desc atual": f.get("Defect Desc", ""), "Repair Comment atual": f.get("Repair Comment", ""),
                    "Hist Seq": hist.get("Defect Hist Seq", ""),
                    "Defect Location 2114": hist.get("Defect Location", ""),
                    "Defect Material ID 2114": hist.get("Defect Material ID", ""),
                    "Defect Code 2114": hist.get("Defect Code", ""),
                    "Defect Description 2114": hist.get("Defect Description", ""),
                    "Defect Oper 2114": hist.get("Defect Oper", ""),
                    "Repair Status 2114": hist.get("Repair Status", ""),
                    "Defect Type 2114": hist.get("Defect Type", ""),
                    "Interpretação Defect Type": _defect_type_meaning(hist.get("Defect Type", "")),
                    "Relação com falha atual": _relation_label(cmp["relation"]),
                    "Motivo da relação": cmp["reason"],
                })

    # Gate contra regressão silenciosa: se o motor detectou reuso, a view principal
    # precisa obrigatoriamente conter a mesma quantidade de cards de material.
    if len(analise_rows) != len(falhas):
        raise RuntimeError(f"Relatório inconsistente: {len(falhas)} falhas, mas {len(analise_rows)} linhas em ANALISE.")
    if len(reusos) != cards_total:
        raise RuntimeError(
            f"Relatório inconsistente: 3074 detectou {len(reusos)} materiais de reuso, "
            f"mas ANALISE montou {cards_total}. A geração foi interrompida para não criar view vazia/incorreta."
        )
    if reusos and not hist_reuso_rows:
        raise RuntimeError("Relatório inconsistente: existem materiais de reuso, mas HIST_REUSO ficou vazio.")

    return {
        "falhas": falhas,
        "materiais": materiais,
        "reusos": reusos,
        "analise_rows": analise_rows,
        "hist_pcba_rows": hist_pcba_rows,
        "hist_reuso_rows": hist_reuso_rows,
        "max_reusos": max_reusos,
        "stats": {
            "falhas": len(falhas),
            "materiais": len(materiais),
            "materiais_reuso": len(reusos),
            "linhas_hist_pcba": len(hist_pcba_rows),
            "linhas_hist_reuso": len(hist_reuso_rows),
            "max_reusos_por_falha": max_reusos,
        },
    }


def _write_sheet(wb, name, rows, headers, widths=None, empty_message=None, table_name=None):
    ws = wb.create_sheet(title=name)
    ws.append(headers)

    for row in rows:
        ws.append([row.get(h, "") for h in headers])

    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.font = Font(bold=True, color=HEADER_FONT)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=LIGHT_BORDER)

    # O usuário rejeitou a barra de divisão/congelamento. Mantemos tudo livre.
    ws.freeze_panes = None
    ws.sheet_view.pane = None
    ws.sheet_view.zoomScale = 85
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(rows)+1)}"

    if rows:
        ref = f"A1:{get_column_letter(len(headers))}{len(rows)+1}"
        table = Table(displayName=table_name or f"T_{name[:20]}", ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False,
        )
        ws.add_table(table)
    elif empty_message:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
        c = ws.cell(2, 1)
        c.value = empty_message
        c.fill = PatternFill("solid", fgColor=FILL_NONE)
        c.font = Font(bold=True, color="666666")
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[2].height = 34

    for idx, header in enumerate(headers, 1):
        width = (widths or {}).get(header)
        if width is None:
            if str(header).startswith("Reuso "):
                width = 42
            elif "Description" in header or "Desc" in header or "Comment" in header:
                width = 30
            elif "PCBA" in header or "Material SN" in header:
                width = 27
            elif "Interpretação" in header or "Motivo" in header or "Situação" in header:
                width = 28
            else:
                width = 16
        ws.column_dimensions[get_column_letter(idx)].width = width

    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = Border(bottom=LIGHT_BORDER)
    return ws


def _color_relation_column(ws, headers, header_name):
    if header_name not in headers or ws.max_row < 2:
        return
    col = headers.index(header_name) + 1
    for r in range(2, ws.max_row + 1):
        cell = ws.cell(r, col)
        value = str(cell.value or "").upper()
        if "MESMA FALHA" in value or "RECORRÊNCIA DA MESMA" in value:
            color = FILL_SAME
        elif "SIMILAR" in value:
            color = FILL_SIMILAR
        elif "DIFERENTE" in value or "OUTRAS FALHAS" in value:
            color = FILL_DIFF
        elif "SEM HIST" in value or "SEM REPET" in value:
            color = FILL_NONE
        else:
            continue
        cell.fill = PatternFill("solid", fgColor=color)
        cell.font = Font(bold=True)


def gerar_excel_integrado(resultado_3074, resultado_2114, output_path):
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    views = _build_views(resultado_3074, resultado_2114)

    wb = Workbook()
    wb.remove(wb.active)

    headers_analise = [
        "Linha", "Modelo", "PCBA", "Defect Time", "Defect Code", "Defect Desc",
        "Repair Comment", "Reason Type", "Reason Desc",
        "Históricos 2114 da PCBA", "Repetições da mesma falha",
        "Falhas similares na PCBA", "Falhas diferentes na PCBA", "Situação da PCBA",
        "Materiais na falha", "Materiais de reuso",
    ] + [f"Reuso {i}" for i in range(1, views["max_reusos"] + 1)]

    ws_main = _write_sheet(
        wb, "ANALISE", views["analise_rows"], headers_analise,
        widths={
            "Linha": 13, "Modelo": 15, "PCBA": 27, "Defect Time": 20,
            "Defect Code": 15, "Defect Desc": 32, "Repair Comment": 30,
            "Reason Type": 14, "Reason Desc": 28, "Históricos 2114 da PCBA": 16,
            "Repetições da mesma falha": 16, "Falhas similares na PCBA": 15,
            "Falhas diferentes na PCBA": 15, "Situação da PCBA": 28,
            "Materiais na falha": 14, "Materiais de reuso": 14,
        },
        empty_message="Nenhuma falha foi processada neste recorte.",
        table_name="T_ANALISE_V012",
    )
    _color_relation_column(ws_main, headers_analise, "Situação da PCBA")
    for r in range(2, ws_main.max_row + 1):
        has_reuse = ws_main.cell(r, headers_analise.index("Materiais de reuso") + 1).value or 0
        ws_main.row_dimensions[r].height = 104 if has_reuse else 36

    headers_pcba = [
        "Linha", "Modelo", "PCBA", "Defect Time atual", "Defect Code atual", "Defect Desc atual",
        "Repair Comment atual", "Hist Seq", "Defect Location 2114", "Defect Material ID 2114",
        "Defect Code 2114", "Defect Description 2114", "Defect Oper 2114", "Repair Status 2114",
        "Defect Type 2114", "Interpretação Defect Type", "Relação com falha atual", "Motivo da relação",
    ]
    ws_pcba = _write_sheet(
        wb, "HIST_PCBA", views["hist_pcba_rows"], headers_pcba,
        widths={"Defect Desc atual": 30, "Repair Comment atual": 28, "Defect Description 2114": 34,
                "Interpretação Defect Type": 34, "Relação com falha atual": 22, "Motivo da relação": 28},
        empty_message="Nenhum histórico 2114 foi retornado para as PCBAs processadas.",
        table_name="T_HIST_PCBA_V012",
    )
    _color_relation_column(ws_pcba, headers_pcba, "Relação com falha atual")

    headers_reuso = [
        "Linha", "Modelo", "PCBA atual", "Defect Time atual", "Defect Code atual", "Defect Desc atual",
        "Tipo material", "Material SN", "Uso na falha", "Reusos anteriores", "Usos conhecidos agora",
        "Reusos conhecidos agora", "PCBA anterior", "Ordem de uso anterior", "Hist Seq",
        "Defect Location antigo", "Defect Material ID antigo", "Defect Code antigo",
        "Defect Description antigo", "Defect Oper antigo", "Repair Status antigo", "Defect Type antigo",
        "Interpretação Defect Type", "Relação com falha atual", "Motivo da relação",
    ]
    ws_reuso = _write_sheet(
        wb, "HIST_REUSO", views["hist_reuso_rows"], headers_reuso,
        widths={"Defect Desc atual": 30, "Material SN": 29, "Defect Description antigo": 34,
                "Interpretação Defect Type": 34, "Relação com falha atual": 22, "Motivo da relação": 28},
        empty_message=(
            "Nenhum material de segundo uso/reuso foi encontrado nas PCBAs processadas. "
            "Esta view está vazia de forma intencional para este recorte."
        ),
        table_name="T_HIST_REUSO_V012",
    )
    _color_relation_column(ws_reuso, headers_reuso, "Relação com falha atual")

    # ANALISE é sempre a primeira tela operacional.
    wb.active = wb.sheetnames.index("ANALISE")
    wb.save(output_path)
    return views["stats"]


# Compatibilidade com chamadas antigas/tests externos.
def _history_line_compat(hist):
    parts = []
    seq = hist.get("Defect Hist Seq", "")
    if seq not in (None, ""):
        parts.append(f"Seq {seq}")
    if hist.get("Defect Code"):
        parts.append(str(hist.get("Defect Code")))
    if hist.get("Defect Description"):
        parts.append(str(hist.get("Defect Description")))
    if hist.get("Defect Type"):
        parts.append(f"Type={hist.get('Defect Type')}")
    if hist.get("Repair Status"):
        parts.append(f"Repair={hist.get('Repair Status')}")
    return " | ".join(parts)


def extrair_reuso_2114(resultado_3074, resultado_2114):
    """Compatibilidade com as APIs V0.10/V0.11 sem afetar as 3 views V0.12."""
    _, _, reusos = extrair_linhas(resultado_3074)
    pcba_map = _pcba_result_map(resultado_2114)
    resumo_rows = []
    detail_rows = []
    max_old = 0

    for r in reusos:
        antigas = _split_pipe(r.get("PCBAs antigas"))
        max_old = max(max_old, len(antigas))
        row = {
            "Linha": r.get("Linha", ""),
            "Modelo": r.get("Modelo", ""),
            "PCBA da falha": r.get("PCBA da falha", ""),
            "Defect Time": r.get("Defect Time", ""),
            "Defect Desc atual": r.get("Defect Desc", ""),
            "Repair Comment atual": r.get("Repair Comment", ""),
            "Tipo material": r.get("Tipo material", ""),
            "Material SN": r.get("Material SN", ""),
            "Usos totais conhecidos": r.get("Usos totais conhecidos"),
            "Reusos totais conhecidos": r.get("Reusos totais conhecidos"),
            "PCBAs atuais Y": r.get("PCBAs atuais Y", ""),
            "Qtde PCBAs antigas": len(antigas),
        }
        for idx, pcba in enumerate(antigas, 1):
            row[f"PCBA antiga {idx}"] = pcba
            histories = _histories(pcba_map, pcba)
            row[f"Historico 2114 - PCBA {idx}"] = (
                "\n".join(_history_line_compat(h) for h in histories)
                if histories else "SEM HISTORICO NA 2114"
            )
        resumo_rows.append(row)

    for ctx in (resultado_2114 or {}).get("contexts", []) or []:
        old_pcba = ctx.get("historical_pcba", "")
        state = pcba_map.get(old_pcba) or {}
        query = state.get("result") or {}
        histories = query.get("rows") or []
        if not histories:
            histories = [None]
        for hist in histories:
            h = hist or {}
            detail_rows.append({
                "Material SN": ctx.get("material_sn", ""),
                "Tipo material": ctx.get("material_type", ""),
                "PCBA da falha atual": ctx.get("current_pcba", ""),
                "Defect Time atual": ctx.get("defect_time", ""),
                "Defect Code atual": ctx.get("defect_code", ""),
                "Defect Desc atual": ctx.get("defect_desc", ""),
                "Repair Comment atual": ctx.get("repair_comment", ""),
                "PCBA antiga": old_pcba,
                "Ordem uso antigo": ctx.get("historical_use_index", ""),
                "Status consulta 2114": state.get("status", "PENDING"),
                "Erro consulta 2114": state.get("error") or "",
                "Shift 2114": query.get("shift", ""),
                "Hist Seq": h.get("Defect Hist Seq", ""),
                "Defect Location": h.get("Defect Location", ""),
                "Defect Material ID": h.get("Defect Material ID", ""),
                "Defect Code antigo": h.get("Defect Code", ""),
                "Defect Description antigo": h.get("Defect Description", ""),
                "Defect Oper": h.get("Defect Oper", ""),
                "Repair Status": h.get("Repair Status", ""),
                "Manual/Auto": h.get("TestTools Auto Defect Or Defect By Hand", ""),
                "Defect Type": h.get("Defect Type", ""),
                "Defect Type Class": h.get("Defect Type Class", ""),
            })
    return resumo_rows, detail_rows, max_old
