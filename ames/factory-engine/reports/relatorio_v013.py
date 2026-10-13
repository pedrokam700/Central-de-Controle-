from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

from reports.relatorio_3028_3074 import extrair_linhas
from reports.relatorio_integrado_2114 import _pcba_result_map, _histories, _material_name, _split_pipe
from core.correlacao_falhas import analisar_historico_pcba_atual, comparar_falhas, familia_falha
from core.uso_pcba import reconstruir_usos_pcba

HEADER_FILL = "1F4E78"
HEADER_FONT = "FFFFFF"
SECTION_FILL = "D9EAF7"
KPI_FILL = "EAF2F8"
FILL_SAME = "E2F0D9"
FILL_SIMILAR = "FFF2CC"
FILL_DIFF = "DDEBF7"
FILL_NONE = "E7E6E6"
FILL_ALERT = "FCE4D6"
LIGHT_BORDER = Side(style="thin", color="D9E1F2")


def _join(values):
    return " | ".join(str(x) for x in (values or []) if x not in (None, ""))


def _failure_key(pcba, defect_time, defect_code):
    return (str(pcba or "").strip(), str(defect_time or "").strip(), str(defect_code or "").strip())


def _relation_label(v):
    return {
        "MESMA_FALHA": "MESMA FALHA",
        "FALHA_SIMILAR": "MESMA FAMÍLIA",
        "FALHA_DIFERENTE": "FALHA DIFERENTE",
        "SEM_HISTORICO": "SEM HISTÓRICO 2114",
        "SEM_HISTORICO_2114": "SEM HISTÓRICO 2114",
    }.get(v, str(v or ""))


def _own_status(a):
    if a.get("same_repetitions", 0) > 0:
        return "RECORRÊNCIA DA MESMA FALHA"
    if a.get("similar_count", 0) > 0:
        return "RECORRÊNCIA DA MESMA FAMÍLIA"
    if a.get("total", 0) > 1:
        return "HISTÓRICO COM OUTRAS FALHAS"
    if a.get("total", 0) == 1:
        return "SEM REPETIÇÃO IDENTIFICADA"
    return "SEM HISTÓRICO 2114"


def _process_state(repair_status, defect_type):
    rs = str(repair_status or "").strip().upper()
    dt_raw = str(defect_type or "").strip()
    dt = dt_raw.lower().replace(" ", "_")
    if rs == "N" and not dt_raw:
        return "AGUARDANDO ANÁLISE DO AUXILIAR"
    if dt in ("mainboard", "main_board"):
        return "REPARO DE PLACA CONCLUÍDO" if rs == "Y" else "ENCAMINHADO / AGUARDANDO REPARO DE PLACA"
    if dt in ("phone_disassembly", "phonedisassembly"):
        return "PHONE DISASSEMBLY CONCLUÍDO" if rs == "Y" else "PHONE DISASSEMBLY PENDENTE"
    if dt_raw:
        return ("FINALIZADO" if rs == "Y" else "PENDENTE / EM FLUXO") + f" — {dt_raw}"
    if rs == "Y":
        return "FINALIZADO — TIPO NÃO INFORMADO"
    return "STATUS NÃO CLASSIFICADO"


def _manual_auto(hist):
    return str(hist.get("TestTools Auto Defect Or Defect By Hand") or "").strip()


def _pcba_item_map(resultado_3074):
    return {str(x.get("pcba_sn") or "").strip(): x for x in (resultado_3074 or {}).get("items", []) or [] if x.get("pcba_sn")}


def _reuse_analysis(row, pcba_map):
    old_pcbas = _split_pipe(row.get("PCBAs antigas"))
    inactive = set(_split_pipe(row.get("PCBAs inativas agora")))
    active = set(_split_pipe(row.get("PCBAs atuais Y")))
    details = []
    for use_index, pcba in enumerate(old_pcbas, 1):
        state = "DESVINCULADA" if pcba in inactive else ("VINCULADA AGORA" if pcba in active else "HISTÓRICA")
        histories = _histories(pcba_map, pcba)
        if not histories:
            details.append({"pcba": pcba, "use_index": use_index, "state": state, "history": None,
                            "relation": "SEM_HISTORICO_2114", "reason": "PCBA sem histórico retornado na 2114"})
            continue
        for hist in histories:
            cmp = comparar_falhas(row.get("Defect Code", ""), row.get("Defect Desc", ""),
                                  hist.get("Defect Code", ""), hist.get("Defect Description", ""))
            details.append({"pcba": pcba, "use_index": use_index, "state": state, "history": hist, **cmp})
    same = sum(1 for x in details if x.get("relation") == "MESMA_FALHA")
    similar = sum(1 for x in details if x.get("relation") == "FALHA_SIMILAR")
    diff = sum(1 for x in details if x.get("relation") == "FALHA_DIFERENTE")
    strongest = "MESMA_FALHA" if same else ("FALHA_SIMILAR" if similar else ("FALHA_DIFERENTE" if diff else "SEM_HISTORICO_2114"))
    return {"old_pcbas": old_pcbas, "details": details, "same": same, "similar": similar, "different": diff, "strongest": strongest}


def _compact_reuse_card(r, ra):
    prev = int(r.get("Reusos anteriores a falha") or 0)
    use = prev + 1
    return (
        f"{_material_name(r.get('Tipo material'))}\n"
        f"SN: {r.get('Material SN') or '-'}\n"
        f"Uso na falha: {use}º | Reusos antes: {prev}\n"
        f"PCBAs desvinculadas: {len(ra['old_pcbas'])}\n"
        f"Histórico: {_relation_label(ra['strongest'])}"
    )



def _usage_cycles_summary(usage):
    parts=[]
    for c in usage.get("cycles", []) or []:
        start=str(c.get("start_utc") or "").replace("+00:00", " UTC")
        end=str(c.get("end_utc") or "").replace("+00:00", " UTC") or "ATIVO"
        parts.append(f"{c.get('use_index')}º: {start} -> {end} [{c.get('link_state','')}]" )
    return "\n".join(parts)


def _build_views(resultado_3074, resultado_2114):
    falhas, materiais, reusos = extrair_linhas(resultado_3074)
    pcba_map = _pcba_result_map(resultado_2114)
    item_map = _pcba_item_map(resultado_3074)

    reuso_by_key = defaultdict(list)
    for r in reusos:
        reuso_by_key[_failure_key(r.get("PCBA da falha"), r.get("Defect Time"), r.get("Defect Code"))].append(r)

    analise_rows, hist_pcba_rows, hist_material_rows = [], [], []
    max_reusos = 0
    pcba_usage_by_context = {}

    for analise_idx, f in enumerate(falhas, 1):
        pcba = str(f.get("PCBA") or "").strip()
        defect_time = f.get("Defect Time", "")
        defect_code = f.get("Defect Code", "")
        key = _failure_key(pcba, defect_time, defect_code)
        histories = _histories(pcba_map, pcba)
        own = analisar_historico_pcba_atual(defect_code, f.get("Defect Desc", ""), histories)
        base_rows = ((item_map.get(pcba) or {}).get("base_query") or {}).get("rows") or []
        usage = reconstruir_usos_pcba(base_rows, pcba, defect_time)
        pcba_usage_by_context[key] = usage

        related = sorted(reuso_by_key.get(key, []), key=lambda x:(str(x.get("Tipo material") or ""), str(x.get("Material SN") or "")))
        max_reusos = max(max_reusos, len(related))

        row = {
            "Linha": f.get("Linha", ""), "Modelo": f.get("Modelo", ""), "PCBA": pcba,
            "Defect Time": defect_time, "Defect Code": defect_code, "Defect Desc": f.get("Defect Desc", ""),
            "Repair Comment": f.get("Repair Comment", ""), "Reason Type": f.get("Reason Type", ""), "Reason Desc": f.get("Reason Desc", ""),
            "Uso da PCBA na falha": f"{usage.get('use_index_at_defect') or 0}º uso" if usage.get('use_index_at_defect') else "NÃO DETERMINADO",
            "Usos conhecidos da PCBA": usage.get("known_use_count", 0),
            "Reusos anteriores da PCBA": usage.get("reuses_before_defect", 0),
            "_pcba_use_index": int(usage.get("use_index_at_defect") or 0),
            "Históricos da PCBA": own.get("total", 0),
            "Mesma falha anterior": own.get("same_repetitions", 0),
            "Mesma família na PCBA": own.get("similar_count", 0),
            "Falhas diferentes na PCBA": own.get("different_count", 0),
            "Situação da PCBA": _own_status(own),
            "Materiais na falha": f.get("Materiais na falha", 0),
            "Materiais 2º+ uso": len(related),
            "_analysis_row": analise_idx + 1,
            "_failure_key": key,
        }
        for idx, r in enumerate(related, 1):
            ra = _reuse_analysis(r, pcba_map)
            row[f"Reuso {idx}"] = _compact_reuse_card(r, ra)
            row[f"_reuso_{idx}_sn"] = r.get("Material SN", "")
            row[f"_reuso_{idx}_key"] = (key, r.get("Material SN", ""))

            prev = int(r.get("Reusos anteriores a falha") or 0)
            use_no = prev + 1
            for detail in ra["details"]:
                hist = detail.get("history") or {}
                hist_material_rows.append({
                    "_failure_key": key, "_material_key": (key, r.get("Material SN", "")), "_analysis_row": analise_idx + 1,
                    "Linha": r.get("Linha", ""), "Modelo": r.get("Modelo", ""), "PCBA atual": r.get("PCBA da falha", ""),
                    "Defect Time atual": r.get("Defect Time", ""), "Defect Code atual": r.get("Defect Code", ""), "Defect Desc atual": r.get("Defect Desc", ""),
                    "Tipo material": _material_name(r.get("Tipo material", "")), "Material SN": r.get("Material SN", ""),
                    "Uso material na falha": f"{use_no}º uso", "Reusos anteriores material": prev,
                    "Usos conhecidos material": r.get("Usos totais conhecidos"), "Reusos conhecidos material": r.get("Reusos totais conhecidos"),
                    "PCBA desvinculada": detail.get("pcba", ""), "Estado vínculo atual": detail.get("state", ""), "Ordem de uso anterior": detail.get("use_index", ""),
                    "Hist Seq": hist.get("Defect Hist Seq", ""), "Defect Code antigo": hist.get("Defect Code", ""),
                    "Defect Description antigo": hist.get("Defect Description", "") or ("SEM HISTÓRICO 2114" if not hist else ""),
                    "Defect Oper antigo": hist.get("Defect Oper", ""), "Registro falha antigo": _manual_auto(hist),
                    "Repair Status antigo": hist.get("Repair Status", ""), "Defect Type antigo": hist.get("Defect Type", ""),
                    "Situação reparo antiga": _process_state(hist.get("Repair Status", ""), hist.get("Defect Type", "")),
                    "Família falha antiga": familia_falha(hist.get("Defect Description", ""), hist.get("Defect Code", "")) if hist else "",
                    "Relação com falha atual": _relation_label(detail.get("relation")), "Motivo da relação": detail.get("reason", ""),
                })

        analise_rows.append(row)

        if not histories:
            hist_pcba_rows.append({
                "_failure_key": key, "_analysis_row": analise_idx + 1,
                "Linha": f.get("Linha", ""), "Modelo": f.get("Modelo", ""), "PCBA": pcba,
                "Uso PCBA na falha": row["Uso da PCBA na falha"], "Usos conhecidos PCBA": usage.get("known_use_count",0),
                "Reusos anteriores PCBA": usage.get("reuses_before_defect",0), "Ciclos montagem PCBA 3074": _usage_cycles_summary(usage), "Defect Time atual": defect_time,
                "Defect Code atual": defect_code, "Defect Desc atual": f.get("Defect Desc", ""), "Repair Comment atual": f.get("Repair Comment", ""),
                "Hist Seq": "", "Defect Code 2114": "", "Defect Description 2114": "SEM HISTÓRICO 2114",
                "Defect Oper 2114": "", "Registro falha 2114": "", "Repair Status 2114": "", "Defect Type 2114": "",
                "Situação reparo": "SEM HISTÓRICO 2114", "Família 2114": "", "Relação com falha atual": "SEM HISTÓRICO 2114", "Motivo da relação": "PCBA sem histórico retornado na 2114",
            })
        else:
            for hist in histories:
                cmp = comparar_falhas(defect_code, f.get("Defect Desc", ""), hist.get("Defect Code", ""), hist.get("Defect Description", ""))
                hist_pcba_rows.append({
                    "_failure_key": key, "_analysis_row": analise_idx + 1,
                    "Linha": f.get("Linha", ""), "Modelo": f.get("Modelo", ""), "PCBA": pcba,
                    "Uso PCBA na falha": row["Uso da PCBA na falha"], "Usos conhecidos PCBA": usage.get("known_use_count",0),
                    "Reusos anteriores PCBA": usage.get("reuses_before_defect",0), "Ciclos montagem PCBA 3074": _usage_cycles_summary(usage), "Defect Time atual": defect_time,
                    "Defect Code atual": defect_code, "Defect Desc atual": f.get("Defect Desc", ""), "Repair Comment atual": f.get("Repair Comment", ""),
                    "Hist Seq": hist.get("Defect Hist Seq", ""), "Defect Code 2114": hist.get("Defect Code", ""),
                    "Defect Description 2114": hist.get("Defect Description", ""), "Defect Oper 2114": hist.get("Defect Oper", ""),
                    "Registro falha 2114": _manual_auto(hist), "Repair Status 2114": hist.get("Repair Status", ""), "Defect Type 2114": hist.get("Defect Type", ""),
                    "Situação reparo": _process_state(hist.get("Repair Status", ""), hist.get("Defect Type", "")),
                    "Família 2114": familia_falha(hist.get("Defect Description", ""), hist.get("Defect Code", "")),
                    "Relação com falha atual": _relation_label(cmp.get("relation")), "Motivo da relação": cmp.get("reason", ""),
                })

    # Gates: toda falha atual precisa aparecer em ANALISE e todo histórico retornado precisa estar acessível.
    if len(analise_rows) != len(falhas):
        raise RuntimeError("ANALISE perdeu falhas durante a consolidação.")
    expected_hist = 0
    for f in falhas:
        hs = _histories(pcba_map, f.get("PCBA"))
        expected_hist += max(1, len(hs))
    if len(hist_pcba_rows) != expected_hist:
        raise RuntimeError(f"HIST_PCBA inconsistente: esperado {expected_hist}, obtido {len(hist_pcba_rows)}.")
    if reusos and not hist_material_rows:
        raise RuntimeError("Existem materiais reutilizados, mas HIST_MATERIAL ficou vazio.")

    return {
        "falhas": falhas, "materiais": materiais, "reusos": reusos,
        "analise_rows": analise_rows, "hist_pcba_rows": hist_pcba_rows, "hist_material_rows": hist_material_rows,
        "max_reusos": max_reusos, "pcba_usage": pcba_usage_by_context,
    }


def _unique_histories(hist_rows):
    seen, out = set(), []
    for r in hist_rows:
        key = (r.get("PCBA"), r.get("Hist Seq"), r.get("Defect Code 2114"), r.get("Defect Description 2114"), r.get("Defect Type 2114"), r.get("Repair Status 2114"))
        if key in seen: continue
        seen.add(key); out.append(r)
    return out


def _build_indicators(views):
    rows = views["analise_rows"]
    unique_pcba = {}
    for r in rows:
        pcba = r.get("PCBA")
        if pcba not in unique_pcba or str(r.get("Defect Time")) > str(unique_pcba[pcba].get("Defect Time")):
            unique_pcba[pcba] = r
    pcba_rows = list(unique_pcba.values())

    pcba_first = sum(1 for r in pcba_rows if int(r.get("_pcba_use_index") or 0) <= 1)
    pcba_second = sum(1 for r in pcba_rows if int(r.get("_pcba_use_index") or 0) == 2)
    pcba_multi = sum(1 for r in pcba_rows if int(r.get("_pcba_use_index") or 0) >= 3)
    pcba_hist2 = sum(1 for r in pcba_rows if int(r.get("Históricos da PCBA") or 0) >= 2)
    pcba_same = sum(1 for r in pcba_rows if int(r.get("Mesma falha anterior") or 0) > 0)
    pcba_sim = sum(1 for r in pcba_rows if int(r.get("Mesma família na PCBA") or 0) > 0)
    pcba_hist3 = sum(1 for r in pcba_rows if int(r.get("Históricos da PCBA") or 0) >= 3)

    # Materiais únicos, para não inflar o indicador por múltiplas linhas de histórico.
    material_by_sn = {}
    for r in views["reusos"]:
        sn = r.get("Material SN")
        if sn: material_by_sn[sn] = r
    material_rows = list(material_by_sn.values())
    mat_second = sum(1 for r in material_rows if int(r.get("Reusos anteriores a falha") or 0) == 1)
    mat_multi = sum(1 for r in material_rows if int(r.get("Reusos anteriores a falha") or 0) >= 2)
    mat_types = Counter(_material_name(r.get("Tipo material")) for r in material_rows)

    hist_unique = _unique_histories(views["hist_pcba_rows"])
    process = Counter(r.get("Situação reparo") or "" for r in hist_unique)
    types = Counter(str(r.get("Defect Type 2114") or "").strip() or "(vazio)" for r in hist_unique)
    families = Counter(r.get("Família 2114") or "OUTRO" for r in hist_unique if r.get("Defect Description 2114") != "SEM HISTÓRICO 2114")

    # Sinais investigativos de materiais reutilizados.
    material_signal = defaultdict(lambda: {"same":0,"similar":0,"different":0})
    for r in views["hist_material_rows"]:
        sn = r.get("Material SN")
        rel = r.get("Relação com falha atual")
        if rel == "MESMA FALHA": material_signal[sn]["same"] += 1
        elif rel == "MESMA FAMÍLIA": material_signal[sn]["similar"] += 1
        elif rel == "FALHA DIFERENTE": material_signal[sn]["different"] += 1
    mat_same = sum(1 for v in material_signal.values() if v["same"] > 0)
    mat_sim = sum(1 for v in material_signal.values() if v["same"] == 0 and v["similar"] > 0)
    screen_count = sum(v for k,v in mat_types.items() if "screen" in k.lower() or "lcd" in k.lower())
    camera_count = sum(v for k,v in mat_types.items() if "camera" in k.lower())

    # Repetição após reparo de placa concluído: sinal de atenção, não causalidade.
    attention_pcba = set()
    for r in views["hist_pcba_rows"]:
        if r.get("Relação com falha atual") in ("MESMA FALHA", "MESMA FAMÍLIA") and r.get("Situação reparo") == "REPARO DE PLACA CONCLUÍDO":
            attention_pcba.add(r.get("PCBA"))

    line_data = defaultdict(lambda: {"falhas":0,"pcbas":set(),"pcba2":set(),"pcba3":set(),"mat2":set(),"mat3":set(),"same":set(),"family":set()})
    for r in rows:
        l = r.get("Linha") or "SEM LINHA"
        d=line_data[l]; d["falhas"] += 1; d["pcbas"].add(r.get("PCBA"))
        uses=int(r.get("_pcba_use_index") or 0)
        if uses==2: d["pcba2"].add(r.get("PCBA"))
        if uses>=3: d["pcba3"].add(r.get("PCBA"))
        if int(r.get("Mesma falha anterior") or 0)>0: d["same"].add(r.get("PCBA"))
        if int(r.get("Mesma família na PCBA") or 0)>0: d["family"].add(r.get("PCBA"))
    for r in views["reusos"]:
        l=r.get("Linha") or "SEM LINHA"; d=line_data[l]; sn=r.get("Material SN")
        prev=int(r.get("Reusos anteriores a falha") or 0)
        if prev==1: d["mat2"].add(sn)
        if prev>=2: d["mat3"].add(sn)

    lines=[]
    for line,d in sorted(line_data.items()):
        lines.append({"Linha":line,"Falhas":d["falhas"],"PCBAs únicas":len(d["pcbas"]),"PCBA 2º uso":len(d["pcba2"]),"PCBA 3º+ uso":len(d["pcba3"]),"Materiais 2º uso":len(d["mat2"]),"Materiais 3º+ uso":len(d["mat3"]),"PCBAs mesma falha":len(d["same"]),"PCBAs mesma família":len(d["family"])})

    return {
        "kpis": [
            ("Falhas atuais 3028", len(rows)), ("PCBAs únicas com falha", len(pcba_rows)),
            ("PCBA em 1º uso", pcba_first), ("PCBA em 2º uso", pcba_second), ("PCBA em 3º+ uso", pcba_multi),
            ("PCBAs com 2+ históricos 2114", pcba_hist2), ("PCBAs com 3+ históricos 2114", pcba_hist3),
            ("PCBAs que repetiram a mesma falha", pcba_same), ("PCBAs com recorrência da mesma família", pcba_sim),
            ("Materiais únicos em 2º uso", mat_second), ("Materiais únicos em 3º+ uso", mat_multi),
            ("Screens/LCDs reutilizados", screen_count), ("Câmeras reutilizadas", camera_count),
            ("Materiais com mesma falha em PCBA desvinculada", mat_same), ("Materiais com mesma família em PCBA desvinculada", mat_sim),
            ("PCBAs com recorrência após MainBoard + Y", len(attention_pcba)),
        ],
        "lines": lines,
        "material_types": sorted(mat_types.items(), key=lambda x:(-x[1],x[0])),
        "repair_states": sorted(process.items(), key=lambda x:(-x[1],x[0])),
        "defect_types": sorted(types.items(), key=lambda x:(-x[1],x[0])),
        "families": sorted(families.items(), key=lambda x:(-x[1],x[0])),
    }


def _style_header(ws, row, col_start, col_end):
    for c in ws.iter_cols(min_col=col_start,max_col=col_end,min_row=row,max_row=row):
        cell=c[0]; cell.fill=PatternFill("solid",fgColor=HEADER_FILL); cell.font=Font(bold=True,color=HEADER_FONT); cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); cell.border=Border(bottom=LIGHT_BORDER)


def _write_table_sheet(wb, name, rows, headers, table_name, widths=None, empty_message=None):
    ws=wb.create_sheet(name)
    ws.append(headers)
    for r in rows: ws.append([r.get(h,"") for h in headers])
    _style_header(ws,1,1,len(headers))
    ws.freeze_panes=None; ws.sheet_view.pane=None; ws.sheet_view.zoomScale=85
    if rows:
        ref=f"A1:{get_column_letter(len(headers))}{len(rows)+1}"
        tab=Table(displayName=table_name,ref=ref); tab.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showFirstColumn=False,showLastColumn=False,showRowStripes=True,showColumnStripes=False); ws.add_table(tab)
        ws.auto_filter.ref=ref
    elif empty_message:
        ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=max(1,len(headers)))
        ws.cell(2,1).value=empty_message; ws.cell(2,1).fill=PatternFill("solid",fgColor=FILL_NONE); ws.cell(2,1).font=Font(bold=True,color="666666")
    for i,h in enumerate(headers,1):
        w=(widths or {}).get(h)
        if w is None:
            if h.startswith("Reuso "): w=27
            elif "Description" in h or "Desc" in h or "Comment" in h or "Motivo" in h or "Situação" in h: w=30
            elif "PCBA" in h or "Material SN" in h: w=27
            else: w=16
        ws.column_dimensions[get_column_letter(i)].width=w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment=Alignment(vertical="top",wrap_text=True); c.border=Border(bottom=LIGHT_BORDER)
    return ws


def _write_indicators(wb, ind):
    ws=wb.create_sheet("INDICADORES")
    ws.sheet_view.zoomScale=90; ws.freeze_panes=None; ws.sheet_view.pane=None
    ws["A1"]="A-MES — Indicadores de rastreabilidade e recorrência"; ws["A1"].font=Font(bold=True,size=16,color="FFFFFF"); ws["A1"].fill=PatternFill("solid",fgColor=HEADER_FILL); ws.merge_cells("A1:H1")
    ws["A2"]="Indicadores são sinais de investigação. Não atribuem causa automaticamente ao material/PCBA."; ws.merge_cells("A2:H2"); ws["A2"].font=Font(italic=True,color="666666")
    r=4
    ws.cell(r,1).value="Indicadores principais"; ws.cell(r,1).font=Font(bold=True,size=12); ws.cell(r,1).fill=PatternFill("solid",fgColor=SECTION_FILL); ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=4); r+=1
    for i,(label,value) in enumerate(ind["kpis"]):
        rr=r+i//2; cc=1+(i%2)*2
        ws.cell(rr,cc).value=label; ws.cell(rr,cc).font=Font(bold=True); ws.cell(rr,cc).fill=PatternFill("solid",fgColor=KPI_FILL)
        ws.cell(rr,cc+1).value=value; ws.cell(rr,cc+1).font=Font(bold=True,size=13); ws.cell(rr,cc+1).alignment=Alignment(horizontal="center")
    r += (len(ind["kpis"])+1)//2 + 2
    # por linha
    ws.cell(r,1).value="Por linha"; ws.cell(r,1).font=Font(bold=True,size=12); ws.cell(r,1).fill=PatternFill("solid",fgColor=SECTION_FILL); ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=8); r+=1
    headers=list(ind["lines"][0].keys()) if ind["lines"] else ["Linha","Falhas","PCBAs únicas","PCBA 2º uso","PCBA 3º+ uso","Materiais 2º uso","Materiais 3º+ uso","PCBAs mesma falha","PCBAs mesma família"]
    for c,h in enumerate(headers,1): ws.cell(r,c).value=h
    _style_header(ws,r,1,len(headers))
    for row in ind["lines"]:
        r+=1
        for c,h in enumerate(headers,1): ws.cell(r,c).value=row.get(h,"")
    r+=3
    # material types left, repair states right
    start=r
    ws.cell(r,1).value="Materiais reutilizados por tipo"; ws.cell(r,1).font=Font(bold=True,size=12); ws.cell(r,1).fill=PatternFill("solid",fgColor=SECTION_FILL); ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=2)
    ws.cell(r,4).value="Situação do reparo — históricos 2114"; ws.cell(r,4).font=Font(bold=True,size=12); ws.cell(r,4).fill=PatternFill("solid",fgColor=SECTION_FILL); ws.merge_cells(start_row=r,start_column=4,end_row=r,end_column=5)
    r+=1
    ws.cell(r,1).value="Tipo material"; ws.cell(r,2).value="Qtd únicos"; _style_header(ws,r,1,2)
    ws.cell(r,4).value="Situação"; ws.cell(r,5).value="Qtd"; _style_header(ws,r,4,5)
    maxn=max(len(ind["material_types"]),len(ind["repair_states"]),1)
    for i in range(maxn):
        rr=r+1+i
        if i<len(ind["material_types"]): ws.cell(rr,1).value=ind["material_types"][i][0]; ws.cell(rr,2).value=ind["material_types"][i][1]
        if i<len(ind["repair_states"]): ws.cell(rr,4).value=ind["repair_states"][i][0]; ws.cell(rr,5).value=ind["repair_states"][i][1]
    r=r+maxn+3
    ws.cell(r,1).value="Famílias de falha 2114"; ws.cell(r,1).font=Font(bold=True,size=12); ws.cell(r,1).fill=PatternFill("solid",fgColor=SECTION_FILL); ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=2)
    ws.cell(r,4).value="Defect Type (todos os valores preservados)"; ws.cell(r,4).font=Font(bold=True,size=12); ws.cell(r,4).fill=PatternFill("solid",fgColor=SECTION_FILL); ws.merge_cells(start_row=r,start_column=4,end_row=r,end_column=5)
    r+=1; ws.cell(r,1).value="Família"; ws.cell(r,2).value="Qtd"; _style_header(ws,r,1,2); ws.cell(r,4).value="Defect Type"; ws.cell(r,5).value="Qtd"; _style_header(ws,r,4,5)
    maxn=max(len(ind["families"]),len(ind["defect_types"]),1)
    for i in range(maxn):
        rr=r+1+i
        if i<len(ind["families"]): ws.cell(rr,1).value=ind["families"][i][0]; ws.cell(rr,2).value=ind["families"][i][1]
        if i<len(ind["defect_types"]): ws.cell(rr,4).value=ind["defect_types"][i][0]; ws.cell(rr,5).value=ind["defect_types"][i][1]
    for col,w in {"A":38,"B":14,"C":4,"D":48,"E":14,"F":4,"G":18,"H":18}.items(): ws.column_dimensions[col].width=w
    # navegação
    ws["G4"]="Abrir ANALISE ↗"; ws["G4"].hyperlink="#'ANALISE'!A1"; ws["G4"].style="Hyperlink"
    ws["G5"]="Histórico PCBA ↗"; ws["G5"].hyperlink="#'HIST_PCBA'!A1"; ws["G5"].style="Hyperlink"
    ws["G6"]="Histórico materiais ↗"; ws["G6"].hyperlink="#'HIST_MATERIAL'!A1"; ws["G6"].style="Hyperlink"
    return ws


def _color_relations(ws, headers, name):
    if name not in headers: return
    c=headers.index(name)+1
    for r in range(2,ws.max_row+1):
        cell=ws.cell(r,c); v=str(cell.value or "").upper()
        if "MESMA FALHA" in v or "RECORRÊNCIA DA MESMA FALHA" in v: color=FILL_SAME
        elif "MESMA FAMÍLIA" in v: color=FILL_SIMILAR
        elif "DIFERENTE" in v or "OUTRAS" in v: color=FILL_DIFF
        elif "SEM HIST" in v or "SEM REPET" in v: color=FILL_NONE
        else: continue
        cell.fill=PatternFill("solid",fgColor=color); cell.font=Font(bold=True)


def gerar_excel_integrado(resultado_3074, resultado_2114, output_path):
    views=_build_views(resultado_3074,resultado_2114); ind=_build_indicators(views)
    wb=Workbook(); wb.remove(wb.active)
    _write_indicators(wb,ind)

    headers_a=["Linha","Modelo","PCBA","Defect Time","Defect Code","Defect Desc","Repair Comment","Reason Type","Reason Desc",
               "Uso da PCBA na falha","Usos conhecidos da PCBA","Reusos anteriores da PCBA","Históricos da PCBA","Mesma falha anterior","Mesma família na PCBA","Falhas diferentes na PCBA","Situação da PCBA","Materiais na falha","Materiais 2º+ uso"] + [f"Reuso {i}" for i in range(1,views["max_reusos"]+1)]
    ws_a=_write_table_sheet(wb,"ANALISE",views["analise_rows"],headers_a,"T_ANALISE_V013",widths={"Linha":13,"Modelo":15,"PCBA":27,"Defect Time":20,"Defect Desc":32,"Repair Comment":30,"Uso da PCBA na falha":18,"Usos conhecidos da PCBA":16,"Históricos da PCBA":15,"Situação da PCBA":28,"Materiais 2º+ uso":15})
    _color_relations(ws_a,headers_a,"Situação da PCBA")

    headers_p=["Voltar","Linha","Modelo","PCBA","Uso PCBA na falha","Usos conhecidos PCBA","Reusos anteriores PCBA","Ciclos montagem PCBA 3074","Defect Time atual","Defect Code atual","Defect Desc atual","Repair Comment atual","Hist Seq","Defect Code 2114","Defect Description 2114","Defect Oper 2114","Registro falha 2114","Repair Status 2114","Defect Type 2114","Situação reparo","Família 2114","Relação com falha atual","Motivo da relação"]
    rows_p=[]
    for r in views["hist_pcba_rows"]:
        x={k:v for k,v in r.items() if not k.startswith("_")}; x["Voltar"]="← ANALISE"; rows_p.append(x)
    ws_p=_write_table_sheet(wb,"HIST_PCBA",rows_p,headers_p,"T_HIST_PCBA_V013",widths={"Voltar":12,"PCBA":27,"Ciclos montagem PCBA 3074":46,"Defect Desc atual":30,"Defect Description 2114":34,"Situação reparo":38,"Relação com falha atual":22,"Motivo da relação":28},empty_message="Nenhum histórico 2114 foi retornado.")
    _color_relations(ws_p,headers_p,"Relação com falha atual")

    headers_m=["Voltar","Linha","Modelo","PCBA atual","Defect Time atual","Defect Code atual","Defect Desc atual","Tipo material","Material SN","Uso material na falha","Reusos anteriores material","Usos conhecidos material","Reusos conhecidos material","PCBA desvinculada","Estado vínculo atual","Ordem de uso anterior","Hist Seq","Defect Code antigo","Defect Description antigo","Defect Oper antigo","Registro falha antigo","Repair Status antigo","Defect Type antigo","Situação reparo antiga","Família falha antiga","Relação com falha atual","Motivo da relação"]
    rows_m=[]
    for r in views["hist_material_rows"]:
        x={k:v for k,v in r.items() if not k.startswith("_")}; x["Voltar"]="← ANALISE"; rows_m.append(x)
    ws_m=_write_table_sheet(wb,"HIST_MATERIAL",rows_m,headers_m,"T_HIST_MATERIAL_V013",widths={"Voltar":12,"Material SN":29,"PCBA desvinculada":27,"Defect Desc atual":30,"Defect Description antigo":34,"Situação reparo antiga":38,"Relação com falha atual":22,"Motivo da relação":28},empty_message="Nenhum material de segundo uso/reuso foi encontrado neste recorte.")
    _color_relations(ws_m,headers_m,"Relação com falha atual")

    # Mapas de âncora para navegação interna.
    pcba_anchor={}; mat_anchor={}; failure_mat_anchor={}
    for idx,r in enumerate(views["hist_pcba_rows"],2): pcba_anchor.setdefault(r.get("_failure_key"),idx)
    for idx,r in enumerate(views["hist_material_rows"],2):
        mat_anchor.setdefault(r.get("_material_key"),idx); failure_mat_anchor.setdefault(r.get("_failure_key"),idx)
    # Voltar
    for idx,r in enumerate(views["hist_pcba_rows"],2): ws_p.cell(idx,1).hyperlink=f"#'ANALISE'!A{r.get('_analysis_row',2)}"; ws_p.cell(idx,1).style="Hyperlink"
    for idx,r in enumerate(views["hist_material_rows"],2): ws_m.cell(idx,1).hyperlink=f"#'ANALISE'!A{r.get('_analysis_row',2)}"; ws_m.cell(idx,1).style="Hyperlink"
    # Links na ANALISE.
    hmap={h:i+1 for i,h in enumerate(headers_a)}
    for rnum,row in enumerate(views["analise_rows"],2):
        key=row.get("_failure_key")
        if key in pcba_anchor:
            for h in ("Uso da PCBA na falha","Históricos da PCBA"):
                c=ws_a.cell(rnum,hmap[h]); c.hyperlink=f"#'HIST_PCBA'!A{pcba_anchor[key]}"; c.style="Hyperlink"
        if key in failure_mat_anchor:
            c=ws_a.cell(rnum,hmap["Materiais 2º+ uso"]); c.hyperlink=f"#'HIST_MATERIAL'!A{failure_mat_anchor[key]}"; c.style="Hyperlink"
        for i in range(1,views["max_reusos"]+1):
            mk=row.get(f"_reuso_{i}_key")
            if mk in mat_anchor:
                c=ws_a.cell(rnum,hmap[f"Reuso {i}"]); c.hyperlink=f"#'HIST_MATERIAL'!A{mat_anchor[mk]}"; c.style="Hyperlink"
        ws_a.row_dimensions[rnum].height=82 if int(row.get("Materiais 2º+ uso") or 0)>0 else 32

    wb.active=wb.sheetnames.index("INDICADORES")
    output_path=Path(output_path); output_path.parent.mkdir(parents=True,exist_ok=True); wb.save(output_path)
    return {
        "falhas":len(views["analise_rows"]), "pcbas_unicas":len({r.get('PCBA') for r in views['analise_rows']}),
        "materiais_reuso":len(views["reusos"]), "linhas_hist_pcba":len(views["hist_pcba_rows"]), "linhas_hist_material":len(views["hist_material_rows"]),
        "pcba_segundo_uso": next((v for k,v in ind['kpis'] if k=='PCBA em 2º uso'),0),
        "pcba_multiuso": next((v for k,v in ind['kpis'] if k=='PCBA em 3º+ uso'),0),
    }
