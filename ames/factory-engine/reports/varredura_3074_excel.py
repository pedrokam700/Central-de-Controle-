from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter


def gerar_excel_varredura(resultado, output_path):
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    summary = dict(resultado.get("summary", {}))
    summary["pcba_query"] = resultado.get("pcba_query")
    summary["scanned_at"] = resultado.get("scanned_at")
    summary["pcba_from_cache"] = resultado.get("pcba_from_cache")

    rows = []
    for r in resultado.get("results", []):
        rows.append({
            "Item SN": r.get("item_sn"),
            "Tipo": r.get("item_type"),
            "Status": r.get("status"),
            "Qtd PCBAs": r.get("distinct_pcba_count"),
            "PCBAs": " | ".join(r.get("pcbas", [])),
            "Cache": "SIM" if r.get("from_cache") else "NAO",
            "Erro": r.get("error") or "",
        })

    df = pd.DataFrame(rows)
    df_reuso = df[df["Status"].isin(["SEGUNDO_USO", "REUTILIZACAO_MULTIPLA"])] if not df.empty else df
    df_erros = df[df["Status"] == "ERRO_CONSULTA"] if not df.empty else df

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        pd.DataFrame([summary]).to_excel(writer, sheet_name="Resumo", index=False)
        df.to_excel(writer, sheet_name="Componentes", index=False)
        df_reuso.to_excel(writer, sheet_name="Reutilizados", index=False)
        df_erros.to_excel(writer, sheet_name="Erros", index=False)

    wb = load_workbook(output_path)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        if ws.max_row >= 1 and ws.max_column >= 1:
            ws.auto_filter.ref = ws.dimensions
        for c in ws[1]:
            c.font = Font(bold=True)
            c.alignment = Alignment(horizontal="center", vertical="center")
        for col_idx in range(1, ws.max_column + 1):
            max_len = 0
            for row_idx in range(1, min(ws.max_row, 500) + 1):
                v = ws.cell(row_idx, col_idx).value
                max_len = max(max_len, len("" if v is None else str(v)))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 55)
    wb.save(output_path)
    return output_path
