from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter

from reports.relatorio_3028_3074 import extrair_linhas
from reports.relatorio_integrado_2114 import _pcba_result_map, _histories, _material_name, _split_pipe
from core.correlacao_falhas import analisar_historico_pcba_atual, comparar_falhas, familia_falha, _hist_seq_value
from core.uso_pcba import reconstruir_usos_pcba

HEADER_FILL = "1F4E78"
HEADER_FONT = "FFFFFF"
SECTION_FILL = "D9EAF7"
KPI_FILL = "EAF2F8"
FILL_SAME = "E2F0D9"
FILL_SIMILAR = "FFF2CC"
FILL_DIFF = "DDEBF7"
FILL_NONE = "E7E6E6"
FILL_ALERT = "FCE4D6"
FILL_SUMMARY = "EAF2F8"
FILL_GROUP = "F5F9FC"
FILL_NOFAIL = "E2F0D9"
LIGHT_BORDER = Side(style="thin", color="D9E1F2")


def _join(values):
    return " | ".join(str(x) for x in (values or []) if x not in (None, ""))


def _failure_key(pcba, defect_time, defect_code):
    return (str(pcba or "").strip(), str(defect_time or "").strip(), str(defect_code or "").strip())


def _relation_label(v):
    return {
        "MESMA_FALHA": "MESMA FALHA",
        "FALHA_SIMILAR": "MESMA FAMÍLIA",
        "FALHA_DIFERENTE": "FALHA DIFERENTE",
        "SEM_HISTORICO": "SEM HISTÓRICO 2114",
        "SEM_HISTORICO_2114": "SEM HISTÓRICO 2114",
    }.get(v, str(v or ""))


def _own_status(a):
    if a.get("same_repetitions", 0) > 0:
        return "RECORRÊNCIA DA MESMA FALHA"
    if a.get("prior_similar_count", 0) > 0:
        return "RECORRÊNCIA DA MESMA FAMÍLIA"
    if a.get("previous_total", 0) > 0:
        return "HISTÓRICO ANTERIOR COM OUTRAS FALHAS"
    if a.get("total", 0) > 0:
        return "SEM REPETIÇÃO ANTERIOR IDENTIFICADA"
    return "SEM HISTÓRICO 2114"


def _process_state(repair_status, defect_type):
    rs = str(repair_status or "").strip().upper()
    dt_raw = str(defect_type or "").strip()
    dt = dt_raw.lower().replace(" ", "_")
    if rs == "N" and not dt_raw:
        return "AGUARDANDO ANÁLISE DO AUXILIAR"
    if dt in ("mainboard", "main_board"):
        return "REPARO DE PLACA CONCLUÍDO" if rs == "Y" else "ENCAMINHADO / AGUARDANDO REPARO DE PLACA"
    if dt in ("phone_disassembly", "phonedisassembly"):
        return "PHONE DISASSEMBLY CONCLUÍDO" if rs == "Y" else "PHONE DISASSEMBLY PENDENTE"
    if dt_raw:
        return ("FINALIZADO" if rs == "Y" else "PENDENTE / EM FLUXO") + f" — {dt_raw}"
    if rs == "Y":
        return "FINALIZADO — TIPO NÃO INFORMADO"
    return "STATUS NÃO CLASSIFICADO"


def _manual_auto(hist):
    return str(hist.get("TestTools Auto Defect Or Defect By Hand") or "").strip()


def _pcba_item_map(resultado_3074):
    return {str(x.get("pcba_sn") or "").strip(): x for x in (resultado_3074 or {}).get("items", []) or [] if x.get("pcba_sn")}


def _reuse_analysis(row, pcba_map):
    old_pcbas = list(dict.fromkeys(_split_pipe(row.get("PCBAs antigas"))))
    inactive = set(_split_pipe(row.get("PCBAs inativas agora")))
    active = set(_split_pipe(row.get("PCBAs atuais Y")))
    details = []
    pcba_summaries = []

    for use_index, pcba in enumerate(old_pcbas, 1):
        state = "DESVINCULADA" if pcba in inactive else ("VINCULADA AGORA" if pcba in active else "HISTÓRICA")
        histories = _histories(pcba_map, pcba)
        if not histories:
            summary = {
                "pcba": pcba, "use_index": use_index, "state": state,
                "failure_count": 0, "relation": "SEM_HISTORICO_2114",
                "same_occurrences": 0, "similar_occurrences": 0, "different_occurrences": 0,
            }
            pcba_summaries.append(summary)
            details.append({
                "pcba": pcba, "use_index": use_index, "state": state, "history": None,
                "relation": "SEM_HISTORICO_2114", "reason": "PCBA sem histórico retornado na 2114",
                "pcba_relation": "SEM_HISTORICO_2114", "pcba_failure_count": 0,
            })
            continue

        local = []
        for hist in histories:
            cmp = comparar_falhas(row.get("Defect Code", ""), row.get("Defect Desc", ""),
                                  hist.get("Defect Code", ""), hist.get("Defect Description", ""))
            local.append((hist, cmp))

        same_occ = sum(1 for _, cmp in local if cmp.get("relation") == "MESMA_FALHA")
        sim_occ = sum(1 for _, cmp in local if cmp.get("relation") == "FALHA_SIMILAR")
        diff_occ = sum(1 for _, cmp in local if cmp.get("relation") == "FALHA_DIFERENTE")
        strongest = "MESMA_FALHA" if same_occ else ("FALHA_SIMILAR" if sim_occ else ("FALHA_DIFERENTE" if diff_occ else "SEM_HISTORICO_2114"))
        summary = {
            "pcba": pcba, "use_index": use_index, "state": state,
            "failure_count": len(histories), "relation": strongest,
            "same_occurrences": same_occ, "similar_occurrences": sim_occ, "different_occurrences": diff_occ,
        }
        pcba_summaries.append(summary)
        for hist, cmp in local:
            details.append({
                "pcba": pcba, "use_index": use_index, "state": state, "history": hist, **cmp,
                "pcba_relation": strongest, "pcba_failure_count": len(histories),
            })

    same_pcbas = sum(1 for x in pcba_summaries if x.get("relation") == "MESMA_FALHA")
    similar_pcbas = sum(1 for x in pcba_summaries if x.get("relation") == "FALHA_SIMILAR")
    different_pcbas = sum(1 for x in pcba_summaries if x.get("relation") == "FALHA_DIFERENTE")
    no_hist_pcbas = sum(1 for x in pcba_summaries if x.get("relation") == "SEM_HISTORICO_2114")
    same_occ = sum(x.get("same_occurrences", 0) for x in pcba_summaries)
    similar_occ = sum(x.get("similar_occurrences", 0) for x in pcba_summaries)
    different_occ = sum(x.get("different_occurrences", 0) for x in pcba_summaries)
    hist_occurrences = sum(x.get("failure_count", 0) for x in pcba_summaries)
    strongest = "MESMA_FALHA" if same_pcbas else ("FALHA_SIMILAR" if similar_pcbas else ("FALHA_DIFERENTE" if different_pcbas else "SEM_HISTORICO_2114"))
    return {
        "old_pcbas": old_pcbas, "details": details, "pcba_summaries": pcba_summaries,
        "old_pcba_count": len(old_pcbas),
        "same_pcbas": same_pcbas, "similar_pcbas": similar_pcbas,
        "different_pcbas": different_pcbas, "no_history_pcbas": no_hist_pcbas,
        "same_occurrences": same_occ, "similar_occurrences": similar_occ,
        "different_occurrences": different_occ, "history_occurrences": hist_occurrences,
        # Compatibilidade com nomes V0.13.
        "same": same_occ, "similar": similar_occ, "different": different_occ,
        "strongest": strongest,
    }



def _material_evidence(ra):
    """Classifica apenas a EVIDÊNCIA histórica disponível.

    Não afirma que o material causou a falha atual. A distinção principal é:
    - reuso sem falha anterior registrada nas PCBAs desvinculadas;
    - reuso com histórico de falha, separado por mesma falha / mesma família /
      outras falhas.
    """
    if int(ra.get("history_occurrences") or 0) <= 0:
        return "REUSO SEM FALHA ANTERIOR REGISTRADA"
    if int(ra.get("same_pcbas") or 0) > 0:
        return "HISTÓRICO COM A MESMA FALHA"
    if int(ra.get("similar_pcbas") or 0) > 0:
        return "HISTÓRICO DA MESMA FAMÍLIA"
    if int(ra.get("different_pcbas") or 0) > 0:
        return "HISTÓRICO COM OUTRAS FALHAS"
    return "HISTÓRICO 2114 NÃO CLASSIFICADO"


def _compact_reuse_card(r, ra):
    prev = int(r.get("Reusos anteriores a falha") or 0)
    use = prev + 1
    evidence = _material_evidence(ra)
    return (
        f"{_material_name(r.get('Tipo material'))}\n"
        f"SN: {r.get('Material SN') or '-'}\n"
        f"Uso na falha: {use}º | Usos conhecidos: {r.get('Usos totais conhecidos') or '-'}\n"
        f"PCBAs desvinculadas: {ra['old_pcba_count']} | Falhas nelas: {ra['history_occurrences']}\n"
        f"Evidência: {evidence}"
    )

def _usage_cycles_summary(usage):
    parts=[]
    for c in usage.get("cycles", []) or []:
        start=str(c.get("start_utc") or "").replace("+00:00", " UTC")
        end=str(c.get("end_utc") or "").replace("+00:00", " UTC") or "ATIVO"
        parts.append(f"{c.get('use_index')}º: {start} -> {end} [{c.get('link_state','')}]" )
    return "\n".join(parts)


def _hist_signature(hist):
    if not hist:
        return None
    return (
        str(hist.get("Defect Hist Seq") or ""),
        str(hist.get("Defect Code") or ""),
        str(hist.get("Defect Description") or ""),
        str(hist.get("Repair Status") or ""),
        str(hist.get("Defect Type") or ""),
    )


def _build_views(resultado_3074, resultado_2114):
    falhas, materiais, reusos = extrair_linhas(resultado_3074)
    pcba_map = _pcba_result_map(resultado_2114)
    item_map = _pcba_item_map(resultado_3074)

    reuso_by_key = defaultdict(list)
    for r in reusos:
        reuso_by_key[_failure_key(r.get("PCBA da falha"), r.get("Defect Time"), r.get("Defect Code"))].append(r)

    analise_rows, hist_pcba_rows, hist_material_rows, material_summary_rows = [], [], [], []
    max_reusos = 0
    pcba_usage_by_context = {}

    for analise_idx, f in enumerate(falhas, 1):
        pcba = str(f.get("PCBA") or "").strip()
        defect_time = f.get("Defect Time", "")
        defect_code = f.get("Defect Code", "")
        key = _failure_key(pcba, defect_time, defect_code)
        histories = _histories(pcba_map, pcba)
        own = analisar_historico_pcba_atual(defect_code, f.get("Defect Desc", ""), histories)
        current_hist = own.get("current_history") or {}
        current_sig = _hist_signature(current_hist)
        current_seq_num = _hist_seq_value(current_hist, -1) if current_hist else None

        base_rows = ((item_map.get(pcba) or {}).get("base_query") or {}).get("rows") or []
        usage = reconstruir_usos_pcba(base_rows, pcba, defect_time)
        pcba_usage_by_context[key] = usage

        related = sorted(reuso_by_key.get(key, []), key=lambda x:(str(x.get("Tipo material") or ""), str(x.get("Material SN") or "")))
        max_reusos = max(max_reusos, len(related))

        current_process = _process_state(current_hist.get("Repair Status", ""), current_hist.get("Defect Type", "")) if current_hist else "SEM HISTÓRICO 2114"
        row = {
            "Linha": f.get("Linha", ""), "Modelo": f.get("Modelo", ""), "PCBA": pcba,
            "Defect Time": defect_time, "Defect Code": defect_code, "Defect Desc": f.get("Defect Desc", ""),
            "Repair Comment": f.get("Repair Comment", ""), "Reason Type": f.get("Reason Type", ""), "Reason Desc": f.get("Reason Desc", ""),
            "Uso da PCBA na falha": f"{usage.get('use_index_at_defect') or 0}º uso" if usage.get("use_index_at_defect") else "NÃO DETERMINADO",
            "Usos conhecidos da PCBA": usage.get("known_use_count", 0),
            "Reusos anteriores da PCBA": usage.get("reuses_before_defect", 0),
            "_pcba_use_index": int(usage.get("use_index_at_defect") or 0),
            "Falhas registradas 2114": own.get("total", 0),
            "Falhas anteriores 2114": own.get("previous_total", 0),
            "Mesma falha anterior": own.get("same_repetitions", 0),
            "Mesma família anterior": own.get("prior_similar_count", 0),
            "Falhas diferentes anteriores": own.get("prior_different_count", 0),
            "Situação da PCBA": _own_status(own),
            "Repair Status atual": current_hist.get("Repair Status", ""),
            "Defect Type atual": current_hist.get("Defect Type", ""),
            "Situação reparo atual": current_process,
            "Registro falha atual": _manual_auto(current_hist),
            "Hist Seq atual": current_hist.get("Defect Hist Seq", ""),
            "Materiais na falha": f.get("Materiais na falha", 0),
            "Materiais 2º+ uso": len(related),
            "_analysis_row": analise_idx + 1,
            "_failure_key": key,
        }

        for idx, r in enumerate(related, 1):
            ra = _reuse_analysis(r, pcba_map)
            row[f"Reuso {idx}"] = _compact_reuse_card(r, ra)
            row[f"_reuso_{idx}_sn"] = r.get("Material SN", "")
            row[f"_reuso_{idx}_key"] = (key, r.get("Material SN", ""))

            prev = int(r.get("Reusos anteriores a falha") or 0)
            use_no = prev + 1
            material_summary_rows.append({
                "_failure_key": key, "_material_key": (key, r.get("Material SN", "")), "_analysis_row": analise_idx + 1,
                "Linha": r.get("Linha", ""), "Modelo": r.get("Modelo", ""), "PCBA atual": r.get("PCBA da falha", ""),
                "Defect Time atual": r.get("Defect Time", ""), "Defect Code atual": r.get("Defect Code", ""), "Defect Desc atual": r.get("Defect Desc", ""),
                "Tipo material": _material_name(r.get("Tipo material", "")), "Material SN": r.get("Material SN", ""),
                "Uso material na falha": f"{use_no}º uso", "Reusos anteriores material": prev,
                "Usos conhecidos material": r.get("Usos totais conhecidos"), "Reusos conhecidos material": r.get("Reusos totais conhecidos"),
                "PCBAs desvinculadas únicas": ra.get("old_pcba_count", 0),
                "PCBAs com mesma falha": ra.get("same_pcbas", 0), "PCBAs com mesma família": ra.get("similar_pcbas", 0),
                "PCBAs com falha diferente": ra.get("different_pcbas", 0), "PCBAs sem histórico de falha": ra.get("no_history_pcbas", 0),
                "Falhas 2114 nas desvinculadas": ra.get("history_occurrences", 0),
                "Evidência histórica": _material_evidence(ra),
                "_pcba_summaries": ra.get("pcba_summaries", []),
            })
            for detail in ra["details"]:
                hist = detail.get("history") or {}
                hist_material_rows.append({
                    "_failure_key": key, "_material_key": (key, r.get("Material SN", "")), "_analysis_row": analise_idx + 1,
                    "Linha": r.get("Linha", ""), "Modelo": r.get("Modelo", ""), "PCBA atual": r.get("PCBA da falha", ""),
                    "Defect Time atual": r.get("Defect Time", ""), "Defect Code atual": r.get("Defect Code", ""), "Defect Desc atual": r.get("Defect Desc", ""),
                    "Tipo material": _material_name(r.get("Tipo material", "")), "Material SN": r.get("Material SN", ""),
                    "Uso material na falha": f"{use_no}º uso", "Reusos anteriores material": prev,
                    "Usos conhecidos material": r.get("Usos totais conhecidos"), "Reusos conhecidos material": r.get("Reusos totais conhecidos"),
                    "PCBAs desvinculadas únicas": ra.get("old_pcba_count", 0),
                    "PCBAs com mesma falha": ra.get("same_pcbas", 0),
                    "PCBAs com mesma família": ra.get("similar_pcbas", 0),
                    "PCBAs com falha diferente": ra.get("different_pcbas", 0),
                    "Falhas 2114 nas desvinculadas": ra.get("history_occurrences", 0),
                    "Evidência histórica": _material_evidence(ra),
                    "PCBA desvinculada": detail.get("pcba", ""), "Estado vínculo atual": detail.get("state", ""), "Ordem de uso anterior": detail.get("use_index", ""),
                    "Falhas 2114 desta PCBA": detail.get("pcba_failure_count", 0),
                    "Classificação desta PCBA": _relation_label(detail.get("pcba_relation")),
                    "Hist Seq": hist.get("Defect Hist Seq", ""), "Defect Code antigo": hist.get("Defect Code", ""),
                    "Defect Description antigo": hist.get("Defect Description", "") or ("SEM HISTÓRICO 2114" if not hist else ""),
                    "Defect Oper antigo": hist.get("Defect Oper", ""), "Registro falha antigo": _manual_auto(hist),
                    "Repair Status antigo": hist.get("Repair Status", ""), "Defect Type antigo": hist.get("Defect Type", ""),
                    "Situação reparo antiga": _process_state(hist.get("Repair Status", ""), hist.get("Defect Type", "")),
                    "Família falha antiga": familia_falha(hist.get("Defect Description", ""), hist.get("Defect Code", "")) if hist else "",
                    "Relação desta ocorrência": _relation_label(detail.get("relation")), "Motivo da relação": detail.get("reason", ""),
                })

        analise_rows.append(row)

        if not histories:
            hist_pcba_rows.append({
                "_failure_key": key, "_analysis_row": analise_idx + 1,
                "Linha": f.get("Linha", ""), "Modelo": f.get("Modelo", ""), "PCBA": pcba,
                "Uso PCBA na falha": row["Uso da PCBA na falha"], "Usos conhecidos PCBA": usage.get("known_use_count",0),
                "Reusos anteriores PCBA": usage.get("reuses_before_defect",0), "Ciclos de uso PCBA 3074": _usage_cycles_summary(usage),
                "Falhas registradas 2114": 0, "Falhas anteriores 2114": 0,
                "Defect Time atual": defect_time, "Defect Code atual": defect_code, "Defect Desc atual": f.get("Defect Desc", ""), "Repair Comment atual": f.get("Repair Comment", ""),
                "Ocorrência atual?": "", "Posição no histórico": "", "Hist Seq": "", "Defect Code 2114": "", "Defect Description 2114": "SEM HISTÓRICO 2114",
                "Defect Oper 2114": "", "Registro falha 2114": "", "Repair Status 2114": "", "Defect Type 2114": "",
                "Situação reparo": "SEM HISTÓRICO 2114", "Família 2114": "", "Relação com falha atual": "SEM HISTÓRICO 2114", "Motivo da relação": "PCBA sem histórico retornado na 2114",
            })
        else:
            for hist_idx, hist in enumerate(histories):
                is_current = _hist_signature(hist) == current_sig
                seq_num = _hist_seq_value(hist, hist_idx)
                if is_current:
                    position = "ATUAL"
                    rel = "OCORRÊNCIA ATUAL"
                    reason = own.get("selection_method", "")
                else:
                    position = "ANTERIOR" if current_seq_num is None or seq_num < current_seq_num else "POSTERIOR"
                    cmp = comparar_falhas(defect_code, f.get("Defect Desc", ""), hist.get("Defect Code", ""), hist.get("Defect Description", ""))
                    rel = _relation_label(cmp.get("relation"))
                    reason = cmp.get("reason", "")
                hist_pcba_rows.append({
                    "_failure_key": key, "_analysis_row": analise_idx + 1,
                    "Linha": f.get("Linha", ""), "Modelo": f.get("Modelo", ""), "PCBA": pcba,
                    "Uso PCBA na falha": row["Uso da PCBA na falha"], "Usos conhecidos PCBA": usage.get("known_use_count",0),
                    "Reusos anteriores PCBA": usage.get("reuses_before_defect",0), "Ciclos de uso PCBA 3074": _usage_cycles_summary(usage),
                    "Falhas registradas 2114": own.get("total", 0), "Falhas anteriores 2114": own.get("previous_total", 0),
                    "Defect Time atual": defect_time, "Defect Code atual": defect_code, "Defect Desc atual": f.get("Defect Desc", ""), "Repair Comment atual": f.get("Repair Comment", ""),
                    "Ocorrência atual?": "SIM" if is_current else "NÃO", "Posição no histórico": position,
                    "Hist Seq": hist.get("Defect Hist Seq", ""), "Defect Code 2114": hist.get("Defect Code", ""),
                    "Defect Description 2114": hist.get("Defect Description", ""), "Defect Oper 2114": hist.get("Defect Oper", ""),
                    "Registro falha 2114": _manual_auto(hist), "Repair Status 2114": hist.get("Repair Status", ""), "Defect Type 2114": hist.get("Defect Type", ""),
                    "Situação reparo": _process_state(hist.get("Repair Status", ""), hist.get("Defect Type", "")),
                    "Família 2114": familia_falha(hist.get("Defect Description", ""), hist.get("Defect Code", "")),
                    "Relação com falha atual": rel, "Motivo da relação": reason,
                })

    if len(analise_rows) != len(falhas):
        raise RuntimeError("ANALISE perdeu falhas durante a consolidação.")
    expected_hist = 0
    for f in falhas:
        hs = _histories(pcba_map, f.get("PCBA"))
        expected_hist += max(1, len(hs))
    if len(hist_pcba_rows) != expected_hist:
        raise RuntimeError(f"HIST_PCBA inconsistente: esperado {expected_hist}, obtido {len(hist_pcba_rows)}.")
    if reusos and not hist_material_rows:
        raise RuntimeError("Existem materiais reutilizados, mas HIST_MATERIAL ficou vazio.")

    return {
        "falhas": falhas, "materiais": materiais, "reusos": reusos,
        "analise_rows": analise_rows, "hist_pcba_rows": hist_pcba_rows, "hist_material_rows": hist_material_rows,
        "material_summary_rows": material_summary_rows,
        "max_reusos": max_reusos, "pcba_usage": pcba_usage_by_context,
    }


def _unique_histories(hist_rows):
    seen, out = set(), []
    for r in hist_rows:
        key = (r.get("PCBA"), r.get("Hist Seq"), r.get("Defect Code 2114"), r.get("Defect Description 2114"), r.get("Defect Type 2114"), r.get("Repair Status 2114"))
        if key in seen: continue
        seen.add(key); out.append(r)
    return out


def _build_indicators(views):
    rows = views["analise_rows"]
    unique_pcba = {}
    for r in rows:
        pcba = r.get("PCBA")
        if pcba not in unique_pcba or str(r.get("Defect Time")) > str(unique_pcba[pcba].get("Defect Time")):
            unique_pcba[pcba] = r
    pcba_rows = list(unique_pcba.values())

    pcba_first = sum(1 for r in pcba_rows if int(r.get("_pcba_use_index") or 0) <= 1)
    pcba_second = sum(1 for r in pcba_rows if int(r.get("_pcba_use_index") or 0) == 2)
    pcba_multi = sum(1 for r in pcba_rows if int(r.get("_pcba_use_index") or 0) >= 3)
    pcba_fail2 = sum(1 for r in pcba_rows if int(r.get("Falhas registradas 2114") or 0) >= 2)
    pcba_fail3 = sum(1 for r in pcba_rows if int(r.get("Falhas registradas 2114") or 0) >= 3)
    pcba_same = sum(1 for r in pcba_rows if int(r.get("Mesma falha anterior") or 0) > 0)
    pcba_sim = sum(1 for r in pcba_rows if int(r.get("Mesma família anterior") or 0) > 0)
    # Muito importante: dois usos não significam duas falhas. Este KPI evidencia
    # PCBA reutilizada que possui apenas a falha atual registrada na 2114.
    pcba_reused_one_failure = sum(
        1 for r in pcba_rows
        if int(r.get("_pcba_use_index") or 0) >= 2 and int(r.get("Falhas registradas 2114") or 0) == 1
    )
    pcba_reused_with_prior_failure = sum(
        1 for r in pcba_rows
        if int(r.get("_pcba_use_index") or 0) >= 2 and int(r.get("Falhas anteriores 2114") or 0) > 0
    )

    # Todos os componentes do lote, sempre por Material SN único.
    all_materials = {}
    for r in views["materiais"]:
        sn = str(r.get("Material SN") or "").strip()
        if sn:
            all_materials.setdefault(sn, r)

    reused_materials = {}
    for r in views["reusos"]:
        sn = str(r.get("Material SN") or "").strip()
        if sn:
            reused_materials.setdefault(sn, r)

    # Consolida evidência histórica por Material SN. Um mesmo material pode
    # aparecer em mais de um contexto; sempre usamos a evidência mais forte.
    summary_by_sn = defaultdict(lambda: {
        "history": 0, "same": 0, "similar": 0, "different": 0, "no_history_pcbas": 0,
    })
    for s in views.get("material_summary_rows", []):
        sn = str(s.get("Material SN") or "").strip()
        if not sn:
            continue
        d = summary_by_sn[sn]
        d["history"] = max(d["history"], int(s.get("Falhas 2114 nas desvinculadas") or 0))
        d["same"] = max(d["same"], int(s.get("PCBAs com mesma falha") or 0))
        d["similar"] = max(d["similar"], int(s.get("PCBAs com mesma família") or 0))
        d["different"] = max(d["different"], int(s.get("PCBAs com falha diferente") or 0))
        d["no_history_pcbas"] = max(d["no_history_pcbas"], int(s.get("PCBAs sem histórico de falha") or 0))

    by_type = defaultdict(lambda: {
        "all": set(), "reuse": set(), "second": set(), "multi": set(),
        "with_old_failure": set(), "without_old_failure": set(), "same": set(), "similar": set(),
    })
    for sn, r in all_materials.items():
        by_type[_material_name(r.get("Tipo material"))]["all"].add(sn)
    for sn, r in reused_materials.items():
        d = by_type[_material_name(r.get("Tipo material"))]
        d["reuse"].add(sn)
        prev = int(r.get("Reusos anteriores a falha") or 0)
        if prev == 1:
            d["second"].add(sn)
        elif prev >= 2:
            d["multi"].add(sn)
        sig = summary_by_sn.get(sn, {})
        if int(sig.get("history") or 0) > 0:
            d["with_old_failure"].add(sn)
        else:
            d["without_old_failure"].add(sn)
        if int(sig.get("same") or 0) > 0:
            d["same"].add(sn)
        elif int(sig.get("similar") or 0) > 0:
            d["similar"].add(sn)

    component_types = []
    for name, d in sorted(by_type.items(), key=lambda kv: (-len(kv[1]["reuse"]), kv[0])):
        total = len(d["all"]); reuse = len(d["reuse"])
        component_types.append({
            "Tipo material": name,
            "Materiais únicos no lote": total,
            "Reutilizados únicos": reuse,
            "2º uso": len(d["second"]),
            "3º+ uso": len(d["multi"]),
            "Com falha nas PCBAs desvinculadas": len(d["with_old_failure"]),
            "Sem falha nas PCBAs desvinculadas": len(d["without_old_failure"]),
            "Com mesma falha": len(d["same"]),
            "Com mesma família": len(d["similar"]),
            "% reuso": round((reuse / total * 100.0), 1) if total else 0,
        })

    mat_second = sum(1 for r in reused_materials.values() if int(r.get("Reusos anteriores a falha") or 0) == 1)
    mat_multi = sum(1 for r in reused_materials.values() if int(r.get("Reusos anteriores a falha") or 0) >= 2)
    mat_with_old_failure = sum(1 for sn in reused_materials if int(summary_by_sn.get(sn, {}).get("history") or 0) > 0)
    mat_without_old_failure = len(reused_materials) - mat_with_old_failure
    mat_same = sum(1 for sn in reused_materials if int(summary_by_sn.get(sn, {}).get("same") or 0) > 0)
    mat_sim = sum(1 for sn in reused_materials if int(summary_by_sn.get(sn, {}).get("same") or 0) == 0 and int(summary_by_sn.get(sn, {}).get("similar") or 0) > 0)
    mat_diff_only = sum(
        1 for sn in reused_materials
        if int(summary_by_sn.get(sn, {}).get("history") or 0) > 0
        and int(summary_by_sn.get(sn, {}).get("same") or 0) == 0
        and int(summary_by_sn.get(sn, {}).get("similar") or 0) == 0
    )

    screen_count = sum(x["Reutilizados únicos"] for x in component_types if "screen" in x["Tipo material"].lower() or "lcd" in x["Tipo material"].lower())
    camera_count = sum(x["Reutilizados únicos"] for x in component_types if "camera" in x["Tipo material"].lower())

    # Relação por vínculo material↔PCBA, nunca por quantidade de linhas de falha.
    link_relation = {}
    for r in views["hist_material_rows"]:
        sn = r.get("Material SN"); pcba = r.get("PCBA desvinculada"); cls = r.get("Classificação desta PCBA")
        if sn and pcba:
            link_relation[(sn, pcba)] = cls
    links_same = sum(1 for v in link_relation.values() if v == "MESMA FALHA")
    links_sim = sum(1 for v in link_relation.values() if v == "MESMA FAMÍLIA")

    hist_unique = _unique_histories(views["hist_pcba_rows"])
    types = Counter(str(r.get("Defect Type 2114") or "").strip() or "(vazio)" for r in hist_unique)
    families = Counter(r.get("Família 2114") or "OUTRO" for r in hist_unique if r.get("Defect Description 2114") != "SEM HISTÓRICO 2114")
    current_process = Counter((r.get("Situação reparo atual") or "SEM HISTÓRICO 2114") for r in pcba_rows)

    attention_pcba = set()
    for r in views["hist_pcba_rows"]:
        if r.get("Posição no histórico") != "ANTERIOR":
            continue
        if r.get("Relação com falha atual") in ("MESMA FALHA", "MESMA FAMÍLIA") and r.get("Situação reparo") == "REPARO DE PLACA CONCLUÍDO":
            attention_pcba.add(r.get("PCBA"))

    line_data = defaultdict(lambda: {
        "falhas":0,"pcbas":set(),"pcba2":set(),"pcba3":set(),"pcba_one_failure":set(),
        "mat2":set(),"mat3":set(),"mat_old_failure":set(),"mat_no_old_failure":set(),"same":set(),"family":set()
    })
    for r in rows:
        line = r.get("Linha") or "SEM LINHA"
        d = line_data[line]; d["falhas"] += 1; d["pcbas"].add(r.get("PCBA"))
        uses = int(r.get("_pcba_use_index") or 0)
        if uses == 2: d["pcba2"].add(r.get("PCBA"))
        if uses >= 3: d["pcba3"].add(r.get("PCBA"))
        if uses >= 2 and int(r.get("Falhas registradas 2114") or 0) == 1: d["pcba_one_failure"].add(r.get("PCBA"))
        if int(r.get("Mesma falha anterior") or 0) > 0: d["same"].add(r.get("PCBA"))
        if int(r.get("Mesma família anterior") or 0) > 0: d["family"].add(r.get("PCBA"))
    for sn, r in reused_materials.items():
        line = r.get("Linha") or "SEM LINHA"; d = line_data[line]
        prev = int(r.get("Reusos anteriores a falha") or 0)
        if prev == 1: d["mat2"].add(sn)
        if prev >= 2: d["mat3"].add(sn)
        if int(summary_by_sn.get(sn, {}).get("history") or 0) > 0: d["mat_old_failure"].add(sn)
        else: d["mat_no_old_failure"].add(sn)

    lines = []
    for line, d in sorted(line_data.items()):
        lines.append({
            "Linha": line, "Falhas": d["falhas"], "PCBAs únicas": len(d["pcbas"]),
            "PCBA 2º uso": len(d["pcba2"]), "PCBA 3º+ uso": len(d["pcba3"]),
            "PCBA 2º+ uso / só 1 falha": len(d["pcba_one_failure"]),
            "Materiais 2º uso": len(d["mat2"]), "Materiais 3º+ uso": len(d["mat3"]),
            "Materiais com falha anterior": len(d["mat_old_failure"]), "Materiais sem falha anterior": len(d["mat_no_old_failure"]),
            "PCBAs mesma falha": len(d["same"]), "PCBAs mesma família": len(d["family"]),
        })

    return {
        "overview_kpis": [
            ("Falhas atuais 3028", len(rows)), ("PCBAs únicas com falha", len(pcba_rows)),
            ("Linhas no lote", len(line_data)), ("Materiais únicos presentes", len(all_materials)),
        ],
        "pcba_kpis": [
            ("PCBA em 1º uso", pcba_first), ("PCBA em 2º uso", pcba_second), ("PCBA em 3º+ uso", pcba_multi),
            ("PCBAs 2º+ uso com só 1 falha 2114", pcba_reused_one_failure),
            ("PCBAs 2º+ uso com falha anterior", pcba_reused_with_prior_failure),
            ("PCBAs com 2+ falhas 2114", pcba_fail2), ("PCBAs com 3+ falhas 2114", pcba_fail3),
            ("PCBAs que repetiram a mesma falha", pcba_same), ("PCBAs com recorrência da mesma família", pcba_sim),
            ("Recorrência após MainBoard + Y", len(attention_pcba)),
        ],
        "material_kpis": [
            ("Materiais únicos 2º+ uso", len(reused_materials)), ("Materiais em 2º uso", mat_second),
            ("Materiais em 3º+ uso", mat_multi),
            ("Reuso COM falha em PCBA desvinculada", mat_with_old_failure),
            ("Reuso SEM falha em PCBA desvinculada", mat_without_old_failure),
            ("Reuso com a MESMA falha", mat_same), ("Reuso com falha da MESMA família", mat_sim),
            ("Reuso com somente outras falhas", mat_diff_only),
            ("Screens/LCDs reutilizados", screen_count), ("Câmeras reutilizadas", camera_count),
            ("Tipos de componente com reuso", sum(1 for x in component_types if x["Reutilizados únicos"] > 0)),
        ],
        "correlation_kpis": [
            ("Vínculos material↔PCBA com mesma falha", links_same),
            ("Vínculos material↔PCBA com mesma família", links_sim),
        ],
        "lines": lines,
        "component_types": component_types,
        "current_repair_states": sorted(current_process.items(), key=lambda x:(-x[1], x[0])),
        "defect_types": sorted(types.items(), key=lambda x:(-x[1],x[0])),
        "families": sorted(families.items(), key=lambda x:(-x[1],x[0])),
        "kpis": [],
        "material_types": [(x["Tipo material"], x["Reutilizados únicos"]) for x in component_types],
        "repair_states": sorted(current_process.items(), key=lambda x:(-x[1], x[0])),
    }

def _style_header(ws, row, col_start, col_end):
    for c in ws.iter_cols(min_col=col_start,max_col=col_end,min_row=row,max_row=row):
        cell=c[0]; cell.fill=PatternFill("solid",fgColor=HEADER_FILL); cell.font=Font(bold=True,color=HEADER_FONT); cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); cell.border=Border(bottom=LIGHT_BORDER)


def _write_table_sheet(wb, name, rows, headers, table_name, widths=None, empty_message=None):
    ws=wb.create_sheet(name)
    ws.append(headers)
    for r in rows: ws.append([r.get(h,"") for h in headers])
    _style_header(ws,1,1,len(headers))
    ws.freeze_panes=None; ws.sheet_view.pane=None; ws.sheet_view.zoomScale=85
    if rows:
        ref=f"A1:{get_column_letter(len(headers))}{len(rows)+1}"
        tab=Table(displayName=table_name,ref=ref); tab.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showFirstColumn=False,showLastColumn=False,showRowStripes=True,showColumnStripes=False); ws.add_table(tab)
        ws.auto_filter.ref=ref
    elif empty_message:
        ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=max(1,len(headers)))
        ws.cell(2,1).value=empty_message; ws.cell(2,1).fill=PatternFill("solid",fgColor=FILL_NONE); ws.cell(2,1).font=Font(bold=True,color="666666")
    for i,h in enumerate(headers,1):
        w=(widths or {}).get(h)
        if w is None:
            if h.startswith("Reuso "): w=27
            elif "Description" in h or "Desc" in h or "Comment" in h or "Motivo" in h or "Situação" in h: w=30
            elif "PCBA" in h or "Material SN" in h: w=27
            else: w=16
        ws.column_dimensions[get_column_letter(i)].width=w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment=Alignment(vertical="top",wrap_text=True); c.border=Border(bottom=LIGHT_BORDER)
    return ws


def _set_internal_link(cell, sheet_name, cell_ref, text=None):
    """Hyperlink interno XLSX via atributo location.

    O WPS trata target="#'SHEET'!A1" como arquivo externo em algumas versões.
    Usar location evita a mensagem "não é possível abrir o arquivo especificado".
    """
    if text is not None:
        cell.value = text
    display = str(cell.value or text or "")
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet_name}'!{cell_ref}", display=display)
    cell.style = "Hyperlink"


def _write_kpi_section(ws, start_row, title, pairs):
    ws.cell(start_row, 1).value = title
    ws.cell(start_row, 1).font = Font(bold=True, size=12, color="FFFFFF")
    ws.cell(start_row, 1).fill = PatternFill("solid", fgColor=HEADER_FILL)
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=8)
    row = start_row + 1
    for i, (label, value) in enumerate(pairs):
        group = i % 4
        if i and group == 0:
            row += 3
        c1 = 1 + group * 2
        c2 = c1 + 1
        ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
        ws.merge_cells(start_row=row + 1, start_column=c1, end_row=row + 1, end_column=c2)
        label_cell = ws.cell(row, c1); value_cell = ws.cell(row + 1, c1)
        label_cell.value = label; value_cell.value = value
        label_cell.fill = PatternFill("solid", fgColor=KPI_FILL)
        label_cell.font = Font(bold=True, color="44546A")
        label_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        value_cell.fill = PatternFill("solid", fgColor="FFFFFF")
        value_cell.font = Font(bold=True, size=16, color=HEADER_FILL)
        value_cell.alignment = Alignment(horizontal="center", vertical="center")
        for rr in (row, row+1):
            for cc in range(c1, c2+1):
                ws.cell(rr, cc).border = Border(left=LIGHT_BORDER, right=LIGHT_BORDER, top=LIGHT_BORDER, bottom=LIGHT_BORDER)
    groups = (len(pairs) + 3) // 4
    return start_row + 1 + groups * 3


def _write_indicators(wb, ind):
    ws = wb.create_sheet("INDICADORES")
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:H1"); ws["A1"] = "A-MES — Indicadores de rastreabilidade, reuso e recorrência"
    ws["A1"].fill = PatternFill("solid", fgColor=HEADER_FILL); ws["A1"].font = Font(color="FFFFFF", bold=True, size=15)
    ws["A1"].alignment = Alignment(vertical="center"); ws.row_dimensions[1].height = 27
    ws.merge_cells("A2:H2"); ws["A2"] = "Usos = ciclos da 3074. Falhas = registros da 2114. Reuso sem falha anterior não significa peça boa/ruim; significa ausência de falha registrada nas PCBAs desvinculadas."
    ws["A2"].font = Font(italic=True, color="666666"); ws["A2"].alignment = Alignment(wrap_text=True)

    r=4
    r=_write_kpi_section(ws,r,"VISÃO GERAL DO LOTE",ind["overview_kpis"]); r+=1
    r=_write_kpi_section(ws,r,"PCBA — USO x FALHAS 2114",ind["pcba_kpis"]); r+=1
    r=_write_kpi_section(ws,r,"MATERIAIS / COMPONENTES — REUSO E EVIDÊNCIA",ind["material_kpis"]); r+=1
    r=_write_kpi_section(ws,r,"VÍNCULOS MATERIAL ↔ PCBA DESVINCULADA",ind["correlation_kpis"]); r+=2

    # Por linha.
    ws.cell(r,1).value="POR LINHA"; ws.cell(r,1).font=Font(bold=True,size=12); ws.cell(r,1).fill=PatternFill("solid",fgColor=SECTION_FILL); ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=12); r+=1
    line_headers=["Linha","Falhas","PCBAs únicas","PCBA 2º uso","PCBA 3º+ uso","PCBA 2º+ uso / só 1 falha","Materiais 2º uso","Materiais 3º+ uso","Materiais com falha anterior","Materiais sem falha anterior","PCBAs mesma falha","PCBAs mesma família"]
    line_header_row=r
    for c,h in enumerate(line_headers,1): ws.cell(r,c).value=h
    _style_header(ws,r,1,len(line_headers)); r+=1; line_start=r
    for row in ind["lines"]:
        for c,h in enumerate(line_headers,1): ws.cell(r,c).value=row.get(h,"")
        r+=1
    line_end=r-1; r+=2

    # Todos os componentes, por Material SN único.
    ws.cell(r,1).value="TODOS OS COMPONENTES DO LOTE — MATERIAL SN ÚNICO"; ws.cell(r,1).font=Font(bold=True,size=12); ws.cell(r,1).fill=PatternFill("solid",fgColor=SECTION_FILL); ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=10); r+=1
    comp_headers=["Tipo material","Materiais únicos no lote","Reutilizados únicos","2º uso","3º+ uso","Com falha nas PCBAs desvinculadas","Sem falha nas PCBAs desvinculadas","Com mesma falha","Com mesma família","% reuso"]
    comp_header_row=r
    for c,h in enumerate(comp_headers,1): ws.cell(r,c).value=h
    _style_header(ws,r,1,len(comp_headers)); r+=1; comp_start=r
    for row in ind["component_types"]:
        for c,h in enumerate(comp_headers,1): ws.cell(r,c).value=row.get(h,"")
        ws.cell(r,10).number_format='0.0"%"'
        r+=1
    comp_end=r-1; r+=2

    # Situação de reparo e famílias.
    ws.cell(r,1).value="SITUAÇÃO ATUAL DO REPARO"; ws.cell(r,1).font=Font(bold=True,size=12); ws.cell(r,1).fill=PatternFill("solid",fgColor=SECTION_FILL); ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=4)
    ws.cell(r,6).value="FAMÍLIAS DE FALHA 2114"; ws.cell(r,6).font=Font(bold=True,size=12); ws.cell(r,6).fill=PatternFill("solid",fgColor=SECTION_FILL); ws.merge_cells(start_row=r,start_column=6,end_row=r,end_column=9); r+=1
    ws.cell(r,1).value="Situação"; ws.cell(r,2).value="Qtd"; _style_header(ws,r,1,2)
    ws.cell(r,6).value="Família"; ws.cell(r,7).value="Qtd"; _style_header(ws,r,6,7)
    maxn=max(len(ind["current_repair_states"]),len(ind["families"]),1)
    for i in range(maxn):
        rr=r+1+i
        if i<len(ind["current_repair_states"]): ws.cell(rr,1).value=ind["current_repair_states"][i][0]; ws.cell(rr,2).value=ind["current_repair_states"][i][1]
        if i<len(ind["families"]): ws.cell(rr,6).value=ind["families"][i][0]; ws.cell(rr,7).value=ind["families"][i][1]
    r=r+maxn+3

    ws.cell(r,1).value="DEFECT TYPE — TODOS OS VALORES PRESERVADOS"; ws.cell(r,1).font=Font(bold=True,size=12); ws.cell(r,1).fill=PatternFill("solid",fgColor=SECTION_FILL); ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=4); r+=1
    ws.cell(r,1).value="Defect Type"; ws.cell(r,2).value="Qtd"; _style_header(ws,r,1,2)
    for i,(name,count) in enumerate(ind["defect_types"],1): ws.cell(r+i,1).value=name; ws.cell(r+i,2).value=count
    r=r+max(len(ind["defect_types"]),1)+3

    # Gráficos compactos abaixo das tabelas para não cobrir dados.
    if comp_end>=comp_start:
        chart=BarChart(); chart.type="bar"; chart.style=10; chart.title="Reuso por tipo de componente — Material SN único"; chart.y_axis.title="Componente"; chart.x_axis.title="Quantidade"
        data=Reference(ws,min_col=3,min_row=comp_header_row,max_row=comp_end); cats=Reference(ws,min_col=1,min_row=comp_start,max_row=comp_end)
        chart.add_data(data,titles_from_data=True); chart.set_categories(cats); chart.height=max(8,min(16,4+len(ind["component_types"])*0.45)); chart.width=20; ws.add_chart(chart,f"A{r}"); r+=24
    if line_end>=line_start and ind["lines"]:
        chart2=BarChart(); chart2.type="col"; chart2.style=10; chart2.title="Falhas e PCBAs únicas por linha"; chart2.y_axis.title="Quantidade"
        data=Reference(ws,min_col=2,max_col=3,min_row=line_header_row,max_row=line_end); cats=Reference(ws,min_col=1,min_row=line_start,max_row=line_end)
        chart2.add_data(data,titles_from_data=True); chart2.set_categories(cats); chart2.height=8; chart2.width=20; ws.add_chart(chart2,f"A{r}")

    for col,w in {"A":35,"B":18,"C":19,"D":16,"E":18,"F":24,"G":24,"H":18,"I":18,"J":14,"K":18,"L":18}.items(): ws.column_dimensions[col].width=w
    for row in ws.iter_rows():
        for cell in row: cell.alignment=Alignment(vertical="center",wrap_text=True)
    return ws

def _apply_detail_sheet_defaults(ws):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.sheet_properties.outlinePr.summaryRight = False
    ws.freeze_panes = None
    ws.auto_filter.ref = None
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    return ws


def _write_grouped_pcba_sheet(wb, views):
    ws = wb.create_sheet("DETALHE_PCBA")
    headers = [
        "Voltar","Nível","Linha","Modelo","PCBA","Uso PCBA na falha","Usos conhecidos PCBA","Reusos anteriores PCBA",
        "Falhas registradas 2114","Falhas anteriores 2114","Mesma falha anterior","Mesma família anterior","Falhas diferentes anteriores",
        "Defect Time atual","Defect Code atual","Defect Desc atual","Repair Comment atual","Repair Status atual","Defect Type atual","Situação reparo atual",
        "Hist Seq","Defect Code 2114","Defect Description 2114","Defect Oper 2114","Registro falha 2114","Repair Status 2114","Defect Type 2114",
        "Situação reparo","Família 2114","Relação com falha atual","Motivo da relação","Ciclos de uso PCBA 3074"
    ]
    for c,h in enumerate(headers,1): ws.cell(1,c).value=h
    _style_header(ws,1,1,len(headers))
    hmap={h:i+1 for i,h in enumerate(headers)}
    by_key=defaultdict(list)
    for r in views["hist_pcba_rows"]: by_key[r.get("_failure_key")].append(r)
    anchors={}
    current_row=2
    for arow in views["analise_rows"]:
        key=arow.get("_failure_key"); histories=by_key.get(key,[])
        anchors[key]=current_row
        summary={
            "Voltar":"← ANALISE","Nível":"RESUMO PCBA","Linha":arow.get("Linha"),"Modelo":arow.get("Modelo"),"PCBA":arow.get("PCBA"),
            "Uso PCBA na falha":arow.get("Uso da PCBA na falha"),"Usos conhecidos PCBA":arow.get("Usos conhecidos da PCBA"),
            "Reusos anteriores PCBA":arow.get("Reusos anteriores da PCBA"),"Falhas registradas 2114":arow.get("Falhas registradas 2114"),
            "Falhas anteriores 2114":arow.get("Falhas anteriores 2114"),"Mesma falha anterior":arow.get("Mesma falha anterior"),
            "Mesma família anterior":arow.get("Mesma família anterior"),"Falhas diferentes anteriores":arow.get("Falhas diferentes anteriores"),
            "Defect Time atual":arow.get("Defect Time"),"Defect Code atual":arow.get("Defect Code"),"Defect Desc atual":arow.get("Defect Desc"),
            "Repair Comment atual":arow.get("Repair Comment"),"Repair Status atual":arow.get("Repair Status atual"),"Defect Type atual":arow.get("Defect Type atual"),
            "Situação reparo atual":arow.get("Situação reparo atual"),
            "Ciclos de uso PCBA 3074": (histories[0].get("Ciclos de uso PCBA 3074") if histories else ""),
        }
        for h in headers: ws.cell(current_row,hmap[h]).value=summary.get(h,"")
        for c in range(1,len(headers)+1): ws.cell(current_row,c).fill=PatternFill("solid",fgColor=FILL_SUMMARY); ws.cell(current_row,c).font=Font(bold=True)
        ws.row_dimensions[current_row].height=42
        summary_row=current_row; current_row+=1
        detail_start=current_row
        if histories:
            for r in histories:
                vals={"Voltar":"","Nível":"HISTÓRICO 2114",**{k:v for k,v in r.items() if not str(k).startswith("_")}}
                for h in headers: ws.cell(current_row,hmap[h]).value=vals.get(h,"")
                current_row+=1
        else:
            ws.cell(current_row,hmap["Nível"]).value="SEM HISTÓRICO 2114"; current_row+=1
        detail_end=current_row-1
        if detail_end>=detail_start:
            ws.row_dimensions.group(detail_start,detail_end,outline_level=1,hidden=True)
            ws.row_dimensions[summary_row].collapsed=True
        current_row+=1
    _apply_detail_sheet_defaults(ws)
    for col,w in {"A":12,"B":18,"C":13,"D":15,"E":27,"F":18,"G":16,"H":18,"I":17,"J":17,"K":17,"L":18,"M":18,"N":20,"O":16,"P":32,"Q":28,"R":15,"S":20,"T":34,"U":12,"V":16,"W":34,"X":12,"Y":16,"Z":14,"AA":20,"AB":38,"AC":20,"AD":22,"AE":28,"AF":48}.items():
        ws.column_dimensions[col].width=w
    _color_relations(ws,headers,"Relação com falha atual")
    return ws, anchors, hmap


def _write_grouped_material_sheet(wb, views):
    ws = wb.create_sheet("DETALHE_MATERIAL")
    headers = [
        "Voltar","Nível","Linha","Modelo","PCBA atual","Defect Time atual","Defect Code atual","Defect Desc atual",
        "Tipo material","Material SN","Uso material na falha","Reusos anteriores material","Usos conhecidos material","Reusos conhecidos material",
        "PCBAs desvinculadas únicas","PCBAs com mesma falha","PCBAs com mesma família","PCBAs com falha diferente","PCBAs sem histórico de falha",
        "Falhas 2114 nas desvinculadas","Evidência histórica","PCBA desvinculada","Estado vínculo atual","Ordem de uso anterior",
        "Falhas 2114 desta PCBA","Classificação desta PCBA","Hist Seq","Defect Code antigo","Defect Description antigo","Defect Oper antigo",
        "Registro falha antigo","Repair Status antigo","Defect Type antigo","Situação reparo antiga","Família falha antiga","Relação desta ocorrência","Motivo da relação"
    ]
    for c,h in enumerate(headers,1): ws.cell(1,c).value=h
    _style_header(ws,1,1,len(headers)); hmap={h:i+1 for i,h in enumerate(headers)}
    details_by_key=defaultdict(list)
    for r in views["hist_material_rows"]: details_by_key[r.get("_material_key")].append(r)
    anchors={}; current_row=2
    for s in views.get("material_summary_rows",[]):
        mkey=s.get("_material_key"); details=details_by_key.get(mkey,[]); anchors[mkey]=current_row
        summary={"Voltar":"← ANALISE","Nível":"RESUMO MATERIAL",**{k:v for k,v in s.items() if not str(k).startswith("_")}}
        for h in headers: ws.cell(current_row,hmap[h]).value=summary.get(h,"")
        for c in range(1,len(headers)+1): ws.cell(current_row,c).fill=PatternFill("solid",fgColor=FILL_SUMMARY); ws.cell(current_row,c).font=Font(bold=True)
        if str(s.get("Evidência histórica")) == "REUSO SEM FALHA ANTERIOR REGISTRADA":
            ws.cell(current_row,hmap["Evidência histórica"]).fill=PatternFill("solid",fgColor=FILL_NOFAIL)
        summary_row=current_row; current_row+=1
        # Agrupar primeiro por PCBA desvinculada para não contar a mesma PCBA várias vezes.
        by_pcba=defaultdict(list)
        for d in details: by_pcba[d.get("PCBA desvinculada")].append(d)
        pcba_block_start=current_row
        for pcba,rows in by_pcba.items():
            first=rows[0]
            pcba_summary={
                "Nível":"PCBA DESVINCULADA","Linha":first.get("Linha"),"Modelo":first.get("Modelo"),"PCBA atual":first.get("PCBA atual"),
                "Tipo material":first.get("Tipo material"),"Material SN":first.get("Material SN"),"PCBA desvinculada":pcba,
                "Estado vínculo atual":first.get("Estado vínculo atual"),"Ordem de uso anterior":first.get("Ordem de uso anterior"),
                "Falhas 2114 desta PCBA":first.get("Falhas 2114 desta PCBA"),"Classificação desta PCBA":first.get("Classificação desta PCBA"),
            }
            for h in headers: ws.cell(current_row,hmap[h]).value=pcba_summary.get(h,"")
            for c in range(1,len(headers)+1): ws.cell(current_row,c).fill=PatternFill("solid",fgColor=FILL_GROUP)
            ws.cell(current_row,hmap["Nível"]).font=Font(bold=True); ws.cell(current_row,hmap["PCBA desvinculada"]).font=Font(bold=True)
            pcba_summary_row=current_row; current_row+=1
            hist_start=current_row
            for d in rows:
                hist={"Nível":"FALHA 2114",**{k:v for k,v in d.items() if not str(k).startswith("_")}}
                for h in headers: ws.cell(current_row,hmap[h]).value=hist.get(h,"")
                current_row+=1
            hist_end=current_row-1
            if hist_end>=hist_start:
                ws.row_dimensions.group(hist_start,hist_end,outline_level=2,hidden=True)
                ws.row_dimensions[pcba_summary_row].outlineLevel=1
                ws.row_dimensions[pcba_summary_row].hidden=True
                ws.row_dimensions[pcba_summary_row].collapsed=True
        pcba_block_end=current_row-1
        # Todas as PCBAs-resumo começam recolhidas sob o material.
        if pcba_block_end>=pcba_block_start:
            for rr in range(pcba_block_start,pcba_block_end+1):
                if ws.row_dimensions[rr].outlineLevel < 1: ws.row_dimensions[rr].outlineLevel=1
                ws.row_dimensions[rr].hidden=True
            ws.row_dimensions[summary_row].collapsed=True
        current_row+=1
    _apply_detail_sheet_defaults(ws)
    for col,w in {"A":12,"B":20,"C":13,"D":15,"E":27,"F":20,"G":16,"H":32,"I":18,"J":30,"K":18,"L":18,"M":16,"N":18,"O":20,"P":18,"Q":18,"R":18,"S":20,"T":20,"U":34,"V":27,"W":18,"X":18,"Y":18,"Z":22,"AA":12,"AB":16,"AC":34,"AD":12,"AE":16,"AF":14,"AG":38,"AH":20,"AI":22,"AJ":28}.items():
        ws.column_dimensions[col].width=w
    _color_relations(ws,headers,"Classificação desta PCBA"); _color_relations(ws,headers,"Relação desta ocorrência")
    return ws, anchors, hmap


def _color_relations(ws, headers, name):
    if name not in headers: return
    c=headers.index(name)+1
    for r in range(2,ws.max_row+1):
        cell=ws.cell(r,c); v=str(cell.value or "").upper()
        if "MESMA FALHA" in v or "RECORRÊNCIA DA MESMA FALHA" in v: color=FILL_SAME
        elif "MESMA FAMÍLIA" in v: color=FILL_SIMILAR
        elif "DIFERENTE" in v or "OUTRAS" in v: color=FILL_DIFF
        elif "SEM HIST" in v or "SEM REPET" in v: color=FILL_NONE
        else: continue
        cell.fill=PatternFill("solid",fgColor=color); cell.font=Font(bold=True)


def gerar_excel_integrado(resultado_3074, resultado_2114, output_path):
    views = _build_views(resultado_3074, resultado_2114)
    ind = _build_indicators(views)
    wb = Workbook(); wb.remove(wb.active)
    _write_indicators(wb, ind)

    headers_a = [
        "Linha","Modelo","PCBA","Defect Time","Defect Code","Defect Desc","Repair Comment","Reason Type","Reason Desc",
        "Uso da PCBA na falha","Usos conhecidos da PCBA","Reusos anteriores da PCBA",
        "Falhas registradas 2114","Falhas anteriores 2114","Mesma falha anterior","Mesma família anterior","Falhas diferentes anteriores",
        "Situação da PCBA","Repair Status atual","Defect Type atual","Situação reparo atual","Registro falha atual",
        "Materiais na falha","Materiais 2º+ uso"
    ] + [f"Reuso {i}" for i in range(1, views["max_reusos"]+1)]
    ws_a = _write_table_sheet(wb,"ANALISE",views["analise_rows"],headers_a,"T_ANALISE_V015",widths={
        "Linha":13,"Modelo":15,"PCBA":27,"Defect Time":20,"Defect Desc":32,"Repair Comment":30,
        "Uso da PCBA na falha":18,"Usos conhecidos da PCBA":16,"Falhas registradas 2114":16,
        "Situação da PCBA":28,"Situação reparo atual":34,"Materiais 2º+ uso":15
    })
    _color_relations(ws_a, headers_a, "Situação da PCBA")

    ws_p, pcba_anchor, _ = _write_grouped_pcba_sheet(wb, views)
    ws_m, mat_anchor, _ = _write_grouped_material_sheet(wb, views)

    # Mapas de navegação. Clique leva ao RESUMO; a partir dali o usuário expande +/−.
    failure_mat_anchor={}
    for s in views.get("material_summary_rows",[]): failure_mat_anchor.setdefault(s.get("_failure_key"), mat_anchor.get(s.get("_material_key")))

    hmap={h:i+1 for i,h in enumerate(headers_a)}
    for rnum,row in enumerate(views["analise_rows"],2):
        key=row.get("_failure_key")
        if key in pcba_anchor:
            for h in ("Uso da PCBA na falha","Falhas registradas 2114"):
                _set_internal_link(ws_a.cell(rnum,hmap[h]),"DETALHE_PCBA",f"A{pcba_anchor[key]}")
        if failure_mat_anchor.get(key):
            _set_internal_link(ws_a.cell(rnum,hmap["Materiais 2º+ uso"]),"DETALHE_MATERIAL",f"A{failure_mat_anchor[key]}")
        for i in range(1,views["max_reusos"]+1):
            mk=row.get(f"_reuso_{i}_key")
            if mk in mat_anchor:
                _set_internal_link(ws_a.cell(rnum,hmap[f"Reuso {i}"]),"DETALHE_MATERIAL",f"A{mat_anchor[mk]}")
        ws_a.row_dimensions[rnum].height=78 if int(row.get("Materiais 2º+ uso") or 0)>0 else 34

    # Links de retorno dos resumos. Detalhes ficam agrupados/recolhidos.
    for key,rownum in pcba_anchor.items():
        # acha a linha ANALISE pelo próprio key
        ar = next((r.get("_analysis_row") for r in views["analise_rows"] if r.get("_failure_key")==key),2)
        _set_internal_link(ws_p.cell(rownum,1),"ANALISE",f"A{ar}","← ANALISE")
    for mkey,rownum in mat_anchor.items():
        s=next((x for x in views.get("material_summary_rows",[]) if x.get("_material_key")==mkey),None)
        ar=(s or {}).get("_analysis_row",2)
        _set_internal_link(ws_m.cell(rownum,1),"ANALISE",f"A{ar}","← ANALISE")

    # Validações de consistência: jamais gerar dashboard enganoso.
    unique_reuse = len({r.get("Material SN") for r in views["reusos"] if r.get("Material SN")})
    if unique_reuse != len({s.get("Material SN") for s in views.get("material_summary_rows",[]) if s.get("Material SN")}):
        raise RuntimeError("Materiais reutilizados não consolidaram corretamente para DETALHE_MATERIAL.")
    camera_count = next((v for k,v in ind["material_kpis"] if k=="Câmeras reutilizadas"),0)
    if camera_count > unique_reuse:
        raise RuntimeError("Indicador impossível: câmeras reutilizadas maior que total de materiais reutilizados únicos.")

    wb.active=wb.sheetnames.index("INDICADORES")
    output_path=Path(output_path); output_path.parent.mkdir(parents=True,exist_ok=True); wb.save(output_path)
    return {
        "falhas":len(views["analise_rows"]),
        "pcbas_unicas":len({r.get('PCBA') for r in views['analise_rows']}),
        "materiais_reuso":unique_reuse,
        "materiais_reuso_sem_falha_anterior":next((v for k,v in ind['material_kpis'] if k=='Reuso SEM falha em PCBA desvinculada'),0),
        "materiais_reuso_com_falha_anterior":next((v for k,v in ind['material_kpis'] if k=='Reuso COM falha em PCBA desvinculada'),0),
        "linhas_hist_pcba":len(views["hist_pcba_rows"]),
        "linhas_hist_material":len(views["hist_material_rows"]),
        "pcba_segundo_uso":next((v for k,v in ind['pcba_kpis'] if k=='PCBA em 2º uso'),0),
        "pcba_multiuso":next((v for k,v in ind['pcba_kpis'] if k=='PCBA em 3º+ uso'),0),
    }

