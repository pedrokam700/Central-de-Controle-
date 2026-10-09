from __future__ import annotations

from mes_scheduler import MES, mes_call, Cancelled, MonitorSkipped

import argparse
import base64
import mimetypes
import shutil
import socket
import json
import os
import re
import subprocess
import threading
import time
import traceback
import uuid
import zipfile
from datetime import datetime, timedelta
from http import HTTPStatus
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlparse, quote
from urllib.request import Request, urlopen

from ames_3028 import parse_export
from ames_3028_live import collect_live_3028 as _collect_live_3028
collect_live_3028 = mes_call(_collect_live_3028)
from engine_bridge import (choose_current_history, discover_engine, refresh_2114, repair_state,
                           run_full_v016, run_deep_v016_records, run_individual_lookup, regenerate_v016_excel, build_trace_insights)
from store import Store, stable_json, utc_now

AGENT_VERSION = "0.5.23"
BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
DB_PATH = DATA_DIR / "ames_local.sqlite3"
CONFIG_PATH = BASE_DIR / "config.json"
BACKUP_DIR = DATA_DIR / "backups"
WORKSTATION = socket.gethostname()


def load_config():
    default = {
        "host": "127.0.0.1",
        "port": 8765,
        "engine_dir": "",
        "configured_lines": ["TAN10101", "TAN10102", "TAN10103"],
        "line_aliases": {"Linha 1": "TAN10101", "Linha 2": "TAN10102", "Linha 3": "TAN10103"},
        "day_start": "07:00",
        "monitor_interval_minutes": 30,
        "performance": "normal",
        "chrome_cdp": "http://127.0.0.1:9222",
        "ames_host": "172.29.185.215",
        "ames_port": 80,
        "ames_start_url": "http://172.29.185.215/asymes",
        "chrome_profile_dir": "",
        "auto_3028": {"ready": True, "note": "coleta direta ExtJS validada em fabrica: 107/107"},
        "auto_3022": {"ready": False, "note": "view AWIP3022 identificada; adapter individual em validação de fábrica"},
        "setup_complete": False,
        "backup_retention": 10,
        "auto_backup_on_start": True,
        "allowed_origins": ["https://central-cora-v15-1-13-21.vercel.app", "https://central-cora-v15-1-13-21-pedrokam700-6477.vercel.app", "https://central-cora-v2-pedrokam700-6477.vercel.app", "https://central-cora-v2.vercel.app", "https://central-ames-bridge.vercel.app"],
    }
    if CONFIG_PATH.exists():
        try:
            loaded = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
            default.update(loaded)
        except Exception:
            pass
    return default


CONFIG = load_config()


def normalize_line_id(value):
    raw = str(value or "").strip()
    if not raw:
        return ""
    aliases = CONFIG.get("line_aliases") or {}
    for alias, target in aliases.items():
        if raw.casefold() == str(alias).strip().casefold():
            return str(target).strip()
    # Compatibilidade com nomes amigaveis ja salvos em builds anteriores.
    m = re.fullmatch(r"linha\s*0*([1-3])", raw, flags=re.IGNORECASE)
    if m:
        return f"TAN1010{m.group(1)}"
    return raw.upper() if raw.lower().startswith("tan") else raw


def normalize_lines(values):
    out = []
    seen = set()
    for value in values or []:
        line = normalize_line_id(value)
        if line and line not in seen:
            seen.add(line)
            out.append(line)
    return out

CONFIG["configured_lines"] = normalize_lines(CONFIG.get("configured_lines") or [])

STORE = Store(DB_PATH)
STORE.initialize()
ENGINE_DIR = discover_engine(CONFIG, BASE_DIR)

JOBS = {}
JOBS_LOCK = threading.Lock()
MONITOR_STOP = threading.Event()
MONITOR_THREAD = None
JOB_CANCEL = {}




def _tcp_probe(host, port, timeout=1.2):
    try:
        with socket.create_connection((str(host), int(port)), timeout=timeout):
            return True
    except OSError:
        return False


def _find_chrome():
    candidates = [
        Path(os.environ.get("PROGRAMFILES", r"C:\Program Files")) / "Google/Chrome/Application/chrome.exe",
        Path(os.environ.get("PROGRAMFILES(X86)", r"C:\Program Files (x86)")) / "Google/Chrome/Application/chrome.exe",
        Path(os.environ.get("LOCALAPPDATA", "")) / "Google/Chrome/Application/chrome.exe",
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    found = shutil.which("chrome") or shutil.which("chrome.exe")
    return Path(found) if found else None


@mes_call
def start_dedicated_chrome():
    profile_cfg = str(CONFIG.get("chrome_profile_dir") or "").strip()
    profile = Path(profile_cfg) if profile_cfg else Path(os.environ.get("LOCALAPPDATA", str(BASE_DIR))) / "CentralAMES" / "ChromeProfile"
    profile.mkdir(parents=True, exist_ok=True)
    url = str(CONFIG.get("ames_start_url") or "http://172.29.185.215/asymes")
    if _tcp_probe("127.0.0.1", 9222, timeout=0.5):
        # Open the URL inside the already-running dedicated CDP browser, never in the default browser.
        try:
            req = Request(f"http://127.0.0.1:9222/json/new?{quote(url, safe='')}", method="PUT")
            with urlopen(req, timeout=2):
                pass
        except Exception:
            pass
        return {"ok": True, "already_running": True, "cdp": CONFIG.get("chrome_cdp"), "profile": str(profile), "url": url}
    chrome = _find_chrome()
    if not chrome:
        raise RuntimeError("Google Chrome nao encontrado neste computador")
    args = [
        str(chrome),
        "--remote-debugging-port=9222",
        f"--user-data-dir={profile}",
        "--no-first-run",
        "--no-default-browser-check",
        "--disable-background-mode",
        url,
    ]
    subprocess.Popen(args, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    deadline = time.time() + 8
    while time.time() < deadline:
        if _tcp_probe("127.0.0.1", 9222, timeout=0.4):
            return {"ok": True, "already_running": False, "chrome": str(chrome), "profile": str(profile), "url": url}
        time.sleep(0.35)
    return {"ok": False, "already_running": False, "chrome": str(chrome), "profile": str(profile), "url": url, "message": "Chrome abriu, mas a porta CDP 9222 ainda nao respondeu"}

def write_config(config):
    CONFIG_PATH.write_text(json.dumps(config, ensure_ascii=False, indent=2), encoding="utf-8")


def create_backup(reason="manual"):
    info = STORE.backup(BACKUP_DIR, keep=int(CONFIG.get("backup_retention", 10) or 10), reason=reason)
    STORE.record_audit("DATABASE_BACKUP", workstation=WORKSTATION, details=info)
    return info


def maybe_backup_on_start():
    if not CONFIG.get("auto_backup_on_start", True):
        return None
    latest = STORE.latest_backup_info(BACKUP_DIR)
    today = datetime.now().date().isoformat()
    if latest and str(latest.get("modified_at", "")).startswith(today):
        return latest
    return create_backup("startup")


def _json_bytes(obj):
    return json.dumps(obj, ensure_ascii=False, default=str).encode("utf-8")


def _job_update(job_id, **fields):
    with JOBS_LOCK:
        job = JOBS.setdefault(job_id, {})
        job.update(fields)
        return dict(job)


def _job_create(kind, config):
    job_id = uuid.uuid4().hex[:12]
    now = utc_now()
    job = {
        "id": job_id,
        "kind": kind,
        "status": "queued",
        "stage": "queued",
        "progress": 0,
        "message": "Na fila",
        "created_at": now,
        "config": config,
    }
    with JOBS_LOCK:
        JOBS[job_id] = job
        JOB_CANCEL[job_id] = threading.Event()
    with STORE.connect() as con:
        con.execute(
            "INSERT OR REPLACE INTO jobs(id,kind,mode,status,created_at,progress,stage,message,config_json) VALUES(?,?,?,?,?,?,?,?,?)",
            (job_id, kind, config.get("mode"), "queued", now, 0, "queued", "Na fila", stable_json(config)),
        )
    return job


def _persist_job(job):
    with STORE.connect() as con:
        con.execute(
            """UPDATE jobs SET status=?,started_at=COALESCE(started_at,?),finished_at=?,progress=?,stage=?,message=?,result_json=?,error=? WHERE id=?""",
            (
                job.get("status"), job.get("started_at"), job.get("finished_at"), job.get("progress", 0),
                job.get("stage"), job.get("message"), stable_json(job.get("result")) if job.get("result") is not None else None,
                job.get("error"), job.get("id"),
            ),
        )


def ingest_integrated(payload, source_name="integrated.json"):
    snapshot_id = STORE.ingest_integrated_payload(payload, source_name=source_name)
    trace_2114 = payload.get("trace_2114") or {}
    if trace_2114.get("pcbas"):
        _attach_2114_to_current(snapshot_id, {"pcbas": trace_2114.get("pcbas") or [], "shift": trace_2114.get("shift")})
    return {
        "snapshot_id": snapshot_id,
        "dashboard": STORE.dashboard(snapshot_id),
        "rows": len(STORE.base_rows(snapshot_id, limit=20000)),
    }


def ingest_3028_path(path, *, mode="manual", start_at=None, end_at=None, shift=None, lines=None):
    path = Path(path)
    parsed = parse_export(path)
    lines = list(lines or parsed.get("lines") or [])
    window_id = STORE.ensure_window(mode, start_at, end_at, shift, lines)
    snapshot_id = STORE.create_snapshot(
        window_id=window_id,
        source_kind="3028_export",
        source_name=path.name,
        source_bytes=path.read_bytes(),
    )
    STORE.ingest_line_metrics(snapshot_id, parsed.get("line_metrics") or [])
    STORE.ingest_defects(snapshot_id, parsed.get("records") or [], reconcile_window=True)
    STORE.save_snapshot_payload(snapshot_id, "3028_parsed", parsed)
    STORE.rebuild_cora_index(snapshot_id)
    STORE.record_audit("INGEST_3028", workstation=WORKSTATION, snapshot_id=snapshot_id,
                       details={"source": path.name, "mode": mode, "lines": lines, "records": parsed.get("records_count")})
    return {"snapshot_id": snapshot_id, "parsed": parsed, "dashboard": STORE.dashboard(snapshot_id)}


def _persist_3028_live(live, *, mode="today", start_at=None, end_at=None, shift=None):
    """Persist each production line in its own logical window/snapshot.

    This is intentional: Linha 1/2/3 can be refreshed at different moments without
    making one line hide or reconcile defects from another line. A batch collection
    is still serial in the MES, but the local history remains independent per line.
    """
    actual_lines = normalize_lines(live.get("lines") or live.get("requested_lines") or [])
    snapshot_ids = {}
    captures = live.get("captures") or []
    validations = live.get("validations") or []
    for line in actual_lines:
        line_metrics = [m for m in (live.get("line_metrics") or []) if normalize_line_id(m.get("line")) == line]
        line_records = [r for r in (live.get("records") or []) if normalize_line_id(r.get("line")) == line]
        line_caps = [c for c in captures if normalize_line_id(c.get("lineId")) == line]
        line_valid = [v for v in validations if normalize_line_id(v.get("line")) == line]
        line_live = {
            "source": live.get("source"),
            "captured_at": live.get("captured_at"),
            "requested_lines": [line],
            "window": live.get("window"),
            "line_metrics": line_metrics,
            "records": line_records,
            "records_count": len(line_records),
            "lines": [line],
            "validations": line_valid,
            "visible_state": live.get("visible_state") if len(actual_lines) == 1 else None,
            "captures": line_caps,
        }
        window_id = STORE.ensure_window(mode, start_at, end_at, shift, [line])
        raw_bytes = stable_json(line_live).encode("utf-8")
        snapshot_id = STORE.create_snapshot(
            window_id=window_id,
            source_kind="3028_live_extjs",
            source_name=f"AWIP3028 direct ExtJS · {line}",
            source_bytes=raw_bytes,
            notes="Coleta direta por linha; sem Excel intermediario",
        )
        STORE.ingest_line_metrics(snapshot_id, line_metrics)
        STORE.ingest_defects(snapshot_id, line_records, reconcile_window=True)
        STORE.save_snapshot_payload(snapshot_id, "3028_live_raw", line_live)
        STORE.rebuild_cora_index(snapshot_id)
        STORE.record_audit(
            "INGEST_3028_LIVE", workstation=WORKSTATION, snapshot_id=snapshot_id, line=line,
            details={
                "mode": mode, "line": line, "records": len(line_records),
                "window": live.get("window"), "validations": line_valid,
            },
        )
        snapshot_ids[line] = snapshot_id
    latest = max(snapshot_ids.values()) if snapshot_ids else None
    return {
        "snapshot_id": latest,
        "snapshot_ids": snapshot_ids,
        "live": live,
        "dashboard": STORE.team_dashboard(actual_lines),
    }


def ingest_3028_live(*, mode="today", start_at=None, end_at=None, shift=None, lines=None):
    if not start_at or not end_at:
        raise RuntimeError("Janela 3028 incompleta")
    live = collect_live_3028(
        cdp_url=str(CONFIG.get("chrome_cdp") or "http://127.0.0.1:9222"),
        ames_host=str(CONFIG.get("ames_host") or "172.29.185.215"),
        lines=normalize_lines(lines or CONFIG.get("configured_lines") or []),
        start_at=start_at,
        end_at=end_at,
    )
    return _persist_3028_live(live, mode=mode, start_at=start_at, end_at=end_at, shift=shift)


def enrich_snapshot_from_integrated(snapshot_id, payload, source_path=None):
    """Attach validated V0.16 3074/2114 evidence to an existing 3028 snapshot."""
    trace_3074 = payload.get("trace_3074") or {}
    trace_2114 = payload.get("trace_2114") or {}
    STORE.save_snapshot_payload(snapshot_id, "integrated_v016", payload)
    STORE.ingest_material_reuse(snapshot_id, trace_3074)
    STORE.ingest_2114_contexts(snapshot_id, trace_2114)
    updated = _attach_2114_to_current(snapshot_id, trace_2114)
    STORE.rebuild_cora_index(snapshot_id)
    if source_path:
        with STORE.connect() as con:
            con.execute(
                "INSERT INTO artifacts(snapshot_id,artifact_type,path,created_at,metadata_json) VALUES(?,?,?,?,?)",
                (snapshot_id, "integrated_json_v016", str(source_path), utc_now(), stable_json({"version": payload.get("version")})),
            )
    return {
        "snapshot_id": snapshot_id,
        "material_rows": len(STORE.dataset_rows("material_reuse", snapshot_id=snapshot_id, limit=100000)),
        "history_rows": len(STORE.dataset_rows("pcba_history", snapshot_id=snapshot_id, limit=100000)),
        "status_rows_updated": updated,
    }


def _attach_2114_to_current(snapshot_id, refresh_payload):
    pcbas = refresh_payload.get("pcbas") or []
    STORE.ingest_2114(snapshot_id, pcbas)
    base_rows = STORE.base_rows(snapshot_id, limit=50000)
    by_pcba = {}
    for r in base_rows:
        by_pcba.setdefault(r.get("pcba_sn"), []).append(r)
    updated = 0
    for entry in pcbas:
        sn = entry.get("pcba_sn")
        if entry.get("status") != "COMPLETE" or not entry.get("result"):
            STORE.record_repair_refresh(snapshot_id, sn, "ERROR", error=entry.get("error"))
            continue
        result = entry["result"]
        hist = result.get("rows") or []
        STORE.record_repair_refresh(snapshot_id, sn, "COMPLETE", payload=result)
        for current in by_pcba.get(sn, []):
            chosen = choose_current_history(hist, current.get("defect_code"), current.get("defect_desc"))
            if not chosen:
                continue
            updated += STORE.update_defect_status(
                snapshot_id, sn, current.get("defect_code"), current.get("defect_desc"),
                chosen.get("Repair Status"), chosen.get("Defect Type"), repair_state(chosen),
            )
    STORE.rebuild_cora_index(snapshot_id)
    return updated


def refresh_repairs(snapshot_id=None, line=None, unresolved_only=False):
    snapshot_id = int(snapshot_id or STORE.latest_snapshot_id() or 0)
    if not snapshot_id:
        raise RuntimeError("Nenhum snapshot local carregado")
    engine = discover_engine(CONFIG, BASE_DIR)
    if not engine:
        raise RuntimeError("Motor V0.16 nao encontrado. Configure AMES_ENGINE_DIR ou engine_dir em config.json")
    pcbas = STORE.current_pcbas(snapshot_id, line=line, only_unresolved=unresolved_only)
    payload = refresh_2114(engine, pcbas)
    updated = _attach_2114_to_current(snapshot_id, payload)
    STORE.record_audit("REPAIR_REFRESH", workstation=WORKSTATION, snapshot_id=snapshot_id, line=line,
                       details={"pcbas_consulted": len(pcbas), "rows_updated": updated, "unresolved_only": bool(unresolved_only)})
    return {"snapshot_id": snapshot_id, "pcbas_consulted": len(pcbas), "rows_updated": updated,
            "shift": payload.get("shift"), "dashboard": STORE.dashboard(snapshot_id)}


def _run_full_pipeline_job(job_id, cfg):
    job = _job_update(job_id, status="running", stage="preflight", progress=2,
                      started_at=utc_now(), message="Preparando coleta local")
    _persist_job(job)
    try:
        source_path = cfg.get("source_path")
        if source_path:
            job = _job_update(job_id, stage="3028", progress=10, message="Lendo export detalhado 3028 e FPY oficial")
            _persist_job(job)
            imported = ingest_3028_path(
                source_path, mode=cfg.get("mode", "manual"), start_at=cfg.get("start_at"),
                end_at=cfg.get("end_at"), shift=cfg.get("shift"), lines=cfg.get("lines") or [],
            )
            snapshot_id = imported["snapshot_id"]

            engine = discover_engine(CONFIG, BASE_DIR)
            if not engine:
                # The 3028 snapshot is still valid and usable even if deep trace is unavailable.
                raise RuntimeError("Motor V0.16 nao encontrado para executar 3074/2114. O snapshot 3028 foi preservado.")

            def on_engine_line(line):
                nonlocal job
                low = line.lower()
                if "fase 1" in low or "3074" in low:
                    stage, progress, msg = "3074", 35, "Rastreando componentes e reuso 3074"
                elif "fase 2114" in low or ("2114" in low and "fase" in low):
                    stage, progress, msg = "2114", 72, "Coletando historicos e reparos 2114"
                else:
                    return
                job = _job_update(job_id, stage=stage, progress=progress, message=msg)
                _persist_job(job)

            job = _job_update(job_id, stage="3074", progress=20, message="Executando motor V0.16 com checkpoint")
            _persist_job(job)
            deep = run_full_v016(engine, source_path, line_callback=on_engine_line)
            job = _job_update(job_id, stage="2114", progress=84, message="Persistindo rastreabilidade e estado atual")
            _persist_job(job)
            enrich = enrich_snapshot_from_integrated(snapshot_id, deep["payload"], source_path=deep.get("json_path"))

            # 3022 is intentionally not guessed. Tomorrow's real screen inspection plugs
            # into process_events without changing the rest of this pipeline.
            job = _job_update(job_id, stage="3022", progress=92,
                              message="3022 aguardando adapter validado; snapshot 3028/3074/2114 concluido")
            _persist_job(job)
            job = _job_update(
                job_id, status="done", stage="done", progress=100, finished_at=utc_now(),
                message="3028/3074/2114 concluidos; 3022 sera enriquecido quando validado",
                result={
                    "snapshot_id": snapshot_id, "dashboard": STORE.dashboard(snapshot_id),
                    "integrated_json": deep.get("json_path"), "enrichment": enrich,
                    "3022_status": "PENDING_ADAPTER_VALIDATION",
                },
            )
            _persist_job(job)
            return

        # Coleta direta 3028 validada em fabrica (107/107).
        if not (CONFIG.get("auto_3028") or {}).get("ready"):
            raise RuntimeError("Coleta automatica 3028 esta desativada na configuracao")

        if not _tcp_probe(CONFIG.get("ames_host", "172.29.185.215"), int(CONFIG.get("ames_port", 80)), timeout=1.5):
            raise RuntimeError("Rede A-MES indisponivel. Conecte TAXXX_5G e valide a rota temporaria do A-MES.")
        if not _tcp_probe("127.0.0.1", 9222, timeout=0.7):
            chrome_result = start_dedicated_chrome()
            if not chrome_result.get("ok"):
                raise RuntimeError(chrome_result.get("message") or "Chrome dedicado nao iniciou")

        job = _job_update(job_id, stage="3028", progress=12, message="Consultando 3028 diretamente no A-MES")
        _persist_job(job)
        imported = ingest_3028_live(
            mode=cfg.get("mode", "today"), start_at=cfg.get("start_at"), end_at=cfg.get("end_at"),
            shift=cfg.get("shift"), lines=cfg.get("lines") or [],
        )
        snapshot_id = imported["snapshot_id"]
        snapshot_ids = imported.get("snapshot_ids") or {}
        validations = imported.get("live", {}).get("validations") or []
        rows = imported.get("live", {}).get("records_count") or 0
        lines_done = imported.get("live", {}).get("lines") or []

        job = _job_update(
            job_id, status="done", stage="done", progress=100, finished_at=utc_now(),
            message="3028 direta concluida; FPY/Top 3/base local atualizados",
            result={
                "snapshot_id": snapshot_id,
                "snapshot_ids": snapshot_ids,
                "dashboard": STORE.team_dashboard(lines_done),
                "records_3028": rows,
                "lines": lines_done,
                "validations": validations,
                "start_at": cfg.get("start_at"),
                "end_at": cfg.get("end_at"),
                "requested_lines": normalize_lines(cfg.get("lines") or CONFIG.get("configured_lines") or []),
                "stages": {
                    "3028": "done",
                    "3074": "not_run",
                    "2114": "not_run",
                    "3022": "pending_adapter",
                },
            },
        )
        _persist_job(job)
        return
    except Exception as exc:
        job = _job_update(job_id, status="error", stage="error", finished_at=utc_now(),
                          message=str(exc), error=f"{type(exc).__name__}: {exc}")
        _persist_job(job)


def _launch_mes_job(kind, cfg, worker, monitor=False):
    reservation = MES.reserve_monitor() if monitor else None
    if monitor and reservation is None:
        return None  # No job/status row or queued thread for a skipped tick.
    try:
        job = _job_create(kind, cfg)
        def state_change(state):
            _job_update(job["id"], mes_state=state)
        def run():
            try:
                with MES.job(token=JOB_CANCEL[job["id"]], monitor=monitor,
                             status=state_change, reservation=reservation):
                    worker(job["id"], cfg)
            except (Cancelled, MonitorSkipped) as exc:
                status = "cancelled" if isinstance(exc, Cancelled) else "skipped"
                _persist_job(_job_update(job["id"], status=status, stage=status,
                    finished_at=utc_now(), message=status, mes_state="idle"))
            finally:
                _job_update(job["id"], mes_state="idle")
                with JOBS_LOCK:
                    JOB_CANCEL.pop(job["id"], None)
        threading.Thread(target=run, daemon=True).start()
        return job
    except BaseException:
        if reservation is not None:
            with MES.job(reservation=reservation):
                pass
        raise


def start_job(cfg, *, monitor=False):
    return _launch_mes_job("analysis", cfg, _run_full_pipeline_job, monitor=monitor)




def _select_trace_records(records, defect_codes=None, max_failures=0, max_pcbas=0):
    """Apply operator scope without changing the underlying 3028 snapshot.

    Records are ordered newest-first for bounded diagnostic runs. When max_pcbas is
    used, all selected occurrences belonging to the chosen PCBAs are kept so one
    PCBAs history is not cut in half.
    """
    rows = [dict(r) for r in (records or []) if r.get("present_in_3028", 1) and str(r.get("pcba_sn") or "").strip()]
    codes = {str(x).strip() for x in (defect_codes or []) if str(x).strip()}
    if codes:
        rows = [r for r in rows if str(r.get("defect_code") or "").strip() in codes]
    rows.sort(key=lambda r: str(r.get("defect_time") or ""), reverse=True)
    if max_failures and int(max_failures) > 0:
        rows = rows[:int(max_failures)]
    if max_pcbas and int(max_pcbas) > 0:
        chosen=[]; seen=set()
        for r in rows:
            sn=str(r.get("pcba_sn") or "").strip()
            if sn not in seen:
                if len(chosen) >= int(max_pcbas):
                    continue
                seen.add(sn); chosen.append(sn)
        keep=set(chosen)
        rows=[r for r in rows if str(r.get("pcba_sn") or "").strip() in keep]
    return rows


def _job_live_event(job_id, event, line=None, line_index=1, line_total=1, snapshot_id=None):
    """Update accurate per-stage progress and a small live feed without hammering SQLite."""
    with JOBS_LOCK:
        job = JOBS.setdefault(job_id,{})
        stage=str((event or {}).get("stage") or "")
        cur=int((event or {}).get("current") or 0)
        total=max(0,int((event or {}).get("total") or 0))
        pct=round((cur*100/total),1) if total else 0.0
        global_pct=round((((max(1,line_index)-1)+(pct/100.0))/max(1,line_total))*100,1) if total else 0.0
        sp=dict(job.get("stage_progress") or {})
        if stage in {"3074","2114"}:
            detail=""
            if stage=="3074":
                detail=str((event or {}).get("sn") or "")
                if (event or {}).get("kind")=="component":
                    detail += f" · falha {(event or {}).get('defect_index')}/{(event or {}).get('defect_total')} · comp {(event or {}).get('component_index')}/{(event or {}).get('component_total')}"
            else:
                detail=str((event or {}).get("pcba") or "")
            sp[stage]={"current":cur,"total":total,"percent":global_pct,"line_percent":pct,"line":line,"line_index":line_index,"line_total":line_total,"detail":detail,"status":(event or {}).get("status")}
            job["stage_progress"]=sp
        # Never put the full stage result in the HTTP job payload.
        safe={k:v for k,v in (event or {}).items() if k not in {"result","state","item"}}
        safe["line"]=line; safe["at"]=utc_now()
        feed=list(job.get("live_events") or [])
        feed.append(safe)
        job["live_events"]=feed[-40:]
        JOBS[job_id]=job
        return dict(job)


def build_share_payload():
    """Sanitized, compact payload for authenticated Central users.

    It contains already-collected summaries only. It never exposes A-MES cookies,
    credentials, CDP state or any ability to query the OPPO network remotely.
    """
    lines=CONFIG.get("configured_lines") or ["TAN10101","TAN10102","TAN10103"]
    dash=STORE.team_dashboard(lines)
    out=[]
    for entry in dash.get("lines") or []:
        line=entry.get("line")
        sid=entry.get("snapshot_id")
        catalog=STORE.dataset_catalog(sid) if sid else []
        counts={x.get("dataset"):int(x.get("rows") or 0) for x in catalog}
        metric=entry.get("metrics") or {}
        reuse_summary={}
        try:
            insight=trace_insights(line=line,snapshot_id=sid) if sid else {}
            if insight.get("ready"):
                reuse_summary={
                    "pcba":dict(insight.get("pcba_kpis") or []),
                    "material":dict(insight.get("material_kpis") or []),
                    "correlation":dict(insight.get("correlation_kpis") or []),
                }
        except Exception:
            reuse_summary={}
        out.append({
            "line":line,"snapshot_id":sid,"collected_at":entry.get("collected_at"),
            "fpy":metric.get("fpy"),"check_fpy":metric.get("check_fpy"),"quantity":metric.get("quantity"),
            "defect_rows":int(entry.get("defect_rows") or 0),"unresolved":int(entry.get("unresolved") or 0),
            "removed_from_latest":int(entry.get("removed_from_latest") or 0),"top3":entry.get("top3") or [],
            "pcba_history_count":counts.get("pcba_history",0),"material_reuse_count":counts.get("material_reuse",0),
            "reuse_summary":reuse_summary,
        })
    return {"schema":"ames-share-v1","generated_at":utc_now(),"workstation":WORKSTATION,"agent_version":AGENT_VERSION,"lines":out}

def trace_insights(line=None, snapshot_id=None):
    """Return approved V0.16 reuse/correlation indicators for one physical line."""
    line = normalize_line_id(line)
    if not snapshot_id and line:
        snapshot_id = STORE.latest_snapshot_id_for_line(line)
    snapshot_id = int(snapshot_id or STORE.latest_snapshot_id() or 0)
    if not snapshot_id:
        return {"schema":"ames-insights-v1","line":line or None,"snapshot_id":None,"ready":False,"message":"Nenhum snapshot 3028 disponível.","drilldowns":{}}
    payload = STORE.get_snapshot_payload(snapshot_id, kinds=["integrated_v016"])
    if not payload:
        return {"schema":"ames-insights-v1","line":line or None,"snapshot_id":snapshot_id,"ready":False,"message":"Este snapshot ainda não terminou o enriquecimento 3074/2114.","drilldowns":{}}
    engine = discover_engine(CONFIG, BASE_DIR)
    if not engine:
        raise RuntimeError("Motor V0.16 não encontrado para gerar indicadores de reuso")
    out = build_trace_insights(engine, payload, line=line or None)
    out["snapshot_id"] = snapshot_id
    out["ready"] = True
    return out


def _verify_xlsx_integrity(path, expected_sheets=None):
    """Fail closed if an XLSX ZIP is incomplete/corrupt before it reaches the user."""
    path = Path(path)
    try:
        with zipfile.ZipFile(path, "r") as zf:
            bad = zf.testzip()
            if bad:
                raise RuntimeError(f"XLSX corrompido no arquivo interno: {bad}")
            names = set(zf.namelist())
            required = {"[Content_Types].xml","xl/workbook.xml","xl/_rels/workbook.xml.rels"}
            missing = sorted(required - names)
            if missing:
                raise RuntimeError("XLSX incompleto: " + ", ".join(missing))
            sheet_xml = [n for n in names if n.startswith("xl/worksheets/sheet") and n.endswith(".xml")]
            if expected_sheets and len(sheet_xml) < len(expected_sheets):
                raise RuntimeError(f"XLSX incompleto: {len(sheet_xml)} planilhas físicas para {len(expected_sheets)} esperadas")
    except zipfile.BadZipFile as exc:
        raise RuntimeError(f"XLSX inválido/corrompido: {exc}") from exc
    return True


def _raw_3028_value(row, key, default=""):
    """Read fields preserved in raw_json without changing the validated SQLite schema."""
    try:
        raw = json.loads(row.get("raw_json") or "{}")
    except Exception:
        raw = {}
    if key in raw and raw.get(key) not in (None, ""):
        return raw.get(key)
    original = raw.get("raw") if isinstance(raw.get("raw"), dict) else {}
    aliases = {
        "defect_return_category":["defectReturnCategory"],
        "defect_detail_category":["defectDetailGategory","defectDetailCategory"],
        "reason_code":["reasonCode"],
        "reason_desc":["reasonDesc"],
        "defect_material":["defectMaterialId"],
        "defect_material_desc":["defectMaterialDesc"],
        "defect_category":["defectCategory"],
        "defect_source":["defectSource"],
        "color":["color"],
        "time_zone":["timeZone"],
        "order_id":["orderId"],
        "user_id":["lotUserId"],
        "repair_time":["repairTime"],
        "defect_receive_time":["defectReceiveTime"],
    }
    for name in aliases.get(key, []):
        if original.get(name) not in (None, ""):
            return original.get(name)
    return default


def _run_deep_trace_job(job_id, cfg):
    job = _job_update(job_id, status="running", stage="3074", progress=5, started_at=utc_now(),
                      message="Preparando rastreabilidade 3074/2114 dos snapshots atuais",
                      stage_progress={"3074":{"current":0,"total":0,"percent":0},"2114":{"current":0,"total":0,"percent":0}},
                      live_events=[], partial_refresh=0)
    _persist_job(job)
    try:
        engine = discover_engine(CONFIG, BASE_DIR)
        if not engine:
            raise RuntimeError("Motor V0.16 nao encontrado para executar 3074/2114")
        lines = normalize_lines(cfg.get("lines") or CONFIG.get("configured_lines") or [])
        defect_codes=[str(x).strip() for x in (cfg.get("defect_codes") or []) if str(x).strip()]
        max_failures=max(0,int(cfg.get("max_failures") or 0))
        max_pcbas=max(0,int(cfg.get("max_pcbas") or 0))
        performance=str(cfg.get("performance") or "balanced").lower()
        snapshot_ids = STORE.latest_snapshot_ids_by_line(lines)
        if not snapshot_ids:
            raise RuntimeError("Nenhum snapshot 3028 atual encontrado para rastrear")
        results = {}
        active_lines=[ln for ln in lines if snapshot_ids.get(ln)]
        total = max(1, len(active_lines))
        shared_query_memo={}
        for idx, line in enumerate(active_lines, 1):
            snapshot_id = snapshot_ids.get(line)
            records = STORE.base_rows(snapshot_id, limit=100000)
            records = _select_trace_records(records, defect_codes=defect_codes, max_failures=max_failures, max_pcbas=max_pcbas)
            if not records:
                results[line]={"snapshot_id":snapshot_id,"skipped":True,"reason":"Nenhuma falha no escopo selecionado"}
                continue
            unique_pcbas=len({str(r.get("pcba_sn") or "").strip() for r in records if str(r.get("pcba_sn") or "").strip()})
            base_progress = 8 + int((idx - 1) * 82 / total)
            job = _job_update(job_id, stage="3074", progress=base_progress,
                              message=f"{line}: 3074 · {unique_pcbas} PCBAs / {len(records)} ocorrencias")
            _persist_job(job)

            last_persist=[0.0]
            def structured_event(ev, _line=line, _idx=idx, _sid=snapshot_id):
                nonlocal job
                # Progressive persistence without reloading the whole dashboard on
                # every poll/component: partial_refresh only changes when new data
                # actually reached SQLite.
                kind=str((ev or {}).get("kind") or "")
                stage=str((ev or {}).get("stage") or "")
                did_ingest=False
                if stage=="3074" and kind=="pcba_result" and (ev or {}).get("item"):
                    try:
                        STORE.ingest_material_reuse(_sid, {"items":[(ev or {}).get("item")]})
                        did_ingest=True
                    except Exception as ie:
                        print('[LIVE 3074 PCBA INGEST]',ie)
                if stage=="3074" and kind=="result" and (ev or {}).get("result"):
                    try:
                        STORE.ingest_material_reuse(_sid, (ev or {}).get("result") or {})
                        STORE.rebuild_cora_index(_sid)
                        did_ingest=True
                    except Exception as ie:
                        print('[LIVE 3074 INGEST]',ie)
                if stage=="2114" and kind=="pcba" and str((ev or {}).get("status"))=="COMPLETE":
                    st=(ev or {}).get("state") or {}
                    if st.get("result"):
                        try:
                            _attach_2114_to_current(_sid,{"pcbas":[{"pcba_sn":(ev or {}).get("pcba"),"status":"COMPLETE","result":st.get("result"),"error":None}]})
                            did_ingest=True
                        except Exception as ie:
                            print('[LIVE 2114 INGEST]',ie)
                job=_job_live_event(job_id,ev,_line,_idx,total,_sid)
                if did_ingest:
                    job=_job_update(job_id, partial_refresh=int(job.get("partial_refresh") or 0)+1)
                now=time.monotonic()
                if did_ingest or now-last_persist[0] >= 1.0:
                    _persist_job(job); last_persist[0]=now

            def cb(text, _line=line, _idx=idx):
                nonlocal job
                stage = "2114" if str(text).startswith("2114") else "3074"
                job = _job_update(job_id, stage=stage, message=f"{_line}: {text}")

            payload = run_deep_v016_records(
                engine, records, source_key=f"snapshot_{snapshot_id}_{line}",
                work_dir=DATA_DIR / "checkpoints", line_callback=cb, event_callback=structured_event,
                max_pcbas=max_pcbas, shared_query_memo=shared_query_memo, performance=performance,
            )
            enrich = enrich_snapshot_from_integrated(snapshot_id, payload)
            results[line] = {"snapshot_id": snapshot_id, "enrichment": enrich,
                             "scope":{"records":len(records),"unique_pcbas":unique_pcbas,"defect_codes":defect_codes,"max_failures":max_failures,"max_pcbas":max_pcbas},
                             "3074": payload.get("trace_3074", {}), "2114": payload.get("trace_2114", {})}
            # The next line can start while this completed line is already visible in the UI.
            job=_job_update(job_id, partial_refresh=int(job.get("partial_refresh") or 0)+1)
            _persist_job(job)
        job = _job_update(job_id, status="done", stage="done", progress=100, finished_at=utc_now(),
                          message="Rastreabilidade 3074/2114 concluida no escopo selecionado",
                          result={"snapshot_ids": snapshot_ids, "lines": active_lines, "deep_trace": results,
                                  "scope":{"defect_codes":defect_codes,"max_failures":max_failures,"max_pcbas":max_pcbas,"performance":performance},
                                  "stages": {"3028":"preserved", "3074":"done", "2114":"done", "3022":"pending_adapter"}})
        _persist_job(job)
    except Exception as exc:
        job = _job_update(job_id, status="error", stage="error", finished_at=utc_now(),
                          message=str(exc), error=f"{type(exc).__name__}: {exc}")
        _persist_job(job)

def _run_sn_lookup_job(job_id, cfg):
    sn = str(cfg.get("sn") or "").strip()
    include_3022 = bool(cfg.get("include_3022", True))
    job = _job_update(job_id, status="running", stage="3074", progress=5, started_at=utc_now(),
                      message=f"Consultando SN {sn}")
    _persist_job(job)
    try:
        engine = ENGINE_DIR or discover_engine(CONFIG, BASE_DIR)
        if not engine:
            raise RuntimeError("Motor V0.16 não encontrado para consulta individual")

        def cb(stage, progress, message):
            nonlocal job
            job = _job_update(job_id, stage=stage, progress=progress, message=message)
            _persist_job(job)

        result = run_individual_lookup(engine, sn, include_3022=include_3022, line_callback=cb)
        job = _job_update(job_id, status="done", stage="done", progress=100, finished_at=utc_now(),
                          message=f"Consulta individual concluída · {sn}", result=result)
        _persist_job(job)
        STORE.record_audit("SN_LOOKUP", workstation=WORKSTATION,
                           details={"sn": sn, "detected_type": result.get("detected_type"),
                                    "summary": result.get("summary"), "warnings": result.get("warnings")})
    except Exception as exc:
        job = _job_update(job_id, status="error", stage="error", finished_at=utc_now(),
                          message=str(exc), error=f"{type(exc).__name__}: {exc}")
        _persist_job(job)


def start_sn_lookup_job(cfg):
    sn = str(cfg.get("sn") or "").strip()
    if not sn:
        raise RuntimeError("Informe ou bipe uma SN")
    return _launch_mes_job("sn_lookup", {"sn": sn, "include_3022": bool(cfg.get("include_3022", True))}, _run_sn_lookup_job)

def start_deep_trace_job(cfg):
    return _launch_mes_job("deep_trace", cfg, _run_deep_trace_job)

def export_snapshot_excel(snapshot_id=None):
    snapshot_id = int(snapshot_id or STORE.latest_snapshot_id() or 0)
    if not snapshot_id:
        raise RuntimeError("Nenhum snapshot para exportar")

    # Prefer the exact report layout already approved as V0.16 whenever deep
    # integrated evidence is available. No MES query is needed for regeneration.
    payload = STORE.get_snapshot_payload(snapshot_id, kinds=["integrated_v016", "integrated_json"])
    engine = discover_engine(CONFIG, BASE_DIR)
    if payload and engine:
        generated = Path(regenerate_v016_excel(engine, payload, DATA_DIR / "tmp"))
        out_dir = DATA_DIR / "exports"
        out_dir.mkdir(parents=True, exist_ok=True)
        target = out_dir / generated.name
        if generated.resolve() != target.resolve():
            shutil.copy2(generated, target)
        _verify_xlsx_integrity(target)
        with STORE.connect() as con:
            con.execute("INSERT INTO artifacts(snapshot_id,artifact_type,path,created_at,metadata_json) VALUES(?,?,?,?,?)",
                        (snapshot_id,"xlsx_v016",str(target),utc_now(),stable_json({"approved_layout":True,"integrity":"OK"})))
        return {"snapshot_id": snapshot_id, "path": str(target), "layout": "V0.16_APROVADO", "integrity":"OK"}

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as exc:
        raise RuntimeError("openpyxl nao instalado") from exc
    out_dir = DATA_DIR / "exports"
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / f"AMES_MONITOR_SNAPSHOT_{snapshot_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "TOP3_FPY"
    headers = ["Linha","FPY %","Check FPY %","Quantity","Total Defect Qty","Falhas snapshot","N atual","Removidas do export","Top 1","Qtd 1","Top 2","Qtd 2","Top 3","Qtd 3"]
    ws.append(headers)
    dash = STORE.dashboard(snapshot_id)
    for line in dash.get("lines") or []:
        m=line.get("metrics") or {}; top=line.get("top3") or []
        def fmt(i):
            if i>=len(top): return ("","")
            return (f"{top[i].get('defect_code') or ''} · {top[i].get('defect_desc') or ''}",top[i].get('qty') or 0)
        t1=fmt(0);t2=fmt(1);t3=fmt(2)
        ws.append([line.get("line"),m.get("fpy"),m.get("check_fpy"),m.get("quantity"),m.get("total_defect_qty"),line.get("defect_rows"),line.get("unresolved"),line.get("removed_from_latest"),t1[0],t1[1],t2[0],t2[1],t3[0],t3[1]])
    base = STORE.base_rows(snapshot_id, limit=100000)
    wsb = wb.create_sheet("BASE_DADOS")
    cols = ["line","pcba_sn","imei","defect_time","defect_code","defect_desc","defect_oper","defect_oper_desc","repair_comment","defect_reason_type","defect_reason_desc","repair_code","repair_desc","product_model","work_shift","manual_or_auto","repair_status_current","defect_type_current","repair_state_current","present_in_3028","defect_key"]
    wsb.append(cols)
    for r in base: wsb.append([r.get(c) for c in cols])
    with STORE.connect() as con:
        hist=[dict(r) for r in con.execute("SELECT * FROM pcba_history WHERE snapshot_id=? ORDER BY pcba_sn,hist_seq",(snapshot_id,)).fetchall()]
        mats=[dict(r) for r in con.execute("SELECT * FROM material_reuse WHERE snapshot_id=? ORDER BY current_pcba_sn,item_type,item_sn",(snapshot_id,)).fetchall()]
    wsh=wb.create_sheet("HIST_PCBA")
    hcols=["pcba_sn","hist_seq","defect_code","defect_desc","defect_oper","defect_location","defect_material_id","repair_status","manual_or_auto","defect_type","defect_type_class"]
    wsh.append(hcols)
    for r in hist:wsh.append([r.get(c) for c in hcols])
    wsm=wb.create_sheet("HIST_MATERIAL")
    mcols=["current_pcba_sn","current_defect_key","item_sn","item_type","usage_status","previous_pcba_count","total_pcba_count_known_now","previous_pcbas_json","active_now_pcbas_json","inactive_now_pcbas_json","bind_time_utc","unbind_time_utc"]
    wsm.append(mcols)
    for r in mats:wsm.append([r.get(c) for c in mcols])
    for sh in wb.worksheets:
        if sh.max_row and sh.max_column:
            for cell in sh[1]:
                cell.fill=PatternFill("solid",fgColor="1F4E78");cell.font=Font(color="FFFFFF",bold=True);cell.alignment=Alignment(vertical="center")
            sh.auto_filter.ref=sh.dimensions
            for idx in range(1,sh.max_column+1):
                width=min(42,max(10,max(len(str(sh.cell(r,idx).value or "")) for r in range(1,min(sh.max_row,200)+1))+2))
                sh.column_dimensions[get_column_letter(idx)].width=width
    wb.save(out)
    _verify_xlsx_integrity(out, wb.sheetnames)
    with STORE.connect() as con:
        con.execute("INSERT INTO artifacts(snapshot_id,artifact_type,path,created_at,metadata_json) VALUES(?,?,?,?,?)",(snapshot_id,"xlsx_monitor",str(out),utc_now(),stable_json({"sheets":wb.sheetnames,"integrity":"OK"})))
    return {"snapshot_id":snapshot_id,"path":str(out),"sheets":wb.sheetnames}

def export_team_excel(lines=None):
    """Team workbook: clean operational tabs + validated V0.16 reuse/correlation drill-downs."""
    team_lines = normalize_lines(lines or CONFIG.get("configured_lines") or ["TAN10101", "TAN10102", "TAN10103"])
    ids = STORE.latest_snapshot_ids_by_line(team_lines)
    unique_ids = sorted(set(ids.values()))
    if not unique_ids:
        raise RuntimeError("Nenhum snapshot das linhas configuradas para exportar")

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("openpyxl nao instalado") from exc

    out_dir = DATA_DIR / "exports"
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / f"AMES_EQUIPE_LINHAS_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "TOP3_FPY"
    headers = ["Linha","Snapshot","Coleta","FPY %","Check FPY %","Quantity","Total Defect Qty","Ocorrencias","N atual","Removidas","Top 1","Qtd 1","Top 2","Qtd 2","Top 3","Qtd 3"]
    ws.append(headers)
    dash = STORE.team_dashboard(team_lines)
    for line in dash.get("lines") or []:
        m = line.get("metrics") or {}
        top = line.get("top3") or []
        def fmt(i):
            if i >= len(top): return ("", "")
            return (f"{top[i].get('defect_code') or ''} · {top[i].get('defect_desc') or ''}", top[i].get('qty') or 0)
        t1,t2,t3=fmt(0),fmt(1),fmt(2)
        ws.append([line.get("line"),line.get("snapshot_id"),line.get("collected_at"),m.get("fpy"),m.get("check_fpy"),m.get("quantity"),m.get("total_defect_qty"),line.get("defect_rows"),line.get("unresolved"),line.get("removed_from_latest"),t1[0],t1[1],t2[0],t2[1],t3[0],t3[1]])

    # 3028: expose normalized columns plus important evidence that was already preserved in raw_json.
    base_cols = ["line","pcba_sn","imei","defect_time","defect_code","defect_desc","defect_oper","defect_oper_desc",
                 "repair_user","repair_comment","defect_reason_type","defect_reason_desc","repair_code","repair_desc",
                 "product_model","work_shift","manual_or_auto","repair_status_current","defect_type_current","repair_state_current",
                 "present_in_3028","defect_key","snapshot_id"]
    raw_cols = ["defect_return_category","defect_detail_category","reason_code","reason_desc","defect_material","defect_material_desc",
                "defect_category","defect_source","color","time_zone","order_id","user_id","defect_receive_time","repair_time"]
    sh=wb.create_sheet("BASE_DADOS")
    sh.append(base_cols+raw_cols)
    defect_rows = STORE.team_dataset_rows("defects", lines=team_lines, limit=100000)
    for row in defect_rows:
        sh.append([row.get(c) for c in base_cols]+[_raw_3028_value(row,c) for c in raw_cols])

    # Audit sheet: the normalized view stays clean, but the exact evidence captured from
    # 3028 remains reachable instead of silently disappearing from the workbook.
    raw_sh=wb.create_sheet("RAW_3028")
    raw_sh.append(["line","pcba_sn","imei","defect_key","snapshot_id","raw_json"])
    for row in defect_rows:
        raw_text=str(row.get("raw_json") or "")
        # XLSX text-cell hard limit is 32,767 chars. Normal 3028 rows are much smaller;
        # keep a marker if a future MES release ever exceeds it.
        if len(raw_text)>32760:
            raw_text=raw_text[:32720]+"...[TRUNCADO NO EXCEL; COMPLETO NO SQLITE]"
        raw_sh.append([row.get("line"),row.get("pcba_sn"),row.get("imei"),row.get("defect_key"),row.get("snapshot_id"),raw_text])

    datasets = [
        ("HIST_PCBA", "pcba_history", ["line","pcba_sn","hist_seq","defect_code","defect_desc","defect_oper","defect_location","defect_material_id","repair_status","manual_or_auto","defect_type","defect_type_class","snapshot_id"]),
        ("HIST_MATERIAL", "material_reuse", ["line","current_pcba_sn","current_defect_key","item_sn","item_type","usage_status","previous_pcba_count","total_pcba_count_known_now","previous_pcbas_json","active_now_pcbas_json","inactive_now_pcbas_json","bind_time_utc","unbind_time_utc","snapshot_id"]),
        ("PROCESSO_3022", "process_events", ["line","pcba_sn","station","operation_code","operation_name","event_time","event_group","source","snapshot_id"]),
    ]
    for sheet_name, dataset, cols in datasets:
        sh = wb.create_sheet(sheet_name); sh.append(cols)
        for row in STORE.team_dataset_rows(dataset, lines=team_lines, limit=100000):
            sh.append([row.get(c) for c in cols])

    # Derived dashboards reuse the same V0.16 correlation engine that drives the approved report.
    dash_rows=[]; pcba_rows=[]; mat_rows=[]; corr_rows=[]; comp_rows=[]
    for line in team_lines:
        sid=ids.get(line)
        if not sid: continue
        try:
            ins=trace_insights(line=line,snapshot_id=sid)
        except Exception as exc:
            dash_rows.append([line,"STATUS","Erro ao gerar indicadores",str(exc),sid]); continue
        if not ins.get("ready"):
            dash_rows.append([line,"STATUS","Rastreabilidade pendente",ins.get("message") or "",sid]); continue
        for group,key in [("GERAL","overview_kpis"),("TAXAS","rate_kpis"),("PCBA","pcba_kpis"),("MATERIAL","material_kpis"),("CORRELACAO","correlation_kpis")]:
            for label,value in ins.get(key) or []:
                dash_rows.append([line,group,label,value,sid])
        for r in ins.get("drilldowns",{}).get("pcba_analysis",[]):
            x=dict(r); x["Linha"]=line; x["Snapshot"]=sid; pcba_rows.append(x)
        for r in ins.get("drilldowns",{}).get("material_analysis",[]):
            x=dict(r); x["Linha"]=line; x["Snapshot"]=sid; mat_rows.append(x)
        for r in ins.get("drilldowns",{}).get("material_history",[]):
            x=dict(r); x["Linha"]=line; x["Snapshot"]=sid; corr_rows.append(x)
        for r in ins.get("component_types") or []:
            x=dict(r); x["Linha"]=line; x["Snapshot"]=sid; comp_rows.append(x)

    sh=wb.create_sheet("DASH_REUSO"); sh.append(["Linha","Grupo","Indicador","Valor","Snapshot"])
    for r in dash_rows: sh.append(r)

    def append_dict_sheet(name, rows, preferred):
        sh=wb.create_sheet(name)
        all_keys=[]
        for k in preferred:
            if any(k in r for r in rows): all_keys.append(k)
        for r in rows:
            for k in r:
                if k not in all_keys: all_keys.append(k)
        if not all_keys: all_keys=preferred
        sh.append(all_keys)
        for r in rows: sh.append([r.get(k) for k in all_keys])
        return sh

    append_dict_sheet("PCBAS_REUSO",pcba_rows,["Linha","Snapshot","PCBA","Modelo","Defect Code","Defect Desc","Defect Time","Uso da PCBA na falha","Usos conhecidos da PCBA","Falhas registradas 2114","Falhas anteriores 2114","Mesma falha anterior","Mesma família anterior","Repair Status atual","Defect Type atual"])
    append_dict_sheet("MATERIAIS_REUSO",mat_rows,["Linha","Snapshot","PCBA atual","Defect Code atual","Defect Desc atual","Tipo material","Material SN","Uso material na falha","Usos conhecidos material","PCBAs desvinculadas únicas","PCBAs com mesma falha","PCBAs com mesma família","Falhas 2114 nas desvinculadas","Evidência histórica"])
    append_dict_sheet("CORRELACOES",corr_rows,["Linha","Snapshot","PCBA atual","Defect Code atual","Tipo material","Material SN","PCBA desvinculada","Classificação desta PCBA","Defect Code antigo","Defect Description antigo","Repair Status antigo","Defect Type antigo","Relação desta ocorrência","Motivo da relação"])
    append_dict_sheet("TIPOS_COMPONENTE",comp_rows,["Linha","Snapshot","Tipo material","Materiais únicos no lote","Reutilizados únicos","2º uso","3º+ uso","Com falha nas PCBAs desvinculadas","Sem falha nas PCBAs desvinculadas","Com mesma falha","Com mesma família","% reuso"])

    for sh in wb.worksheets:
        if sh.max_row and sh.max_column:
            sh.freeze_panes="A2"
            for cell in sh[1]:
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            sh.auto_filter.ref = sh.dimensions
            for idx in range(1, sh.max_column + 1):
                width = min(46, max(10, max(len(str(sh.cell(r,idx).value or "")) for r in range(1, min(sh.max_row,250)+1)) + 2))
                sh.column_dimensions[get_column_letter(idx)].width = width
    wb.save(out)
    _verify_xlsx_integrity(out, wb.sheetnames)
    STORE.record_audit("EXPORT_TEAM_EXCEL", workstation=WORKSTATION, details={"lines": team_lines, "snapshot_ids": ids, "path": str(out), "sheets":wb.sheetnames,"integrity":"OK"})
    return {"snapshot_id": max(unique_ids), "snapshot_ids": ids, "path": str(out), "layout": "TEAM_LATEST_BY_LINE_V0522", "sheets": wb.sheetnames, "integrity":"OK"}


def previous_day_window():
    start_h, start_m = [int(x) for x in CONFIG.get("day_start", "07:00").split(":")[:2]]
    now = datetime.now()
    today_start = now.replace(hour=start_h, minute=start_m, second=0, microsecond=0)
    if now < today_start:
        today_start -= timedelta(days=1)
    start = today_start - timedelta(days=1)
    end = today_start
    return start.isoformat(timespec="minutes"), end.isoformat(timespec="minutes")


def today_window():
    start_h, start_m = [int(x) for x in CONFIG.get("day_start", "07:00").split(":")[:2]]
    now = datetime.now()
    start = now.replace(hour=start_h, minute=start_m, second=0, microsecond=0)
    if now < start:
        start -= timedelta(days=1)
    return start.isoformat(timespec="minutes"), now.isoformat(timespec="minutes")


def monitor_loop():
    while not MONITOR_STOP.wait(5):
        with STORE.connect() as con:
            row = con.execute("SELECT * FROM monitor_profiles WHERE id=1").fetchone()
        if not row or not row["enabled"]:
            continue
        cfg = json.loads(row["config_json"] or "{}")
        interval = max(5, int(row["interval_minutes"] or 30))
        next_run = row["next_run_at"]
        due = True
        if next_run:
            try:
                due = datetime.fromisoformat(next_run) <= datetime.now()
            except ValueError:
                due = True
        if not due:
            continue
        if cfg.get("mode") == "today":
            start_at, end_at = today_window()
        else:
            start_at, end_at = previous_day_window()
        run_cfg = dict(cfg)
        run_cfg.update({"start_at": start_at, "end_at": end_at, "mode": cfg.get("mode", "today")})
        start_job(run_cfg, monitor=True)
        nxt = datetime.now() + timedelta(minutes=interval)
        with STORE.connect() as con:
            con.execute("UPDATE monitor_profiles SET next_run_at=?,updated_at=? WHERE id=1",
                        (nxt.isoformat(timespec="seconds"), utc_now()))


def ensure_monitor_thread():
    global MONITOR_THREAD
    if MONITOR_THREAD and MONITOR_THREAD.is_alive():
        return
    MONITOR_STOP.clear()
    MONITOR_THREAD = threading.Thread(target=monitor_loop, daemon=True)
    MONITOR_THREAD.start()


class Handler(BaseHTTPRequestHandler):
    server_version = "AMESLocalAgent/0.5"

    def log_message(self, fmt, *args):
        print("[HTTP]", fmt % args)

    def _cors(self):
        origin = self.headers.get("Origin")
        active_port = int(getattr(self.server, "server_port", CONFIG.get("port", 8765)))
        allowed = {
            f"http://127.0.0.1:{active_port}",
            f"http://localhost:{active_port}",
        }
        allowed.update(str(x).rstrip("/") for x in (CONFIG.get("allowed_origins") or []) if str(x).strip())
        if origin and origin.rstrip("/") in allowed:
            self.send_header("Access-Control-Allow-Origin", origin)
            self.send_header("Vary", "Origin")
            self.send_header("Access-Control-Allow-Private-Network", "true")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")
        self.send_header("Access-Control-Allow-Methods", "GET,POST,OPTIONS")
        self.send_header("Cache-Control", "no-store")

    def _send(self, obj, status=200):
        body = _json_bytes(obj)
        self.send_response(status)
        self._cors()
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def _body(self):
        length = int(self.headers.get("Content-Length") or 0)
        if length > 60 * 1024 * 1024:
            raise RuntimeError("Payload maior que 60 MB")
        raw = self.rfile.read(length) if length else b"{}"
        return json.loads(raw.decode("utf-8")) if raw else {}

    def _send_static(self, path: Path):
        if not path.exists() or not path.is_file():
            self._send({"error": "static_not_found"}, 404)
            return
        data = path.read_bytes()
        mime = mimetypes.guess_type(str(path))[0] or "application/octet-stream"
        self.send_response(200)
        self._cors()
        self.send_header("Content-Type", mime + ("; charset=utf-8" if mime.startswith("text/") or mime in {"application/javascript"} else ""))
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def _send_download(self, path: Path, download_name: str | None = None):
        path = Path(path)
        if not path.exists() or not path.is_file():
            self._send({"error": "download_not_found", "path": str(path)}, 404)
            return
        data = path.read_bytes()
        name = Path(download_name or path.name).name
        mime = mimetypes.guess_type(str(path))[0] or "application/octet-stream"
        self.send_response(200)
        self._cors()
        self.send_header("Content-Type", mime)
        self.send_header("Content-Disposition", f'attachment; filename="{name}"')
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def do_OPTIONS(self):
        self.send_response(204)
        self._cors()
        self.end_headers()

    def do_GET(self):
        u = urlparse(self.path)
        q = parse_qs(u.query)
        try:
            ui_root = BASE_DIR.parent
            static_map = {
                "/": "ames-offline-v2.html",
                "/ames-offline": "ames-offline-v2.html",
                "/ames-offline-v2.html": "ames-offline-v2.html",
                "/ames-offline-v2.css": "ames-offline-v2.css",
                "/ames-offline-v2.js": "ames-offline-v2.js",
            }
            if u.path in static_map:
                self._send_static(ui_root / static_map[u.path])
                return
            if u.path == "/api/v1/health":
                engine = discover_engine(CONFIG, BASE_DIR)
                self._send({
                    "ok": True,
                    "agent_version": AGENT_VERSION,
                    "mes_scheduler": MES.snapshot(),
                    "db": str(DB_PATH),
                    "engine_found": bool(engine),
                    "engine_dir": str(engine) if engine else None,
                    "chrome_cdp": CONFIG.get("chrome_cdp"),
                    "chrome_cdp_reachable": _tcp_probe("127.0.0.1", 9222, timeout=0.5),
                    "ames_host": CONFIG.get("ames_host", "172.29.185.215"),
                    "ames_port": int(CONFIG.get("ames_port", 80)),
                    "ames_reachable": _tcp_probe(CONFIG.get("ames_host", "172.29.185.215"), int(CONFIG.get("ames_port", 80)), timeout=1.0),
                    "auto_3028_ready": bool((CONFIG.get("auto_3028") or {}).get("ready")),
                    "auto_3022_ready": bool((CONFIG.get("auto_3022") or {}).get("ready")),
                    "latest_snapshot_id": STORE.latest_snapshot_id(),
                    "configured_lines": CONFIG.get("configured_lines") or [],
                    "day_start": CONFIG.get("day_start", "07:00"),
                    "monitor_interval_minutes": int(CONFIG.get("monitor_interval_minutes", 30)),
                    "setup_complete": bool(CONFIG.get("setup_complete")),
                    "latest_backup": STORE.latest_backup_info(BACKUP_DIR),
                    "workstation": WORKSTATION,
                    "readiness": {
                        "agent": True,
                        "engine": bool(engine),
                        "chrome": _tcp_probe("127.0.0.1", 9222, timeout=0.5),
                        "ames_network": _tcp_probe(CONFIG.get("ames_host", "172.29.185.215"), int(CONFIG.get("ames_port", 80)), timeout=1.0),
                        "auto_3028": bool((CONFIG.get("auto_3028") or {}).get("ready")),
                        "auto_3022": bool((CONFIG.get("auto_3022") or {}).get("ready")),
                    },
                })
                return
            if u.path == "/api/v1/export/excel/download":
                sid = (q.get("snapshot_id") or [None])[0]
                team = str((q.get("team") or [""])[0]).lower() in {"1","true","yes"}
                result = export_team_excel(CONFIG.get("configured_lines") or []) if team else export_snapshot_excel(int(sid) if sid else None)
                file_path = Path(result.get("path") or "")
                self._send_download(file_path, file_path.name)
                return
            if u.path == "/api/v1/config":
                self._send(CONFIG)
                return
            if u.path == "/api/v1/share/export":
                self._send(build_share_payload())
                return
            if u.path == "/api/v1/insights":
                line = normalize_line_id((q.get("line") or [None])[0])
                sid = (q.get("snapshot_id") or [None])[0]
                self._send(trace_insights(line=line, snapshot_id=int(sid) if sid else None))
                return
            if u.path == "/api/v1/dashboard":
                sid = (q.get("snapshot_id") or [None])[0]
                self._send(STORE.dashboard(int(sid) if sid else None))
                return
            if u.path == "/api/v1/team-dashboard":
                self._send(STORE.team_dashboard(CONFIG.get("configured_lines") or ["TAN10101","TAN10102","TAN10103"]))
                return
            if u.path == "/api/v1/trends":
                sid = (q.get("snapshot_id") or [None])[0]
                line = normalize_line_id((q.get("line") or [None])[0])
                if not sid and line:
                    sid = STORE.latest_snapshot_id_for_line(line)
                limit = int((q.get("limit") or [80])[0])
                self._send({"rows": STORE.trend_rows(int(sid) if sid else None, line=line, limit=limit)})
                return
            if u.path == "/api/v1/base":
                sid = (q.get("snapshot_id") or [None])[0]
                line = normalize_line_id((q.get("line") or [None])[0])
                dataset = (q.get("dataset") or ["defects"])[0]
                limit = int((q.get("limit") or [5000])[0])
                if sid:
                    rows = STORE.dataset_rows(dataset, int(sid), line=line, limit=limit)
                else:
                    rows = STORE.team_dataset_rows(dataset, lines=CONFIG.get("configured_lines") or [], line=line, limit=limit)
                self._send({"dataset": dataset, "rows": rows})
                return
            if u.path == "/api/v1/base/catalog":
                sid = (q.get("snapshot_id") or [None])[0]
                datasets = STORE.dataset_catalog(int(sid)) if sid else STORE.team_dataset_catalog(CONFIG.get("configured_lines") or [])
                self._send({"datasets": datasets})
                return
            if u.path == "/api/v1/cora/search":
                query = (q.get("q") or [""])[0]
                line = (q.get("line") or [None])[0]
                self._send({"results": STORE.search_cora(query, line=line, limit=30)})
                return
            if u.path == "/api/v1/jobs":
                with JOBS_LOCK:
                    jobs = sorted(JOBS.values(), key=lambda x: x.get("created_at", ""), reverse=True)
                self._send({"jobs": jobs[:50]})
                return
            if u.path.startswith("/api/v1/jobs/"):
                jid = u.path.rsplit("/", 1)[-1]
                with JOBS_LOCK:
                    job = JOBS.get(jid)
                if not job:
                    self._send({"error": "job_not_found"}, 404)
                else:
                    self._send(job)
                return
            if u.path == "/api/v1/monitor":
                with STORE.connect() as con:
                    row = con.execute("SELECT * FROM monitor_profiles WHERE id=1").fetchone()
                self._send(dict(row) if row else {"enabled": 0})
                return
            self._send({"error": "not_found", "path": u.path}, 404)
        except Exception as exc:
            self._send({"error": type(exc).__name__, "message": str(exc)}, 500)

    def do_POST(self):
        u = urlparse(self.path)
        try:
            body = self._body()
            if u.path.startswith("/api/v1/jobs/") and u.path.endswith("/cancel"):
                jid = u.path.split("/")[-2]
                with JOBS_LOCK:
                    token = JOB_CANCEL.get(jid)
                    if token: token.set()
                MES.wake()
                self._send({"ok": bool(token), "id": jid, "cancel_requested": bool(token)})
                return
            if u.path == "/api/v1/chrome/start":
                result = start_dedicated_chrome()
                STORE.record_audit("CHROME_START", workstation=WORKSTATION, details={k:v for k,v in result.items() if k != "profile"})
                self._send(result)
                return
            if u.path == "/api/v1/import/integrated":
                payload = body.get("payload") if isinstance(body, dict) and "payload" in body else body
                source_name = body.get("source_name", "browser-import.json") if isinstance(body, dict) else "browser-import.json"
                self._send(ingest_integrated(payload, source_name))
                return
            if u.path == "/api/v1/import/3028":
                self._send(ingest_3028_path(
                    body.get("path"), mode=body.get("mode", "manual"), start_at=body.get("start_at"),
                    end_at=body.get("end_at"), shift=body.get("shift"), lines=body.get("lines") or [],
                ))
                return
            if u.path == "/api/v1/upload/3028":
                name = Path(str(body.get("filename") or "Detailed_Report_of_Defective_Passthrough_Rate.xlsx")).name
                raw = base64.b64decode(body.get("data_b64") or "", validate=True)
                if not raw or len(raw) > 25 * 1024 * 1024:
                    raise RuntimeError("Arquivo 3028 vazio ou maior que 25 MB")
                imports = DATA_DIR / "imports"
                imports.mkdir(parents=True, exist_ok=True)
                stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
                target = imports / f"{stamp}_{name}"
                target.write_bytes(raw)
                parsed = parse_export(target)  # fail fast before accepting the upload
                self._send({"ok": True, "path": str(target), "records": parsed.get("records_count"),
                            "lines": parsed.get("lines") or []})
                return
            if u.path == "/api/v1/sn-lookup":
                self._send(start_sn_lookup_job({"sn": body.get("sn"), "include_3022": body.get("include_3022", True)}), 202)
                return
            if u.path == "/api/v1/deep-trace":
                self._send(start_deep_trace_job({
                    "lines": body.get("lines") or CONFIG.get("configured_lines") or [],
                    "defect_codes": body.get("defect_codes") or [],
                    "max_failures": body.get("max_failures") or 0,
                    "max_pcbas": body.get("max_pcbas") or 0,
                    "performance": body.get("performance") or CONFIG.get("performance") or "balanced",
                }), 202)
                return
            if u.path == "/api/v1/repairs/refresh":
                self._send(refresh_repairs(body.get("snapshot_id"), body.get("line"), body.get("unresolved_only", False)))
                return
            if u.path == "/api/v1/export/excel":
                self._send(export_snapshot_excel(body.get("snapshot_id")))
                return
            if u.path == "/api/v1/runs":
                cfg = dict(body)
                if cfg.get("preset") == "previous_day":
                    cfg["start_at"], cfg["end_at"] = previous_day_window()
                    cfg["mode"] = "previous_day"
                elif cfg.get("preset") == "today":
                    cfg["start_at"], cfg["end_at"] = today_window()
                    cfg["mode"] = "today"
                self._send(start_job(cfg), 202)
                return
            if u.path == "/api/v1/monitor/start":
                interval = max(5, int(body.get("interval_minutes") or CONFIG.get("monitor_interval_minutes", 30)))
                cfg = dict(body)
                cfg["mode"] = cfg.get("mode", "today")
                nxt = datetime.now().isoformat(timespec="seconds")
                with STORE.connect() as con:
                    con.execute(
                        """UPDATE monitor_profiles SET enabled=1,interval_minutes=?,mode=?,config_json=?,next_run_at=?,updated_at=? WHERE id=1""",
                        (interval, cfg["mode"], stable_json(cfg), nxt, utc_now()),
                    )
                ensure_monitor_thread()
                STORE.record_audit("MONITOR_START", workstation=WORKSTATION, details={"interval_minutes": interval, "mode": cfg.get("mode")})
                self._send({"ok": True, "enabled": True, "interval_minutes": interval})
                return
            if u.path == "/api/v1/monitor/stop":
                with STORE.connect() as con:
                    con.execute("UPDATE monitor_profiles SET enabled=0,next_run_at=NULL,updated_at=? WHERE id=1", (utc_now(),))
                STORE.record_audit("MONITOR_STOP", workstation=WORKSTATION, details={})
                self._send({"ok": True, "enabled": False})
                return
            if u.path == "/api/v1/backup":
                self._send({"ok": True, "backup": create_backup(str(body.get("reason") or "manual"))})
                return
            if u.path == "/api/v1/config":
                allowed = {"configured_lines","line_aliases","day_start","monitor_interval_minutes","performance","chrome_profile_dir",
                           "ames_host","ames_port","ames_start_url","setup_complete","backup_retention","auto_backup_on_start","allowed_origins"}
                changed = {k: v for k, v in body.items() if k in allowed}
                if "configured_lines" in changed:
                    changed["configured_lines"] = normalize_lines(changed.get("configured_lines") or [])
                CONFIG.update(changed)
                write_config(CONFIG)
                STORE.record_audit("CONFIG_UPDATE", workstation=WORKSTATION, details={"keys": sorted(changed.keys())})
                self._send({"ok": True, "config": CONFIG})
                return
            self._send({"error": "not_found", "path": u.path}, 404)
        except Exception as exc:
            traceback.print_exc()
            self._send({"error": type(exc).__name__, "message": str(exc)}, 500)


def main():
    try:
        maybe_backup_on_start()
    except Exception as exc:
        print(f"[BACKUP] aviso: {exc}")
    STORE.record_audit("AGENT_START", workstation=WORKSTATION, details={"agent_version": AGENT_VERSION})
    ap = argparse.ArgumentParser(description="A-MES Local Agent - Central de Trabalho")
    ap.add_argument("--host", default=CONFIG.get("host", "127.0.0.1"))
    ap.add_argument("--port", type=int, default=int(CONFIG.get("port", 8765)))
    args = ap.parse_args()
    ensure_monitor_thread()
    server = ThreadingHTTPServer((args.host, args.port), Handler)
    print(f"A-MES Local Agent v{AGENT_VERSION}")
    print(f"Tela: http://{args.host}:{args.port}/")
    print(f"API: http://{args.host}:{args.port}/api/v1/health")
    print(f"Banco: {DB_PATH}")
    print(f"Motor V0.16: {discover_engine(CONFIG, BASE_DIR) or 'NAO ENCONTRADO'}")
    print("Ctrl+C para encerrar.")
    try:
        server.serve_forever(poll_interval=0.4)
    except KeyboardInterrupt:
        pass
    finally:
        MONITOR_STOP.set()
        server.server_close()


if __name__ == "__main__":
    main()
